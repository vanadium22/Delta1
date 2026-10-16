from datetime import datetime
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from shared_data import DataPaths
from shared_data.io import sha256
from shared_data.readers import LocalData
from index_futures_basis.demo import snapshot
from index_futures_basis.config import load_config
from index_futures_basis.service import BasisService
from index_futures_basis.store import Store


class SharedBasisTests(unittest.TestCase):
    def test_custom_market_directory_does_not_redefine_the_shared_catalog(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "index_futures_basis"
            standard = BasisService(root, "offline")
            config = load_config(root)
            config["data_directory"] = str(standard.paths.root / "custom_market")
            custom = BasisService(root, "offline", config)
            self.assertNotEqual(custom.store.path, standard.store.path)
            again = BasisService(root, "offline")
            self.assertEqual(again.store.path, standard.store.path)

    def test_old_database_and_research_dividend_migrate_without_network_or_original_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            repo = Path(directory)
            root = repo / "index_futures_basis"
            old = Store(root / "data/offline/market.sqlite")
            old.import_raw(snapshot(), {"kind": "offline", "label": "synthetic migration fixture"})
            original_quotes = old.quotes("2026-09-01", "2026-10-08")
            digest = sha256(old.path)
            excel = repo / "research/index_futures_basis/inputs/基差情况_2026-09-18.xlsx"
            excel.parent.mkdir(parents=True)
            book = Workbook()
            sheet = book.active
            sheet.title = "Sheet1"
            sheet["B1"], sheet["G3"] = datetime(2026, 9, 18), "分红"
            for row, code in enumerate([f"{r}0{s}.CFE" for r in ("IF", "IC", "IM", "IH") for s in range(4)], 4):
                sheet.cell(row, 2, code)
                sheet.cell(row, 7, 3.5)
            book.save(excel)
            book.close()
            excel_hash = sha256(excel)
            (root / "config.local.json").write_text(json.dumps({"data_directory": "data", "dividends": {"directories": ["inputs", "input"]}}), encoding="utf-8")
            with patch("index_futures_basis.service.WindProvider", side_effect=AssertionError("network")):
                service = BasisService(root, "offline")
                result = service.query("IH", 4)
                reopened = BasisService(root, "offline")
                self.assertEqual(result["current"], reopened.query("IH", 4)["current"])
            self.assertEqual(service.store.path, repo / "data/market/index_futures/offline/market.sqlite")
            self.assertEqual(service.store.quotes("2026-09-01", "2026-10-08"), original_quotes)
            self.assertEqual(sha256(old.path), digest)
            self.assertEqual(sha256(excel), excel_hash)
            self.assertEqual(sha256(service.paths.dividend_inputs / excel.name), excel_hash)
            self.assertTrue(Path(result["provenance"]["dividends"]["archived_input"]).is_relative_to(repo / "data"))
            self.assertEqual(len(LocalData(DataPaths.for_repo(repo)).daily_quotes("2026-09-01", "2026-10-08", mode="offline")["records"]), len(original_quotes))
