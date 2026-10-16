"""Report freshness, one-session fallback, and the Shanghai 15:05 cutoff."""
import copy
from datetime import datetime, timezone
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from index_futures_basis.config import DEFAULTS, closed_date, load_config, validate_config
from index_futures_basis.service import BasisService
from index_futures_basis.sync import Cancelled
from index_futures_basis.tests.test_sync import FakeProvider


class CutoffTests(unittest.TestCase):
    def test_1505_boundary_and_utc_conversion(self):
        for now, expected in [(datetime(2026, 10, 15, 15, 4, 59), "2026-10-14"),
                              (datetime(2026, 10, 15, 15, 5), "2026-10-15"),
                              (datetime(2026, 10, 15, 7, 5, tzinfo=timezone.utc), "2026-10-15")]:
            self.assertEqual(closed_date(DEFAULTS, now), expected)
        with self.assertRaisesRegex(ValueError, "15:05"):
            validate_config({**DEFAULTS, "close_minute": 4})

    def test_old_default_migrates_but_custom_and_v2_settings_survive(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "config.local.json"
            for settings, expected in [({"version": 1, "close_hour": 18}, (15, 5)),
                                       ({"version": 1, "close_hour": 16}, (16, 0)),
                                       ({"version": 2, "close_hour": 18, "close_minute": 30}, (18, 30))]:
                original = json.dumps(settings)
                path.write_text(original, encoding="utf-8")
                config = load_config(root)
                self.assertEqual((config["close_hour"], config["close_minute"]), expected)
                self.assertEqual(config["version"], 2)
                self.assertEqual(path.read_text(encoding="utf-8"), original)


class ReportUpdateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.service = BasisService(self.root, "wind")
        self.provider = FakeProvider()
        self.seed("2026-10-07")
        self.clock = patch("index_futures_basis.service.closed_date", return_value="2026-10-08")
        self.clock.start()
        self.wind = patch("index_futures_basis.service.WindProvider", return_value=self.provider)
        self.wind.start()

    def tearDown(self):
        self.wind.stop()
        self.clock.stop()
        self.temp.cleanup()

    def seed(self, end):
        raw = copy.deepcopy(self.provider.raw)
        raw["calendar"] = [d for d in raw["calendar"] if d <= end]
        raw["records"] = [r for r in raw["records"] if r["trade_date"] <= end]
        raw["actual_end"] = end
        self.service.store.import_raw(raw, {"kind": "wind", "label": "synthetic test fixture"})

    def export(self):
        inputs = self.root / "inputs"
        inputs.mkdir(exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet["B1"], sheet["G3"] = datetime(2026, 9, 18), "分红"
        for row, code in enumerate([f"{root}0{s}.CFE" for root in ("IF", "IC", "IM", "IH") for s in range(4)], 4):
            sheet.cell(row, 2, code)
            sheet.cell(row, 7, 3.0)
        book.save(inputs / "基差情况_2026-09-18.xlsx")
        book.close()
        # Exercise orchestration and provenance without rendering 25 unchanged charts.
        with patch("index_futures_basis.report.plot_all", return_value={}), \
             patch("index_futures_basis.report.build_docx", return_value={}) as builder:
            folder = Path(self.service.export("report"))
        return json.loads((folder / "run_manifest.json").read_text(encoding="utf-8")), builder

    def test_report_fetches_new_session_then_rechecks_with_zero_downloads(self):
        manifest, builder = self.export()
        self.assertEqual(manifest["actual_end"], "2026-10-08")
        self.assertEqual(manifest["market_update_check"]["state"], "complete")
        self.assertEqual(str(builder.call_args.args[9].date()), "2026-10-08")
        self.assertTrue(self.provider.calls)
        for call in self.provider.calls:
            if call[0] in {"mapping", "series"}:
                self.assertEqual(call[-2:], ("2026-10-08", "2026-10-08"))
        self.provider.calls.clear()
        self.service.prepare_report()
        self.assertEqual(self.provider.calls, [])

    def test_missing_today_spot_falls_back_and_records_exact_reason(self):
        self.provider.missing.add(("000905.SH", "2026-10-08", "close"))
        manifest, builder = self.export()
        check = manifest["market_update_check"]
        self.assertEqual(check["state"], "fallback")
        self.assertEqual(check["target_trade_date"], "2026-10-08")
        self.assertEqual(manifest["actual_end"], "2026-10-07")
        self.assertIn("000905.SH.close", check["reason"])
        self.assertEqual(str(builder.call_args.args[9].date()), "2026-10-07")
        self.assertNotIn(("000905.SH", "2026-10-08", "close"), self.service.store.quotes("2026-10-08", "2026-10-08"))

    def test_failed_calendar_request_uses_verified_previous_session(self):
        self.provider.fail_after = 1
        actual, check = self.service.prepare_report()
        self.assertEqual(actual, "2026-10-07")
        self.assertEqual(check["sync_state"], "failed")
        self.assertIn("模拟网络中断", check["message"])
        self.assertIsNone(check["target_trade_date"])

    def test_fallback_does_not_skip_an_incomplete_previous_session(self):
        self.provider.missing.update({("000905.SH", "2026-10-08", "close"), ("000905.SH", "2026-10-07", "close")})
        with self.service.store.connection() as con:
            con.execute("DELETE FROM quotes WHERE day='2026-10-07' AND code='000905.SH'")
        with self.assertRaisesRegex(ValueError, "前一交易日2026-10-07的数据也不完整"):
            self.service.export("report")
        self.assertFalse(self.service.exports_root.exists())

    def test_weekend_uses_latest_exchange_session_without_fallback_warning(self):
        self.seed("2026-10-08")
        self.service.store.save_calendar("2026-10-09", "2026-10-11", set())
        with patch("index_futures_basis.service.closed_date", return_value="2026-10-11"):
            actual, check = self.service.prepare_report()
        self.assertEqual(actual, "2026-10-08")
        self.assertEqual(check["state"], "complete")
        self.assertEqual(check["requested_cutoff"], "2026-10-11")
        self.assertEqual(self.provider.calls, [])

    def test_explicit_history_stays_selected_and_does_not_fetch_today(self):
        actual, check = self.service.prepare_report("2026-10-06")
        self.assertEqual(actual, "2026-10-06")
        self.assertEqual(check["state"], "complete")
        self.assertEqual(self.provider.calls, [])
        self.provider.missing.add(("000905.SH", "2026-10-08", "close"))
        with self.assertRaisesRegex(ValueError, "指定历史日期不自动改期"):
            self.service.prepare_report("2026-10-08")

    def test_cancellation_and_unknown_previous_calendar_do_not_export(self):
        with patch.object(self.service, "sync", side_effect=Cancelled("用户取消")):
            with self.assertRaises(Cancelled):
                self.service.export("report")
        self.provider.fail_after = 1
        with self.service.store.connection() as con:
            con.execute("DELETE FROM calendar WHERE day='2026-10-07'")
        with self.assertRaisesRegex(ValueError, "无法确认前一交易日"):
            self.service.export("report")
        self.assertFalse(self.service.exports_root.exists())

    def test_offline_report_does_not_call_wind_or_freshness_check(self):
        self.service = BasisService(self.root, "offline")
        self.seed("2026-10-07")
        with patch.object(self.service, "prepare_report", side_effect=AssertionError("unexpected update")):
            manifest, _ = self.export()
        self.assertEqual(manifest["actual_end"], "2026-10-07")
        self.assertNotIn("market_update_check", manifest)
        self.assertEqual(self.provider.calls, [])
