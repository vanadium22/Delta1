"""Independent invariants for distribution windows, dividend refresh and pairing."""
from pathlib import Path
import tempfile
import unittest

import numpy as np
import pandas as pd
from openpyxl import Workbook

from index_futures_basis.distributions import annualized_distributions, cross_index_basis, cross_index_distributions, distribution_rows
from index_futures_basis.analytics import attach_dividends, read_dividends
from index_futures_basis.tests.test_analytics import sample_panel, dividend_fixture


class DistributionTests(unittest.TestCase):
    def test_calendar_month_boundary_and_current_exclusion(self):
        end = pd.Timestamp("2026-10-08")
        g = pd.DataFrame({"trade_date": pd.to_datetime(["2025-10-01", "2026-07-07", "2026-07-08",
                                                       "2026-08-01", "2026-10-07", "2026-10-08", "2026-10-09"]),
                          "value": [-1000, -999, 0, 10, 20, 10000, 99999]})
        row = distribution_rows(g, end, "value")[0]
        self.assertEqual(row["n"], 3)
        self.assertEqual(row["p05"], 1)
        self.assertEqual(row["p25"], 5)
        self.assertEqual(row["p50"], 10)
        self.assertEqual(row["p75"], 15)
        self.assertEqual(row["p95"], 19)
        self.assertEqual(row["percentile"], 100)
        self.assertEqual(row["window_start"], pd.Timestamp("2026-07-08"))

    def test_ties_half_weight_and_incomplete_history(self):
        g = pd.DataFrame({"trade_date": pd.to_datetime(["2026-06-01", "2026-07-08", "2026-08-01", "2026-10-07", "2026-10-08"]),
                          "value": [-999, 0, 10, 10, 10]})
        rows = distribution_rows(g, pd.Timestamp("2026-10-08"), "value")
        self.assertAlmostEqual(rows[0]["percentile"], 100 * 2 / 3)
        self.assertFalse(rows[2]["complete"])
        self.assertTrue(np.isnan(rows[2]["p50"]))
        g.loc[2, "value"] = np.nan
        self.assertFalse(distribution_rows(g, pd.Timestamp("2026-10-08"), "value")[0]["complete"])

    def test_annual_quantiles_never_use_display_fill(self):
        p = sample_panel()
        p.loc[p.slot.eq(1), "days_to_expiry"] = [5, 1, 0]
        a = attach_dividends(p, dividend_fixture(p), p.trade_date.min())
        self.assertTrue(a.annualized_plot_filled.any())
        q = annualized_distributions(a, a.trade_date.max())
        self.assertTrue(q.loc[q.slot.eq(1), "current"].isna().all())
        self.assertTrue(q.loc[q.slot.eq(1), "n"].eq(1).all())

    def cross_fixture(self):
        p = sample_panel()
        legs = []
        for root, basis in [("IM", -30), ("IC", -10)]:
            g = p.copy()
            g["root"] = root
            g["instrument_id"] = g.instrument_id.str.replace("IF", root, regex=False)
            g["continuous_code"] = g.continuous_code.str.replace("IF", root, regex=False)
            g["basis_points"] = basis
            g["roll_day"] = False
            legs.append(g)
        return pd.concat(legs, ignore_index=True)

    def test_cross_is_basis_difference_and_exact_expiry_pairing(self):
        p = self.cross_fixture()
        cross = cross_index_basis(p)
        self.assertTrue(cross.basis_difference_points.eq(-20).all())
        self.assertEqual(len(cross), 12)
        self.assertEqual(len(cross_index_distributions(cross, cross.trade_date.max())), 12)
        p.loc[p.root.eq("IC") & p.slot.eq(1), "expiry_date"] += pd.Timedelta(days=1)
        with self.assertRaisesRegex(ValueError, "到期日不一致"):
            cross_index_basis(p)

    def test_cross_missing_leg_not_silently_dropped(self):
        with self.assertRaisesRegex(ValueError, "缺少同日同期限"):
            cross_index_basis(self.cross_fixture().iloc[1:])


class LatestDividendTests(unittest.TestCase):
    @staticmethod
    def workbook(path, date, value):
        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet["B1"] = pd.Timestamp(date).to_pydatetime()
        sheet["G3"] = "分红"
        for i in range(4):
            sheet.cell(i+4, 2, f"IF0{i}.CFE")
            sheet.cell(i+4, 7, value+i)
        book.save(path)
        book.close()

    def test_select_by_internal_date_and_reread_updated_values(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            old = folder / "指数基差情况_2099-01-01.xlsx"
            new = folder / "指数基差情况_2000-01-01.xlsx"
            self.workbook(old, "2026-09-18", 1)
            self.workbook(new, "2026-10-15", 2)
            config = {"selection": "latest_snapshot", "layout": "snapshot_continuous", "directory": "."}
            d, info = read_dividends(config, folder)
            self.assertEqual(info["snapshot_date"], "2026-10-15")
            self.assertEqual(d.dividend_points.iloc[0], 2)
            self.workbook(new, "2026-10-15", 3)
            d, updated = read_dividends(config, folder)
            self.assertEqual(d.dividend_points.iloc[0], 3)
            self.assertNotEqual(info["sha256"], updated["sha256"])
            self.workbook(old, "2026-10-15", 99)
            with self.assertRaisesRegex(ValueError, "内容不同"):
                read_dividends(config, folder)

    def test_input_alias_filename_and_same_day_priority(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            (folder / "input").mkdir()
            (folder / "inputs").mkdir()
            self.workbook(folder / "inputs/指数基差情况_2026-09-18.xlsx", "2026-09-18", 1)
            new = folder / "input/基差情况_2026-10-15.xlsx"
            self.workbook(new, "2026-10-15", 2)
            config = {"selection": "latest_snapshot", "layout": "snapshot_continuous", "directories": ["input", "inputs"]}
            d, info = read_dividends(config, folder)
            self.assertEqual(Path(info["path"]), new)
            self.assertEqual(d.dividend_points.iloc[0], 2)
            # An old-folder copy for the same date must not override the new input.
            self.workbook(folder / "inputs/指数基差情况_2026-10-15.xlsx", "2026-10-15", 99)
            d, info = read_dividends(config, folder)
            self.assertEqual(d.dividend_points.iloc[0], 2)
            self.workbook(folder / "input/基差情况_2026-10-16.xlsx", "2026-10-16", 3)
            d, info = read_dividends(config, folder)
            self.assertEqual(info["snapshot_date"], "2026-10-16")
            self.assertEqual(d.dividend_points.iloc[0], 3)

    def test_newer_legacy_date_wins_and_excel_lock_files_are_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            (folder / "input").mkdir()
            (folder / "inputs").mkdir()
            self.workbook(folder / "input/基差情况_2026-10-14.xlsx", "2026-10-14", 1)
            self.workbook(folder / "inputs/指数基差情况_2026-10-15.xlsx", "2026-10-15", 2)
            (folder / "input/~$基差情况_2026-10-16.xlsx").write_bytes(b"not a workbook")
            (folder / "input/unrelated.xlsx").write_bytes(b"not a workbook")
            config = {"selection": "latest_snapshot", "layout": "snapshot_continuous", "directories": ["input", "inputs"]}
            d, info = read_dividends(config, folder)
            self.assertEqual(info["snapshot_date"], "2026-10-15")
            self.assertEqual(d.dividend_points.iloc[0], 2)


if __name__ == "__main__":
    unittest.main()
