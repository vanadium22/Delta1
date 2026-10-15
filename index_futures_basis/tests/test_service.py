import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd

from index_futures_basis.config import DEFAULTS
from index_futures_basis.demo import snapshot
from index_futures_basis.service import BasisService


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.service = BasisService(self.root, "demo")

    def tearDown(self):
        self.temp.cleanup()

    def test_query_statistics_agree_with_report_engine(self):
        q = self.service.query("IC", 4)
        a = self.service.analyze()
        expected = a.history[a.history.root.eq("IC") & a.history.slot.eq(4)]
        np.testing.assert_allclose(q["statistics"]["mean"], expected["mean"])
        np.testing.assert_allclose(q["statistics"].p50, expected["median"])
        np.testing.assert_allclose(q["statistics"].percentile, expected.percentile)
        self.assertEqual(q["current_contract"], "IC2703.CFE")
        self.assertIs(a, self.service.analyze())

    def test_no_network_in_demo_queries_and_exports(self):
        with patch("index_futures_basis.service.WindProvider", side_effect=AssertionError("network")):
            for metric in ["basis_points", "raw_basis_points", "im_ic", "fixed_60d", "calendar_spread", "oi_change"]:
                q = self.service.query("IC", 4, metric)
                self.assertEqual(len(q["statistics"]), 3)
            folder = Path(self.service.export("data"))
            self.assertEqual(len(list(folder.glob("*.csv"))), 8)
            manifest = json.loads((folder / "run_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["mode"], "demo")
        with self.assertRaisesRegex(ValueError, "不访问 Wind"):
            self.service.sync()

    def test_single_chart_export_does_not_create_other_charts_or_word(self):
        folder = Path(self.service.export("figures", ["IC_年化基差分位蜡烛图"]))
        self.assertEqual([p.name for p in (folder / "figures").glob("*.png")], ["IC_年化基差分位蜡烛图.png"])
        self.assertFalse(list(folder.glob("*.docx")))

    def test_formal_and_demo_settings_and_stores_are_separate(self):
        self.service.save_settings({"auto_sync": False})
        live = BasisService(self.root, "wind")
        self.assertTrue(live.config["auto_sync"])
        self.assertIsNone(live.status()["latest_complete"])
        self.assertNotEqual(live.store.path, self.service.store.path)

    def test_excel_refresh_invalidates_derived_results_without_touching_market(self):
        service = BasisService(self.root, "offline")
        service.store.import_raw(snapshot(), {"kind": "offline"})
        (self.root / "inputs").mkdir()
        from openpyxl import Workbook
        def workbook(value):
            book = Workbook()
            sheet = book.active
            sheet.title = "Sheet1"
            sheet["B1"] = pd.Timestamp("2026-09-18").to_pydatetime()
            sheet["G3"] = "分红"
            for i, code in enumerate([f"{root}0{s}.CFE" for root in ("IF", "IC", "IM") for s in range(4)], 4):
                sheet.cell(i, 2, code)
                sheet.cell(i, 7, value)
            book.save(self.root / "inputs/基差情况_2026-09-18.xlsx")
            book.close()
        workbook(5)
        a = service.analyze()
        market = service.store.quotes("2025-07-01", "2026-10-08")
        workbook(15)
        b = service.analyze()
        self.assertNotEqual(a.fingerprint, b.fingerprint)
        np.testing.assert_allclose(b.panel.basis_points - a.panel.basis_points, 10)
        self.assertEqual(service.store.quotes("2025-07-01", "2026-10-08"), market)
        self.assertEqual(len(list((service.data_root / "dividends").glob("*.xlsx"))), 2)

    def test_asof_and_interior_gap_are_not_silently_relabelled(self):
        with self.assertRaises(ValueError):
            self.service.query(as_of="2026-10-09")
        with self.service.store.connection() as con:
            con.execute("DELETE FROM quotes WHERE day='2026-09-01' AND code='000905.SH'")
        with self.assertRaises(ValueError):
            self.service.query()
