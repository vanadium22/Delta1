"""Keep usable evidence when market refresh or individual local inputs fail."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from docx import Document
from openpyxl import Workbook
from PIL import Image

from index_futures_basis.report import chart_catalog
from index_futures_basis.service import BasisService


class PartialExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def word(self, service, **kwargs):
        image = self.root / "placeholder.png"
        Image.new("RGB", (120, 60), "white").save(image)
        charts = {name: image for name in chart_catalog() if not name.startswith(tuple(f"IM0{i}_IC0{i}" for i in range(4)))}
        with patch("index_futures_basis.report.plot_all", return_value=charts):
            folder = Path(service.export("report", **kwargs))
        return Document(next(folder.glob("*.docx"))), json.loads((folder / "run_manifest.json").read_text(encoding="utf-8"))

    def test_empty_cache_still_generates_word_and_blank_chart_without_network(self):
        service = BasisService(self.root, "wind")
        with patch("index_futures_basis.service.WindProvider", side_effect=AssertionError("network")):
            doc, manifest = self.word(service, as_of="2026-10-15", refresh=False)
            folder = Path(service.export("figures", ["跨期03-02_绝对基差"], as_of="2026-10-15"))
        self.assertEqual(len(doc.inline_shapes), 46)
        self.assertEqual(len(doc.tables), 17)
        self.assertIsNone(manifest["data_availability"]["latest_local_date"])
        self.assertTrue(manifest["local_only"])
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("缺失项留空", text)
        self.assertIn("分红不可用", text)
        self.assertEqual(doc.tables[0].rows[1].cells[1].text, "—")
        self.assertEqual(len(list((folder / "figures").glob("*.png"))), 1)

    def test_sparse_fields_are_independent_and_missing_totals_never_become_zero(self):
        service = BasisService(self.root, "demo")
        with service.store.connection() as con:
            con.execute("DELETE FROM quotes WHERE code='000905.SH' AND day='2026-10-08'")
            con.execute("DELETE FROM quotes WHERE field='oi' AND code LIKE 'IH%'")
            con.execute("DELETE FROM contracts WHERE code LIKE 'IF%'")
            con.execute("DELETE FROM mapping WHERE root='IM' AND slot=4")
        a = service.analyze(allow_partial=True)
        now = a.panel[a.panel.trade_date.eq(a.end)]
        self.assertTrue(now[now.root.eq("IF")].basis_points.notna().all())
        self.assertTrue(now[now.root.eq("IF")].annualized_basis_pct.isna().all())
        self.assertTrue(a.totals[a.totals.root.eq("IH")].open_interest.isna().all())
        self.assertTrue(now[now.root.eq("IH")].basis_points.notna().all())
        ic = a.calendar_annualized[a.calendar_annualized.root.eq("IC") & a.calendar_annualized.trade_date.eq(a.end)]
        self.assertTrue(ic.spread_points.notna().all())
        self.assertTrue(ic.annualized_basis_pct.isna().all())
        self.assertTrue(a.cross[a.cross.slot.eq(4)].basis_difference_points.isna().all())
        self.word(service, refresh=False)

    def test_timeout_uses_old_local_cache_and_explicit_future_cutoff_has_no_carried_current(self):
        service = BasisService(self.root, "demo")
        service.mode = "wind"  # Keep the synthetic fixture store; no production files.
        with patch.object(service, "sync", side_effect=TimeoutError("Wind无响应")), \
             patch("index_futures_basis.service.closed_date", return_value="2026-10-16"):
            doc, manifest = self.word(service)
            self.assertEqual(manifest["actual_end"], "2026-10-08")
            self.assertIn("Wind无响应", manifest["market_update_check"]["reason"])
            a = service.analyze("2026-10-15", allow_partial=True)
        self.assertTrue(a.panel.loc[a.panel.trade_date.eq(a.end), "futures_close"].isna().all())
        self.assertEqual(a.provenance["data_availability"]["latest_local_date"], "2026-10-08")
        self.assertIn("已使用本地可用数据", "\n".join(p.text for p in doc.paragraphs))

    def test_wind_empty_can_use_imported_offline_cache_without_mixing_demo(self):
        fixture = BasisService(self.root / "fixture", "demo")
        service = BasisService(self.root, "wind")
        offline = BasisService(self.root, "offline")
        from index_futures_basis.demo import snapshot
        offline.store.import_raw(snapshot(), {"kind": "test offline cache"})
        a = service.analyze(allow_partial=True)
        self.assertEqual(a.provenance["data_availability"]["source_database"], str(offline.store.path))
        self.assertNotEqual(a.provenance["data_availability"]["source_database"], str(fixture.store.path))
        self.assertEqual(service.store.latest_available("2026-10-16"), None)

    def test_partial_dividend_cells_and_archived_snapshot_are_usable(self):
        service = BasisService(self.root, "offline")
        from index_futures_basis.demo import snapshot
        service.store.import_raw(snapshot(), {"kind": "test"})
        folder = service.dividend_directories()[0]
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / "基差情况_2026-09-18.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet["B1"], sheet["G3"] = "2026-09-18", "分红"
        for i, code in enumerate([f"{r}0{s}.CFE" for r in ("IF", "IC", "IM", "IH") for s in range(4)], 4):
            sheet.cell(i, 2, code)
            sheet.cell(i, 7, None if code == "IH03.CFE" else 5)
        book.save(path)
        book.close()
        a = service.analyze(allow_partial=True)
        self.assertTrue(a.panel.loc[a.panel.continuous_code.eq("IH03.CFE"), "basis_points"].isna().all())
        self.assertTrue(a.panel.loc[a.panel.root.eq("IF"), "basis_points"].notna().all())
        path.write_bytes(b"damaged workbook")  # Only this test's temporary workbook.
        b = service.analyze(allow_partial=True)
        self.assertIn("已归档快照", b.provenance["dividends"]["fallback_reason"])
        np.testing.assert_allclose(a.panel.basis_points, b.panel.basis_points, equal_nan=True)

    def test_complete_data_has_same_values_as_strict_analysis(self):
        service = BasisService(self.root, "demo")
        strict, partial = service.analyze(), service.analyze(allow_partial=True)
        for column in ("basis_points", "annualized_basis_pct", "oi_change"):
            keys = ["trade_date", "root", "slot"]
            np.testing.assert_allclose(strict.panel.sort_values(keys)[column], partial.panel.sort_values(keys)[column], equal_nan=True)
        for name in ("fixed_quantiles", "calendar_quantiles", "calendar_point_quantiles"):
            keys = ["root", "months"] + (["pair"] if "calendar" in name else [])
            np.testing.assert_allclose(getattr(strict, name).sort_values(keys).percentile,
                                       getattr(partial, name).sort_values(keys).percentile, equal_nan=True)


if __name__ == "__main__":
    unittest.main()
