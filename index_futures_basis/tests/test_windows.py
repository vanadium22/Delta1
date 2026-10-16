import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from docx import Document
from PIL import Image

from index_futures_basis.report import build_docx, chart_catalog, plot_all
from index_futures_basis.service import BasisService
from index_futures_basis.windows import normalize_windows


class WindowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.service = BasisService(self.root, "demo")

    def tearDown(self):
        self.temp.cleanup()

    def test_invalid_settings_are_rejected_before_network_or_output(self):
        for value in ([], [0], [37], [3, 3], [True], [1.5], [1, 2, 3, 4], "", "1,,3", "3.0"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                normalize_windows({"fixed_60_quantile_months": value})
        with self.assertRaises(ValueError):
            normalize_windows({"fixed_90_months": [3]})
        live = BasisService(self.root, "wind")
        with patch.object(live, "prepare_report", side_effect=AssertionError("network")), self.assertRaises(ValueError):
            live.export(windows={"series_months": 0})
        self.assertFalse(live.exports_root.exists())

    def test_custom_natural_month_quantiles_use_real_samples_and_cache_key(self):
        default = self.service.analyze()
        raw = self.service.store.quotes("2025-01-01", "2026-12-31")
        windows = normalize_windows({"fixed_60_quantile_months": [2, 4, 9], "fixed_90_quantile_months": [1, 5, 10],
            "calendar_quantile_months": [2, 8], "cross_quantile_months": [4], "annual_quantile_months": [1, 4, 9]})
        with patch("index_futures_basis.service.WindProvider", side_effect=AssertionError("network")):
            a = self.service.analyze(windows=windows)
        self.assertNotEqual(a.fingerprint, default.fingerprint)
        self.assertIs(a, self.service.analyze(windows=windows))
        for series, quantiles in ((a.fixed, a.fixed_quantiles), (a.fixed90, a.fixed90_quantiles)):
            for row in quantiles.itertuples():
                mask = (series.root.eq(row.root) & series.trade_date.ge(a.end - pd.DateOffset(months=row.months))
                        & series.trade_date.lt(a.end) & series.interpolable)
                history = series.loc[mask, "annualized_basis_pct"].dropna()
                self.assertEqual(row.n, len(history))
                self.assertAlmostEqual(row.p50, np.median(history))
                expected = ((history < row.current).sum() + .5 * (history == row.current).sum()) / len(history) * 100
                self.assertAlmostEqual(row.percentile, expected)
                self.assertLess(row.sample_end, a.end)
        self.assertEqual(set(a.calendar_quantiles.months), {2, 8})
        self.assertEqual(set(a.cross_quantiles.months), {4})
        self.assertEqual(set(a.history.months), {1, 4, 9})
        self.assertEqual(self.service.analyze().fingerprint, default.fingerprint)
        self.assertEqual(self.service.store.quotes("2025-01-01", "2026-12-31"), raw)

    def test_profiles_persist_independently_and_extend_incremental_history(self):
        self.service.save_output_windows("figures", {"fixed_60_quantile_months": [3, 18, 24]})
        self.service.save_output_windows("report", {"fixed_90_quantile_months": [1, 6, 18]})
        loaded = BasisService(self.root, "demo")
        self.assertEqual(loaded.output_windows("figures")["fixed_60_quantile_months"], [3, 18, 24])
        self.assertEqual(loaded.output_windows("report")["fixed_60_quantile_months"], [3, 6, 12])
        self.assertEqual(loaded.output_windows("report")["fixed_90_quantile_months"], [1, 6, 18])
        start, end = loaded.window()
        self.assertEqual(start, str((pd.Timestamp(end) - pd.DateOffset(months=24) - pd.Timedelta(days=15)).date()))
        self.assertEqual(BasisService(self.root, "wind").output_windows("figures"), normalize_windows())

    def test_long_window_missing_history_is_not_reported_as_complete(self):
        a = self.service.analyze(windows={"fixed_60_quantile_months": [3, 36], "fixed_90_quantile_months": [3, 36]})
        for q in (a.fixed_quantiles, a.fixed90_quantiles):
            self.assertTrue(q.loc[q.months.eq(3), "complete"].all())
            self.assertFalse(q.loc[q.months.eq(36), "complete"].any())
            self.assertTrue(q.loc[q.months.eq(36), "percentile"].isna().all())
        with self.assertRaisesRegex(ValueError, "36个月曲线历史不足"):
            self.service.export("figures", ["IC_合并年化基差"], windows={"series_months": 36})

    def test_curve_windows_reach_renderer_and_filenames(self):
        import matplotlib.dates as mdates
        from matplotlib.figure import Figure
        windows = normalize_windows({"series_months": 2, "fixed_60_months": [6, 2], "calendar_months": 4})
        a = self.service.analyze(windows=windows)
        seen = {}
        def capture(fig, filename, **kwargs):
            seen[Path(filename).stem] = [ax.get_xlim() for ax in fig.axes]
        chosen = ["IC_合并年化基差", "60天固定期限_近6个月", "跨期02-00_年化基差"]
        with patch.object(Figure, "savefig", capture):
            plot_all(a.panel, a.fixed, a.spread, a.start, a.end, self.root / "plots", ["IF", "IC", "IM", "IH"],
                     selected=chosen, windows=windows)
        for name, months in zip(chosen, [2, 6, 4]):
            self.assertAlmostEqual(seen[name][0][0], mdates.date2num(a.end - pd.DateOffset(months=months)))
        self.assertIn("60天固定期限_近2个月", chart_catalog(windows))
        self.assertNotIn("60天固定期限_近1个月", chart_catalog(windows))

    def test_custom_word_captions_tables_csv_and_provenance_agree(self):
        windows = normalize_windows({"series_months": 2, "fixed_60_months": [6, 2], "fixed_90_months": [4, 1],
            "calendar_months": 4, "annual_quantile_months": [1, 4, 9], "fixed_60_quantile_months": [1, 2, 9],
            "fixed_90_quantile_months": [2, 5], "calendar_quantile_months": [2, 8], "cross_quantile_months": [4]})
        a = self.service.analyze(windows=windows)
        image = self.root / "placeholder.png"
        Image.new("RGB", (4, 4), "white").save(image)
        class Charts(dict):
            def __getitem__(self, key):
                self[key] = image
                return image
        target = self.root / "custom.docx"
        build_docx(Path(__file__).parents[1] / "assets/report_template.docx", target,
            a.panel, a.fixed, a.spread, a.totals, a.history, Charts(), a.start, a.end,
            ["IF", "IC", "IM", "IH"], a.provenance, annual_quantiles=a.annual_quantiles,
            cross=a.cross, cross_quantiles=a.cross_quantiles,
            terms=(a.fixed90, a.fixed90_quantiles, a.calendar_annualized, a.calendar_quantiles),
            windows=windows, fixed_quantiles=a.fixed_quantiles)
        doc = Document(target)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("60日固定期限年化基差1个月2个月9个月分位", text)
        self.assertIn("90日固定期限年化基差2个月5个月分位", text)
        self.assertIn("近4个月", text)
        self.assertNotIn("3个月6个月12个月分位", text)
        self.assertIn("1月当前", doc.tables[-6].rows[0].cells[2].text)
        self.assertIn("2月当前", doc.tables[-5].rows[0].cells[2].text)
        self.assertEqual(len(doc.tables[-5].rows[0].cells), 6)
        self.assertIn("03-02跨期绝对基差2个月8个月分位", text)
        self.assertIn("2月当前", doc.tables[-1].rows[0].cells[2].text)
        self.assertEqual(doc.tables[-1].rows[0].cells[1].text, "当前\n点")
        output = Path(self.service.export("data", windows=windows))
        manifest = json.loads((output / "run_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["configuration"]["windows"], windows)
        self.assertEqual(set(pd.read_csv(output / "fixed_60d_quantiles.csv").months), {1, 2, 9})
        self.assertEqual(set(pd.read_csv(output / "fixed_90d_quantiles.csv").months), {2, 5})
        self.assertEqual(set(pd.read_csv(output / "calendar_absolute_quantiles.csv").months), {2, 8})
