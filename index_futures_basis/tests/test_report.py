"""Verify the exported front tables retain the analyzed date and inputs."""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
import numpy as np
import pandas as pd
from PIL import Image

from index_futures_basis.report import build_docx
from index_futures_basis.service import BasisService


class ReportTests(unittest.TestCase):
    def test_front_tables_show_actual_date_prices_expiry_and_unfilled_annualization(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            analysis = BasisService(root, "demo").analyze()
            panel = analysis.panel.copy()
            mask = panel.trade_date.eq(analysis.end) & panel.root.eq("IC") & panel.slot.eq(1)
            panel.loc[mask, "annualized_basis_pct"] = np.nan
            panel.loc[mask, "annualized_plot_pct"] = -99.99
            panel.loc[mask, "days_to_expiry"] = 1
            panel.loc[mask, "expiry_date"] = analysis.end + pd.Timedelta(days=1)
            provenance = {**analysis.provenance, "requested_end": "2026-10-15",
                          "generated_at": "2026-10-15T20:00:00+08:00",
                          "market_update_check": {"state": "fallback", "reason": "指数收盘价缺失"}}
            image = root / "chart.png"
            Image.new("RGB", (4, 4), "white").save(image)
            class Charts(dict):
                def __getitem__(self, key):
                    self[key] = image
                    return image
            output = root / "report.docx"
            with patch("index_futures_basis.service.WindProvider", side_effect=AssertionError("network")):
                counts = build_docx(Path(__file__).parents[1] / "assets/report_template.docx", output,
                    panel, analysis.fixed, analysis.spread, analysis.totals, analysis.history, Charts(),
                    analysis.start, analysis.end, ["IF", "IC", "IM"], provenance, preview=True,
                    annual_quantiles=analysis.annual_quantiles, cross=analysis.cross,
                    cross_quantiles=analysis.cross_quantiles)
            doc = Document(output)
            front = doc.tables[:3]
            self.assertEqual(counts["tables"], 9)
            self.assertEqual(counts["figures"], 25)
            text = "\n".join(p.text for p in doc.paragraphs[:12])
            self.assertIn("行情日期：2026-10-08", text)
            self.assertIn("更新未成功：目标截止2026-10-15", text)
            self.assertIn("已回退至前一交易日2026-10-08", text)
            self.assertIn("指数收盘价缺失", text)
            ic = front[1].rows[1].cells
            latest = panel.loc[mask].iloc[0]
            self.assertEqual(ic[0].text, "IC00/IC2610")
            self.assertEqual(ic[1].text, f"{latest.futures_close:,.2f}")
            self.assertEqual(ic[2].text, "2026-10-09")
            self.assertEqual(ic[3].text, "1")
            self.assertEqual(ic[4].text, f"{latest.dividend_points:,.5f}")
            self.assertEqual(ic[6].text, "—")
            captions = [p.text for p in doc.paragraphs if "现货收盘值" in p.text]
            self.assertTrue(any(f"{latest.spot_close:,.4f}" in t for t in captions))
