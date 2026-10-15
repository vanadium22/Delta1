"""Validated dividend-adjusted daily research data. No display fills enter statistics."""
from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd

ROOTS = {"IF": "000300.SH", "IC": "000905.SH", "IM": "000852.SH", "IH": "000016.SH"}
NAMES = {"IF": "沪深300", "IC": "中证500", "IM": "中证1000", "IH": "上证50"}
LABELS = {1: "当月", 2: "次月", 3: "季月一", 4: "半年度"}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str,
                                    allow_nan=False) + "\n", encoding="utf-8")


def normalize_code(value):
    value = str(value).strip().upper()
    if re.fullmatch(r"(?:IF|IC|IM|IH)\d{2}(?:\d{2})?", value):
        value += ".CFE"
    return value


def dates(values, name):
    # Reject serial numbers here: read_excel converts genuinely typed Excel dates.
    if pd.api.types.is_numeric_dtype(values):
        raise ValueError(f"{name} 须为 Excel 日期或 YYYY-MM-DD，不能是未注明口径的数字")
    result = pd.to_datetime(values, errors="coerce").dt.normalize()
    if result.isna().any():
        raise ValueError(f"{name} 含空值或无法识别的日期")
    return result


def unique(frame, keys, label):
    if frame.duplicated(keys).any():
        example = frame.loc[frame.duplicated(keys, keep=False), keys].head(5).to_dict("records")
        raise ValueError(f"{label} 存在重复键：{example}")


def validate_panel(frame, roots, calendar=None):
    required = ["trade_date", "root", "slot", "instrument_id", "expiry_date",
                "futures_close", "spot_close", "open_interest"]
    missing = set(required) - set(frame)
    if missing:
        raise ValueError(f"行情缺少字段：{sorted(missing)}")
    p = frame[required].copy()
    p["trade_date"] = dates(p.trade_date, "行情日期")
    p["expiry_date"] = dates(p.expiry_date, "到期日")
    p["root"] = p.root.astype(str).str.upper().str.strip()
    p = p[p.root.isin(roots)].copy()
    if p.empty:
        raise ValueError("行情为空")
    for col in ["slot", "futures_close", "spot_close", "open_interest"]:
        p[col] = pd.to_numeric(p[col], errors="coerce")
        if not np.isfinite(p[col]).all():
            raise ValueError(f"行情 {col} 含缺失值或非有限值")
    if not p.slot.isin([1, 2, 3, 4]).all():
        raise ValueError("slot 必须为 1、2、3、4")
    p["slot"] = p.slot.astype(int)
    p["instrument_id"] = p.instrument_id.map(normalize_code)
    valid_codes = p.instrument_id.str.fullmatch(r"(?:IF|IC|IM|IH)\d{4}\.CFE")
    if not valid_codes.all() or not p.instrument_id.str[:2].eq(p.root).all():
        raise ValueError("具体合约代码与品种不匹配，不能以连续代码替代实际合约")
    if (p[["futures_close", "spot_close"]] <= 0).any().any():
        raise ValueError("期现货价格必须为正数")
    if (p.open_interest < 0).any() or not np.allclose(p.open_interest, p.open_interest.round()):
        raise ValueError("持仓量必须为非负整数，单位为手")
    unique(p, ["trade_date", "root", "slot"], "行情期限")
    unique(p, ["trade_date", "instrument_id"], "行情具体合约")
    if p.groupby(["trade_date", "root"]).spot_close.nunique().max() != 1:
        raise ValueError("同日同指数现货收盘值不一致")
    expected_days = pd.DatetimeIndex(calendar if calendar is not None else p.trade_date.unique()).normalize().sort_values()
    expected = pd.MultiIndex.from_product([expected_days, roots, [1, 2, 3, 4]],
                                         names=["trade_date", "root", "slot"])
    actual = pd.MultiIndex.from_frame(p[["trade_date", "root", "slot"]])
    if len(expected.difference(actual)) or len(actual.difference(expected)):
        raise ValueError(f"行情不完整：缺少 {len(expected.difference(actual))} 个日期品种期限，禁止静默缩短窗口")
    p = p.sort_values(["trade_date", "root", "slot"]).reset_index(drop=True)
    p["days_to_expiry"] = (p.expiry_date - p.trade_date).dt.days
    if p.days_to_expiry.lt(0).any():
        raise ValueError("交易日不能晚于合约到期日")
    if p.groupby("instrument_id").expiry_date.nunique().max() != 1:
        raise ValueError("同一具体合约出现不同到期日")
    expiry_diff = p.groupby(["trade_date", "root"]).days_to_expiry.diff()
    if expiry_diff.dropna().le(0).any():
        raise ValueError("00 至 03 的实际到期日未严格递增，请核对 Wind 连续映射")
    p["continuous_code"] = p.root + (p.slot - 1).astype(str).str.zfill(2) + ".CFE"
    return p


def read_dividends(config, base):
    """Read exact date × contract remaining dividends in index points, never yields."""
    if config.get("unit", "points") != "points":
        raise ValueError("分红必须是指数点数；分红率、现金金额不可直接加在期货价格上")
    if config.get("selection", "fixed_file") == "latest_snapshot":
        return read_latest_snapshot(config, base)
    path = Path(config["file"])
    if not path.is_absolute():
        path = Path(base) / path
    if not path.is_file():
        raise ValueError(f"未找到分红文件：{path}。请检查设置中的分红目录或 config.local.json")
    if config.get("layout") == "snapshot_continuous":
        return read_continuous_snapshot(path, config)
    if path.suffix.lower() == ".csv":
        raw = pd.read_csv(path, encoding="utf-8-sig")
    else:
        raw = pd.read_excel(path, sheet_name=config.get("sheet", 0),
                            header=int(config.get("header_row", 1)) - 1)
    raw.columns = [str(x).strip() for x in raw.columns]
    columns = config.get("columns", {})
    date_col = columns.get("trade_date", "日期")
    if config.get("layout", "long") == "wide":
        mapping = config.get("wide_columns", {})
        if not mapping or date_col not in raw:
            raise ValueError("宽表需配置日期列及 wide_columns（连续合约代码到分红列名）")
        parts = []
        for code, col in mapping.items():
            if col not in raw:
                raise ValueError(f"分红列不存在：{col}")
            parts.append(pd.DataFrame({"trade_date": raw[date_col], "code": code,
                                       "dividend_points": raw[col]}))
        d = pd.concat(parts, ignore_index=True)
    else:
        mapping = {"trade_date": date_col, "code": columns.get("code", "合约"),
                   "dividend_points": columns.get("dividend_points", "分红")}
        if set(mapping.values()) - set(raw):
            raise ValueError(f"分红表缺少配置列：{set(mapping.values()) - set(raw)}；实际列：{list(raw)}")
        d = raw[list(mapping.values())].rename(columns={v: k for k, v in mapping.items()})
    d = d.dropna(how="all").copy()
    if d.empty:
        raise ValueError("分红表没有数据")
    d["trade_date"] = dates(d.trade_date, "分红日期")
    d["code"] = d.code.map(normalize_code)
    if not d.code.str.fullmatch(r"(?:IF|IC|IM|IH)\d{2}(?:\d{2})?\.CFE").all():
        raise ValueError("分红合约须为 IF00.CFE 等连续代码或 IF2610.CFE 等具体合约")
    d["dividend_points"] = pd.to_numeric(d.dividend_points, errors="coerce")
    if not np.isfinite(d.dividend_points).all() or d.dividend_points.lt(0).any():
        raise ValueError("分红须为非负有限点数；空值不等于零，请检查 Excel 公式缓存")
    unique(d, ["trade_date", "code"], "分红")
    return d, {"path": str(path.resolve()), "sha256": sha256(path), "rows": len(d),
               "unit": "index_points", "semantics": config.get("semantics", "remaining_dividend_to_expiry"),
               "historical_type": config.get("historical_type", "user_supplied_unverified")}


def read_latest_snapshot(config, base):
    """Select by the workbook's source date, never by filename or filesystem mtime."""
    if config.get("layout") != "snapshot_continuous":
        raise ValueError("自动选择最新分红仅支持 snapshot_continuous；逐日历史请明确指定文件")
    folders = config.get("directories", [config.get("directory", "inputs")])
    directories = [Path(base) / folder for folder in folders]
    pattern = config.get("pattern", "*基差情况_*.xlsx")
    paths = []
    priority = {}
    for rank, directory in enumerate(directories):
        for path in sorted(directory.glob(pattern)):
            resolved = path.resolve()
            if path.is_file() and not path.name.startswith("~$") and resolved not in priority:
                paths.append(path)
                priority[resolved] = rank
    if not paths:
        raise ValueError(f"未找到分红表：{directories}。请在 inputs 放入 基差情况_日期.xlsx（也兼容指数基差情况_日期.xlsx）")
    candidates = []
    for path in paths:
        # Invalid candidate files are surfaced, never silently bypassed for an older file.
        data, info = read_continuous_snapshot(path, config)
        candidates.append((data, info))
    latest_date = max(info["snapshot_date"] for _, info in candidates)
    newest = [(data, info) for data, info in candidates if info["snapshot_date"] == latest_date]
    # On the same source date, the primary input folder supersedes legacy copies.
    best_priority = min(priority[Path(info["path"])] for _, info in newest)
    newest = [(data, info) for data, info in newest if priority[Path(info["path"])] == best_priority]
    canonical = [data.sort_values("code").reset_index(drop=True) for data, _ in newest]
    if any(not canonical[0].equals(data) for data in canonical[1:]):
        raise ValueError(f"{latest_date} 存在多份内容不同的最新分红表；请在同一分红目录中只保留一份该日表格")
    data, info = newest[-1]
    info["selection"] = "latest_snapshot_by_workbook_date"
    info["search_directories"] = [str(directory.resolve()) for directory in directories]
    info["same_date_directory_rule"] = "earlier directory wins; conflicting files within that directory fail"
    info["candidates"] = [{"path": i["path"], "snapshot_date": i["snapshot_date"], "sha256": i["sha256"]}
                          for _, i in candidates]
    return data, info


def read_continuous_snapshot(path, config):
    """A supplied static dividend assumption per 00..03 slot, not historical forecasts.

    Excel formulas are never executed and the original workbook is never saved.
    The report must disclose the source date and application across the whole window.
    """
    from openpyxl import load_workbook

    source_bytes = Path(path).read_bytes()
    workbook = load_workbook(io.BytesIO(source_bytes), read_only=True, data_only=True)
    sheet_name = config.get("sheet", "Sheet1")
    try:
        sheet = workbook[sheet_name] if isinstance(sheet_name, str) else workbook.worksheets[sheet_name]
        date_cell = config.get("snapshot_date_cell", "B1")
        snapshot_date = dates(pd.Series([sheet[date_cell].value]), "分红快照日期").iloc[0]
        code_col = config.get("code_column", "B")
        dividend_col = config.get("dividend_column", "G")
        header_row = int(config.get("header_row", 3))
        if sheet[f"{dividend_col}{header_row}"].value != "分红":
            raise ValueError(f"{sheet.title}!{dividend_col}{header_row} 不是“分红”，请核对配置，不能误用期间分红列")
        records, cells = [], {}
        for row in range(header_row + 1, sheet.max_row + 1):
            code = normalize_code(sheet[f"{code_col}{row}"].value)
            if not re.fullmatch(r"(?:IF|IC|IM|IH)0[0-3]\.CFE", code):
                continue
            value = sheet[f"{dividend_col}{row}"].value
            if isinstance(value, bool):
                raise ValueError(f"{sheet.title}!{dividend_col}{row} 不是有效分红点数")
            records.append({"trade_date": snapshot_date, "code": code, "dividend_points": value})
            cells[code] = f"{sheet.title}!{dividend_col}{row}"
        if not records:
            raise ValueError("未在快照中找到 00—03 连续档位及分红列")
        d = pd.DataFrame(records)
        unique(d, ["code"], "固定分红快照")
        d["dividend_points"] = pd.to_numeric(d.dividend_points, errors="coerce")
        if not np.isfinite(d.dividend_points).all() or d.dividend_points.lt(0).any():
            raise ValueError("快照分红须为非负有限点数，空白或公式缓存缺失不能当零")
        info = {"path": str(path.resolve()), "sha256": hashlib.sha256(source_bytes).hexdigest(), "rows": len(d),
                "unit": "index_points", "semantics": "fixed_remaining_dividend_by_continuous_slot",
                "historical_type": "fixed_snapshot_scenario", "application": "snapshot_continuous",
                "snapshot_date": str(snapshot_date.date()), "date_cell": f"{sheet.title}!{date_cell}",
                "cells": cells, "mapping_key": "continuous_code",
                "limitation": "Single-date static assumptions applied to each continuous slot across the entire history; not point-in-time historical dividend forecasts."}
        return d, info
    finally:
        workbook.close()


def attach_dividends(panel, dividends, plot_start, application="exact_date"):
    p = panel.copy()
    if application == "snapshot_continuous":
        unique(dividends, ["code"], "固定分红快照")
        if dividends.trade_date.nunique() != 1 or not dividends.code.str.fullmatch(r"(?:IF|IC|IM|IH)0[0-3]\.CFE").all():
            raise ValueError("固定分红模式要求单一快照日期和 00—03 连续代码")
        p["dividend_points"] = p.continuous_code.map(dividends.set_index("code").dividend_points)
        p["dividend_source_date"] = dividends.trade_date.iloc[0]
        p["dividend_application"] = "fixed_snapshot_scenario"
    elif application == "exact_date":
        lookup = dividends.set_index(["trade_date", "code"]).dividend_points
        concrete = lookup.reindex(pd.MultiIndex.from_arrays([p.trade_date, p.instrument_id])).to_numpy()
        continuous = lookup.reindex(pd.MultiIndex.from_arrays([p.trade_date, p.continuous_code])).to_numpy()
        both = np.isfinite(concrete) & np.isfinite(continuous)
        if not np.allclose(concrete[both], continuous[both], rtol=0, atol=1e-8):
            raise ValueError("同日具体合约与连续合约分红冲突")
        p["dividend_points"] = np.where(np.isfinite(concrete), concrete, continuous)
        p["dividend_source_date"] = p.trade_date.where(p.dividend_points.notna())
        p["dividend_application"] = "exact_date"
    else:
        raise ValueError(f"未知分红匹配方式：{application}")
    missing = p.trade_date.ge(plot_start) & p.dividend_points.isna()
    if missing.any():
        rows = p.loc[missing, ["trade_date", "continuous_code", "instrument_id"]]
        raise ValueError(f"最近三个月缺少 {len(rows)} 条分红，不能用零或当前分红回填历史。示例：{rows.head(6).to_dict('records')}")
    p["raw_basis_points"] = p.futures_close - p.spot_close
    p["adjusted_futures_close"] = p.futures_close + p.dividend_points
    p["basis_points"] = p.adjusted_futures_close - p.spot_close
    p["basis_rate"] = p.basis_points / p.spot_close
    p["annualized_basis_pct"] = (p.basis_rate * 36500 / p.days_to_expiry.replace(0, np.nan)).where(p.days_to_expiry.ge(5))
    # Display-only carry: hold the last valid annualized value until delivery.
    # Reset on each roll; never fill source data or feed display values into statistics.
    ordered = p.sort_values(["root", "slot", "trade_date"])
    segment = ordered.groupby(["root", "slot"]).instrument_id.transform(
        lambda s: s.ne(s.shift()).cumsum())
    held = ordered.annualized_basis_pct.groupby([ordered.root, ordered.slot, segment]).ffill()
    near_expiry = p.days_to_expiry.lt(5) & p.basis_points.notna() & p.annualized_basis_pct.isna()
    p["annualized_plot_pct"] = p.annualized_basis_pct.where(~near_expiry, held.reindex(p.index))
    p["annualized_plot_filled"] = near_expiry & p.annualized_plot_pct.notna()
    p["roll_day"] = p.groupby(["root", "slot"]).instrument_id.transform(lambda s: s.ne(s.shift()) & s.shift().notna())
    # Same concrete contract versus previous exchange trading day, including rolls.
    previous_dates = dict(zip(sorted(p.trade_date.unique())[1:], sorted(p.trade_date.unique())[:-1]))
    oi = p.set_index(["trade_date", "instrument_id"]).open_interest
    prev = oi.reindex(pd.MultiIndex.from_arrays([p.trade_date.map(previous_dates), p.instrument_id])).to_numpy()
    p["oi_change"] = p.open_interest - prev
    return p


def fixed_tenor(panel, days=60):
    records = []
    for (date, root), g in panel.groupby(["trade_date", "root"]):
        g = g[g.days_to_expiry.ge(5)].sort_values("days_to_expiry")
        exact = g[g.days_to_expiry.eq(days)]
        value = np.nan
        if not exact.empty:
            value = exact.iloc[0].basis_rate * 36500 / days
        else:
            low, high = g[g.days_to_expiry.lt(days)], g[g.days_to_expiry.gt(days)]
            if not low.empty and not high.empty:
                left, right = low.iloc[-1], high.iloc[0]
                weight = (right.days_to_expiry - days) / (right.days_to_expiry - left.days_to_expiry)
                value = (weight * left.basis_rate + (1 - weight) * right.basis_rate) * 36500 / days
        records.append({"trade_date": date, "root": root, "annualized_basis_pct": value})
    return pd.DataFrame(records)


def spreads(panel):
    rows = []
    for (date, root), g in panel.groupby(["trade_date", "root"]):
        g = g.set_index("slot")
        for slot in [2, 3, 4]:
            rows.append({"trade_date": date, "root": root, "spread": f"{root}0{slot-1}−{root}00",
                         "spread_points": g.loc[slot, "adjusted_futures_close"] - g.loc[1, "adjusted_futures_close"],
                         "raw_spread_points": g.loc[slot, "futures_close"] - g.loc[1, "futures_close"],
                         "dividend_difference": g.loc[slot, "dividend_points"] - g.loc[1, "dividend_points"]})
    return pd.DataFrame(rows)


def total_oi(panel):
    totals = panel.groupby(["trade_date", "root"], as_index=False).open_interest.sum()
    totals["oi_change"] = totals.groupby("root").open_interest.diff()
    return totals


def historical_summary(panel, end):
    records = []
    latest = panel[panel.trade_date.eq(end)]
    for row in latest.itertuples():
        for months in [3, 6, 12]:
            start = end - pd.DateOffset(months=months)
            history = panel[(panel.root == row.root) & (panel.slot == row.slot) &
                            panel.trade_date.ge(start) & panel.trade_date.lt(end)]
            eligible = history.days_to_expiry.ge(5)
            values = history.loc[eligible, "annualized_basis_pct"].dropna()
            # A 3-month data slice must never be labelled a full-year distribution.
            available_start = panel.trade_date.min()
            complete = available_start <= start and len(values) == int(eligible.sum()) and len(values) > 0
            current = row.annualized_basis_pct
            percentile = 100 * ((values < current).sum() + .5 * (values == current).sum()) / len(values) if complete and np.isfinite(current) else np.nan
            records.append({"root": row.root, "slot": row.slot, "months": months,
                            "current": current, "n": len(values), "complete": bool(complete),
                            "mean": values.mean() if complete else np.nan,
                            "median": values.median() if complete else np.nan, "percentile": percentile})
    return pd.DataFrame(records)
