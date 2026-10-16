"""Charts and Word report generated only from the validated adjusted panel."""
from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pandas as pd

from .analytics import LABELS, NAMES, ROOTS
from .term_structure import PAIRS, term_analysis, term_chart_catalog

COLORS = ["#28577B", "#288D87", "#C7813C", "#8D819E"]


def plot_all(panel, fixed, spread, start, end, output, roots, annual_quantiles=None,
             cross=None, cross_quantiles=None, whiskers="p05_p95", selected=None, watermark="", terms=None):
    os.environ.setdefault("MPLCONFIGDIR", str(Path(output) / ".mplconfig"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates
    from matplotlib import font_manager
    from matplotlib.ticker import MaxNLocator
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    from .distributions import annualized_distributions, cross_index_basis, cross_index_distributions

    if annual_quantiles is None:
        annual_quantiles = annualized_distributions(panel, end)
    if cross is None:
        cross = cross_index_basis(panel)
    if cross_quantiles is None:
        cross_quantiles = cross_index_distributions(cross, end)

    fontfile = Path("C:/Windows/Fonts/msyh.ttc")
    if fontfile.exists():
        font_manager.fontManager.addfont(str(fontfile))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(fontfile)).get_name()
    else:
        plt.rcParams["font.family"] = ["Noto Sans CJK SC", "DejaVu Sans"]
    plt.rcParams.update({"font.size": 13, "axes.unicode_minus": False, "svg.fonttype": "path"})
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    charts = {}
    selected = set(selected) if selected is not None else None
    if selected is not None and (not selected or selected - set(chart_catalog())):
        raise ValueError("请选择有效图表")
    def want(name):
        return selected is None or name in selected
    visible = panel[panel.trade_date.between(start, end)]

    def style(ax, window_start=None):
        ax.set_xlim(start if window_start is None else window_start, end)
        ax.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=4, maxticks=6))
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%m-%d"))
        ax.yaxis.set_major_locator(MaxNLocator(nbins=4))
        ax.grid(axis="y", color="#E1E7ED", linewidth=.7)
        ax.set_axisbelow(True)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.spines["bottom"].set_color("#CAD2DB")
        ax.tick_params(length=0, pad=6, labelcolor="#52606D")
        ax.margins(y=.14)

    def save(fig, name, include_in_doc=True):
        path = output / f"{name}.png"
        if watermark:
            fig.text(.99, .99, watermark, ha="right", va="top", color="#B5473C", fontsize=11,
                     bbox={"facecolor": "white", "alpha": .85, "edgecolor": "none"})
        fig.savefig(path, dpi=210, facecolor="white")
        plt.close(fig)
        if include_in_doc:
            charts[name] = path

    def quantile_panel(ax, data, title, color, unit, show_percentile=False):
        current = data.current.iloc[0]
        y_values = [current] if np.isfinite(current) else []
        for position, row in enumerate(data.sort_values("months").itertuples(), 1):
            if row.complete:
                lo, hi = (row.p05, row.p95) if whiskers == "p05_p95" else (row.minimum, row.maximum)
                y_values.extend([lo, hi])
                ax.bxp([{"q1": row.p25, "q3": row.p75, "med": row.p50,
                         "whislo": lo, "whishi": hi}], positions=[position], widths=.38,
                       showfliers=False, patch_artist=True, manage_ticks=False,
                       boxprops={"facecolor": color, "alpha": .35, "edgecolor": color},
                       medianprops={"color": "#263238", "linewidth": 1.8},
                       whiskerprops={"color": color, "linewidth": 1.6},
                       capprops={"color": color, "linewidth": 1.6})
            else:
                ax.text(position, .05, "覆盖不足", ha="center", fontsize=10,
                        color="#6F7780", transform=ax.get_xaxis_transform())
            if np.isfinite(current):
                ax.scatter(position, current, color="#B5473C", marker="D", s=44, zorder=5)
        ax.set_title(f"{title}  当前 {fmt(current)}{unit}", loc="left", color=color, fontsize=13)
        ax.set_xticks(list(range(1, len(data) + 1)), [f"{int(r.months)}个月\nn={int(r.n)}"
            + (f"\n当前{fmt(r.percentile, 1)}分位" if show_percentile else "")
            for r in data.sort_values("months").itertuples()])
        ax.set_xlim(.45, len(data) + .55)
        ax.set_ylabel(unit)
        ax.yaxis.set_major_locator(MaxNLocator(nbins=5))
        ax.grid(axis="y", color="#E1E7ED", linewidth=.7)
        ax.set_axisbelow(True)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.spines["bottom"].set_color("#CAD2DB")
        ax.tick_params(length=0, pad=6, labelcolor="#52606D")
        # Tight independent limits retain every whisker and the current point.
        # The same panel still uses one common scale for its 3/6/12-month candles.
        if y_values:
            low, high = min(y_values), max(y_values)
            span = high - low
            padding = span * .055 if span > 0 else max(abs(low) * .02, .1)
            ax.set_ylim(low - padding, high + padding)

    def quantile_figure(data, title, label, unit):
        fig, axes = plt.subplots(2, 2, figsize=(12.8, 9.8), layout="constrained")
        whisker_label = "P5—P95" if whiskers == "p05_p95" else "最小—最大"
        fig.suptitle(f"{title}  截至{end:%Y-%m-%d}", x=.02, ha="left", fontsize=15)
        for slot, ax in enumerate(axes.flat, 1):
            quantile_panel(ax, data[data.slot.eq(slot)], label(slot), COLORS[slot-1], unit)
        handles = [Patch(facecolor="#9EB6CB", alpha=.6, label="实体 P25—P75"),
                   Line2D([], [], color="#263238", label="横线 P50"),
                   Line2D([], [], color="#70889C", marker="_", label=f"上下须 {whisker_label}"),
                   Line2D([], [], color="#B5473C", marker="D", linestyle="None", label="当日值")]
        fig.legend(handles=handles, loc="outside lower center", ncol=4, frameon=False, fontsize=12)
        return fig

    for root in roots:
        data = visible[visible.root.eq(root)]
        for col, kind, unit in [("annualized_plot_pct", "年化基差", "%"), ("basis_points", "绝对基差", "点")]:
            display_note = "  临期缺值沿用前值" if col == "annualized_plot_pct" else ""
            if want(f"{root}_连续{kind}"):
                fig, axes = plt.subplots(2, 2, figsize=(12.8, 5.5), layout="constrained")
                fig.suptitle(f"分红调整后{kind}（{unit}）  各期限独立纵轴{display_note}", x=.02, ha="left", fontsize=14)
                for slot, ax in enumerate(axes.flat, 1):
                    g = data[data.slot.eq(slot)]
                    ax.plot(g.trade_date, g[col], color=COLORS[slot-1], linewidth=1.9)
                    ax.set_title(f"{root}0{slot-1}  {LABELS[slot]}", loc="left", color=COLORS[slot-1], fontsize=14)
                    style(ax)
                save(fig, f"{root}_连续{kind}")
            if want(f"{root}_合并{kind}"):
                fig, ax = plt.subplots(figsize=(12.8, 3.5), layout="constrained")
                for slot in [1, 2, 3, 4]:
                    g = data[data.slot.eq(slot)]
                    ax.plot(g.trade_date, g[col], color=COLORS[slot-1], linewidth=1.9,
                            label=f"{root}0{slot-1} {LABELS[slot]}")
                ax.set_title(f"分红调整后{kind}（{unit}）{display_note}", loc="left", fontsize=14, pad=36)
                ax.legend(loc="lower right", bbox_to_anchor=(1, 1.01), ncol=4, frameon=False, fontsize=12)
                style(ax)
                save(fig, f"{root}_合并{kind}")

        if want(f"{root}_年化基差分位蜡烛图"):
            quantile_fig = quantile_figure(annual_quantiles[annual_quantiles.root.eq(root)],
                "分红调整后年化基差历史分位", lambda slot: f"{root}0{slot-1} {LABELS[slot]}", "%")
            save(quantile_fig, f"{root}_年化基差分位蜡烛图")

        if want(f"{root}_基差与持仓变化"):
            fig, axes = plt.subplots(2, 2, figsize=(12.8, 6.2), layout="constrained")
            fig.suptitle("分红调整后绝对基差与持仓日变化  虚线为该期限换月日", fontsize=14, x=.02, ha="left")
            for slot, ax in enumerate(axes.flat, 1):
                g = data[data.slot.eq(slot)]
                right = ax.twinx()
                right.bar(g.trade_date, g.oi_change, width=1.5, color="#9DAEC0", alpha=.6, label="持仓日增减（右）")
                right.axhline(0, color="#C4CDD7", linewidth=.7)
                right.set_ylabel("日增减（手）", fontsize=11, color="#687C93")
                right.spines[["top", "left"]].set_visible(False)
                right.spines["right"].set_color("#D5DCE4")
                right.yaxis.set_major_locator(MaxNLocator(nbins=4))
                right.tick_params(length=0, labelsize=10, labelcolor="#687C93")
                ax.set_zorder(right.get_zorder() + 1)
                ax.patch.set_visible(False)
                ax.plot(g.trade_date, g.basis_points, color=COLORS[slot-1], linewidth=2, label="绝对基差（左）")
                for day in g.loc[g.roll_day, "trade_date"]:
                    ax.axvline(day, color="#ADB8C4", linestyle="--", linewidth=.7)
                ax.set_title(f"{root}0{slot-1}  {LABELS[slot]}", loc="left", color=COLORS[slot-1], fontsize=14)
                ax.set_ylabel("基差（点）", fontsize=11)
                style(ax)
            save(fig, f"{root}_基差与持仓变化")

        if want(f"{root}_跨期价差"):
            g = spread[spread.root.eq(root) & spread.trade_date.between(start, end)]
            fig, ax = plt.subplots(figsize=(12.8, 3.7), layout="constrained")
            for i, (label, curve) in enumerate(g.groupby("spread", sort=True)):
                ax.plot(curve.trade_date, curve.spread_points, linewidth=2, color=COLORS[i+1], label=label)
            ax.set_title("分红调整后跨期价差（点）", loc="left", fontsize=14, pad=35)
            ax.legend(loc="lower right", bbox_to_anchor=(1, 1.01), ncol=3, frameon=False)
            ax.axhline(0, color="#A3ADB9", linewidth=.8)
            for day in data.loc[data.roll_day, "trade_date"].drop_duplicates():
                ax.axvline(day, color="#ADB8C4", linestyle="--", linewidth=.7)
            style(ax)
            save(fig, f"{root}_跨期价差")

    if want("60天固定期限比较"):
        fig, ax = plt.subplots(figsize=(12.8, 4.0), layout="constrained")
        for i, root in enumerate(roots):
            g = fixed[fixed.root.eq(root) & fixed.trade_date.between(start, end)]
            ax.plot(g.trade_date, g.annualized_basis_pct, color=COLORS[i], linewidth=2, label=root)
        ax.set_title("分红调整后60天固定期限年化基差（%）", loc="left", fontsize=14, pad=35)
        ax.legend(loc="lower right", bbox_to_anchor=(1, 1.01), ncol=len(roots), frameon=False)
        style(ax)
        save(fig, "60天固定期限比较")
    if not cross.empty:
        visible_cross = cross[cross.trade_date.between(start, end)]

        def draw_cross(ax, slot):
            g = visible_cross[visible_cross.slot.eq(slot)]
            ax.plot(g.trade_date, g.basis_difference_points, color=COLORS[slot-1], linewidth=2)
            ax.axhline(0, color="#A3ADB9", linewidth=.8)
            for day in g.loc[g.roll_day, "trade_date"]:
                ax.axvline(day, color="#ADB8C4", linestyle="--", linewidth=.7)
            ax.set_title(f"IM0{slot-1}−IC0{slot-1}  {LABELS[slot]}", loc="left", color=COLORS[slot-1], fontsize=14)
            ax.set_ylabel("基差差（点）")
            style(ax)

        if want("IM_IC_绝对基差差分图"):
            fig, axes = plt.subplots(2, 2, figsize=(12.8, 5.8), layout="constrained")
            fig.suptitle("同到期日 IM−IC 分红调整后绝对基差差  各期限独立纵轴", x=.02, ha="left", fontsize=14)
            for slot, ax in enumerate(axes.flat, 1):
                draw_cross(ax, slot)
            save(fig, "IM_IC_绝对基差差分图")
        for slot in [1, 2, 3, 4]:
            if want(f"IM0{slot-1}_IC0{slot-1}_绝对基差差"):
                fig, ax = plt.subplots(figsize=(12.8, 4.0), layout="constrained")
                draw_cross(ax, slot)
                save(fig, f"IM0{slot-1}_IC0{slot-1}_绝对基差差", include_in_doc=False)
        if want("IM_IC_绝对基差差合并图"):
            fig, ax = plt.subplots(figsize=(12.8, 3.8), layout="constrained")
            for slot, g in visible_cross.groupby("slot"):
                ax.plot(g.trade_date, g.basis_difference_points, color=COLORS[slot-1], linewidth=2,
                        label=f"IM0{slot-1}−IC0{slot-1}")
            ax.axhline(0, color="#A3ADB9", linewidth=.8)
            ax.set_title("同到期日 IM−IC 分红调整后绝对基差差（点）", loc="left", fontsize=14, pad=36)
            ax.legend(loc="lower right", bbox_to_anchor=(1, 1.01), ncol=4, frameon=False, fontsize=12)
            style(ax)
            save(fig, "IM_IC_绝对基差差合并图")
        if want("IM_IC_绝对基差差分位蜡烛图"):
            fig = quantile_figure(cross_quantiles, "IM−IC 绝对基差差历史分位",
                                 lambda slot: f"IM0{slot-1}−IC0{slot-1} {LABELS[slot]}", "点")
            save(fig, "IM_IC_绝对基差差分位蜡烛图")
    if selected is None or selected.intersection(term_chart_catalog()):
        fixed90, fixed90_quantiles, calendar, calendar_quantiles = terms if terms is not None else term_analysis(panel, end)

        def term_lines(series, name, title, months, roll=False):
            left = end - pd.DateOffset(months=months)
            fig, ax = plt.subplots(figsize=(12.8, 4.1), layout="constrained")
            for i, root in enumerate(roots):
                g = series[series.root.eq(root) & series.trade_date.between(left, end)].sort_values("trade_date")
                current = g.loc[g.trade_date.eq(end), "annualized_basis_pct"]
                current = current.iloc[0] if len(current) else np.nan
                ax.plot(g.trade_date, g.annualized_basis_pct, linewidth=2, color=COLORS[i],
                        label=f"{root} 当前{fmt(current)}%")
                if np.isfinite(current):
                    ax.scatter([end], [current], color=COLORS[i], s=22, zorder=4)
                if roll:
                    for day in g.loc[g.roll_day, "trade_date"]:
                        ax.axvline(day, color=COLORS[i], linestyle="--", alpha=.25, linewidth=.7)
            ax.set_title(f"{title}  近{months}个自然月（{left:%Y-%m-%d}至{end:%Y-%m-%d}）",
                         loc="left", fontsize=14, pad=35)
            ax.legend(loc="lower right", bbox_to_anchor=(1, 1.01), ncol=len(roots), frameon=False, fontsize=12)
            ax.set_ylabel("年化基差（%）")
            ax.axhline(0, color="#A3ADB9", linewidth=.8)
            style(ax, left)
            save(fig, name)

        def term_quantiles(data, name, title):
            grid_rows, grid_cols = (2, 2) if len(roots) == 4 else (1, len(roots))
            fig, axes = plt.subplots(grid_rows, grid_cols, figsize=(12.8, 9.0 if grid_rows == 2 else 5.8),
                                     squeeze=False, layout="constrained")
            fig.suptitle(f"{title}  截至{end:%Y-%m-%d}", x=.02, ha="left", fontsize=15)
            for i, (root, ax) in enumerate(zip(roots, axes.flat)):
                quantile_panel(ax, data[data.root.eq(root)], root, COLORS[i], "%", show_percentile=True)
            fig.legend(handles=[Patch(facecolor="#9EB6CB", alpha=.6, label="实体 P25—P75"),
                Line2D([], [], color="#263238", label="横线 P50"),
                Line2D([], [], color="#70889C", marker="_",
                       label="上下须 P5—P95" if whiskers == "p05_p95" else "上下须 最小—最大"),
                Line2D([], [], color="#B5473C", marker="D", linestyle="None", label="当日值")],
                loc="outside lower center", ncol=4, frameon=False, fontsize=12)
            save(fig, name)

        for days, series in [(60, fixed), (90, fixed90)]:
            for months in (3, 1):
                name = f"{days}天固定期限_近{months}个月"
                if want(name):
                    term_lines(series, name, f"分红调整后{days}天固定期限年化基差", months)
        if want("90天固定期限_分位蜡烛图"):
            term_quantiles(fixed90_quantiles, "90天固定期限_分位蜡烛图", "90天固定期限年化基差历史分位")
        for pair in PAIRS:
            if want(f"跨期{pair}_年化基差"):
                term_lines(calendar[calendar.pair.eq(pair)], f"跨期{pair}_年化基差",
                           f"{pair} 跨期年化基差（现货分母）", 3, roll=True)
            if want(f"跨期{pair}_年化分位蜡烛图"):
                term_quantiles(calendar_quantiles[calendar_quantiles.pair.eq(pair)],
                               f"跨期{pair}_年化分位蜡烛图", f"{pair} 跨期年化基差历史分位")
    return charts


def fmt(value, decimals=2, signed=False):
    if pd.isna(value):
        return "—"
    return format(value, ("+" if signed else "") + f",.{decimals}f").replace("-", "−")


def delta(series, periods):
    return series.iloc[-1] - series.iloc[-periods-1] if len(series) > periods else np.nan


def root_analysis(root, panel, fixed, spread, totals, start, end):
    p = panel[(panel.root == root) & panel.trade_date.between(start, end)]
    f = fixed[(fixed.root == root) & fixed.trade_date.between(start, end)].sort_values("trade_date")
    oi = totals[(totals.root == root) & totals.trade_date.between(start, end)].sort_values("trade_date")
    latest = p[p.trade_date.eq(end)].sort_values("slot")
    f0, f1 = f.annualized_basis_pct.iloc[0], f.annualized_basis_pct.iloc[-1]
    change5 = delta(f.annualized_basis_pct, 5)
    trend = "上升" if f1 >= f0 else "下降"
    lead = (f"{NAMES[root]} {root}。60天固定期限年化基差由窗口首日的{fmt(f0)}%变为{fmt(f1)}%，"
            f"累计{trend}{fmt(abs(f1-f0))}个百分点，近5个交易日变动{fmt(change5, signed=True)}个百分点。"
            f"当月连续的最新绝对基差为{fmt(latest.iloc[0].basis_points)}点；其分红调整幅度为"
            f"{fmt(latest.iloc[0].dividend_points)}点。")
    maxrow = latest.loc[latest.open_interest.idxmax()]
    last_oi = oi.open_interest.iloc[-1]
    oi5 = delta(oi.open_interest, 5)
    holdings = (f"品种总持仓为{fmt(last_oi/10000)}万手，近5个交易日变动{fmt(oi5, 0, True)}手。"
                f"持仓最多的是{maxrow.instrument_id.replace('.CFE','')}，占比{fmt(maxrow.open_interest/last_oi*100) if last_oi else '—'}%。"
                "各期限柱形为同一具体合约较上一交易日的持仓变化；换月日依据新合约自身前值计算，"
                "新挂牌合约无前值时留空。品种总持仓变化另按四合约合计计算。")
    s = spread[(spread.root == root) & spread.trade_date.between(start, end)]
    spread_text = []
    for label, g in s.groupby("spread"):
        g = g.sort_values("trade_date")
        spread_text.append(f"{label}为{fmt(g.spread_points.iloc[-1])}点，较窗口首日变动{fmt(g.spread_points.iloc[-1]-g.spread_points.iloc[0], signed=True)}点")
    curve = "；".join(spread_text) + "。负值表示相应远期合约在加回各自分红后仍低于当月。换月会改变两端合约，连续价差的跨月变化不等于固定两份合约的盈亏。"
    if np.isfinite(change5) and np.isfinite(oi5):
        coexist = f"近5日观察到总持仓{'增加' if oi5 >= 0 else '减少'}与60天基差{'上升' if change5 >= 0 else '下降'}并存。"
    else:
        coexist = "近期可比样本不足，暂不计算五日联动。"
    interpretation = coexist + "该组合只描述同步变化，不能据此区分多头需求、空头套保或移仓动机。后续可结合固定期限基差与跨期价差是否同向变化，判断变化集中于近月还是延伸至远月。"
    return lead, holdings, curve, interpretation


def quantile_guide(provenance, end, annualized):
    whiskers = provenance.get("configuration", {}).get("quantile_whiskers", "p05_p95")
    ends = "第5和第95百分位" if whiskers == "p05_p95" else "最小值和最大值"
    text = (f"每根蜡烛汇总{end:%Y年%m月%d日}之前向前3、6、12个自然月的日度样本，不包含截止日。"
            f"实体上下缘为P25与P75，中间横线为中位数P50，上下须为{ends}，红色菱形为当日值；"
            "n为有效样本数。蜡烛描述历史分布，颜色不表示开盘收盘涨跌。")
    if annualized:
        text += "年化样本剔除到期不足5天的日期，临期横线填充值不进入分布或当前红点。"
    else:
        text += "点数基差差保留临近交割的实际观察，不作年化，也不使用横线填充值。"
    if provenance.get("dividends", {}).get("historical_type") == "fixed_snapshot_scenario":
        text += "各窗口使用本次选定的同一组固定分红，属于固定假设下的历史比较。"
    return text


def quantile_description(label, data, unit):
    data = data.sort_values("months")
    current = data.current.iloc[0]
    positions = [f"{fmt(row.percentile, 1)}分位" if np.isfinite(row.percentile)
                 else ("覆盖不足" if not row.complete else "当日值缺失") for row in data.itertuples()]
    if not np.isfinite(current):
        return f"{label}当前值留空；3、6、12个月的历史分布仍按各自有效样本展示。"
    return (f"{label}当前{fmt(current)}{unit}；3、6、12个月依次为"
            + "、".join(positions) + "。")


def build_docx(template, target, panel, fixed, spread, totals, history, charts,
               start, end, roots, provenance, preview=False, annual_quantiles=None,
               cross=None, cross_quantiles=None, terms=None):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from .distributions import annualized_distributions, cross_index_basis, cross_index_distributions

    if annual_quantiles is None:
        annual_quantiles = annualized_distributions(panel, end)
    if cross is None:
        cross = cross_index_basis(panel)
    if cross_quantiles is None:
        cross_quantiles = cross_index_distributions(cross, end)

    doc = Document(template)
    for element in list(doc.element.body):
        if element.tag != qn("w:sectPr"):
            doc.element.body.remove(element)
    for rel_id, relation in list(doc.part.rels.items()):
        if relation.reltype.endswith("/image"):
            doc.part.drop_rel(rel_id)
    for name in ["Title", "Heading 1", "Heading 2", "Caption"]:
        s = doc.styles[name]
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.keep_with_next = True
    section = doc.sections[0]
    width = (section.page_width - section.left_margin - section.right_margin) / 914400
    figure_number, table_number = 0, 0

    def p(text, style=None):
        return doc.add_paragraph(text, style)

    def heading(text, page=False):
        para = p(text, "Heading 1")
        para.paragraph_format.page_break_before = page
        return para

    def picture(key, caption):
        nonlocal figure_number
        figure_number += 1
        p(f"图{figure_number}  {caption}", "Caption")
        para = p("")
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.line_spacing = 1
        para.paragraph_format.space_after = Pt(6)
        shape = para.add_run().add_picture(str(charts[key]), width=Inches(width))
        shape._inline.docPr.set("descr", caption)

    def table(caption, headers, rows, widths=None, font_size=None):
        nonlocal table_number
        table_number += 1
        p(f"表{table_number}  {caption}", "Caption")
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = widths or [width / len(headers)] * len(headers)
        for col, w in zip(t.columns, widths):
            col.width = Inches(w)
        for cells, row in [(t.rows[0].cells, headers)] + [(t.add_row().cells, row) for row in rows]:
            for i, (cell, value) in enumerate(zip(cells, row)):
                cell.width = Inches(widths[i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = str(value)
                for para in cell.paragraphs:
                    para.style = "TableText" if "TableText" in doc.styles else "Normal"
                    para.paragraph_format.first_line_indent = Pt(0)
                    para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.space_before = Pt(3)
                    para.paragraph_format.keep_with_next = False
                    if font_size is not None:
                        para.paragraph_format.line_spacing = Pt(font_size + 3)
                        para.paragraph_format.space_before = Pt(1)
                        para.paragraph_format.space_after = Pt(1)
                        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                        for run in para.runs:
                            run.font.size = Pt(font_size)
        for row_idx, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if row_idx == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                tcpr = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EDF2" if row_idx == 0 else "FFFFFF")
                tcpr.append(shade)
                borders = OxmlElement("w:tcBorders")
                for edge in ["top", "bottom", "left", "right"]:
                    border = OxmlElement("w:" + edge)
                    for attr, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                        border.set(qn("w:" + attr), value)
                    borders.append(border)
                tcpr.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ["top", "bottom", "left", "right"]:
                    element = OxmlElement("w:" + edge)
                    element.set(qn("w:w"), "65")
                    element.set(qn("w:type"), "dxa")
                    margins.append(element)
                tcpr.append(margins)
                if row_idx == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        spacer = p("")
        spacer.paragraph_format.space_after = Pt(0)
        if font_size is not None:
            spacer.paragraph_format.line_spacing = Pt(3)

    p("股指基差计算明细" + (" 合成演示" if preview else ""), "Title")
    def front_note(text):
        para = p(text)
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.line_spacing = Pt(15)
        para.paragraph_format.space_after = Pt(5)
        for run in para.runs:
            run.font.size = Pt(10.5)
        return para

    if preview:
        front_note("本报告全部使用合成行情及合成分红，仅用于演示应用功能，不是真实市场研究结果。")
    front_note(f"行情日期：{end:%Y-%m-%d}。报告生成时间：{provenance.get('generated_at', '未记录')}。"
      "下表均为该行情日的实际合约收盘数据。请先核对行情日期，再比较年化结果。")
    requested_end = provenance.get("requested_end")
    update_check = provenance.get("market_update_check", {})
    if update_check.get("state") == "fallback":
        reason = update_check.get("reason", "")
        reason = reason[:90] + "…（完整原因见导出清单）" if len(reason) > 90 else reason
        note = front_note(f"更新未成功：目标截止{requested_end}，已回退至前一交易日{end:%Y-%m-%d}。原因：{reason}")
        for run in note.runs:
            run.bold = True
    elif update_check.get("state") == "complete":
        front_note(f"生成前已补齐并校验行情，最新适用交易日为{end:%Y-%m-%d}。"
                   f"按北京时间{update_check['daily_check_time']}纳入当日；交易日历检查截至{requested_end}。")
    elif requested_end and pd.Timestamp(requested_end).normalize() > pd.Timestamp(end).normalize():
        front_note(f"本次完整日线截止目标为{requested_end}，本报告实际行情截至{end:%Y-%m-%d}。"
          "如需更新，在数据管理页检查并补齐行情后重新导出。")
    dividend_info = provenance.get("dividends", {})
    if dividend_info.get("historical_type") == "fixed_snapshot_scenario":
        front_note(f"分红快照日期：{dividend_info.get('snapshot_date', '未记录')}。"
          "采用表格G列固定点数，按00—03档位匹配；不是行情日期，也不使用J列期间分红。")
    front_note("调整基差B＝期货收盘价F＋分红G−现货收盘值S；年化基差＝B÷S×365÷剩余自然日D×100%。"
      "D＝实际到期日−行情日期；D不足5天时年化留空，记为—。价格和基差单位均为指数点。")
    for root_index, root in enumerate(roots):
        if len(roots) > 3 and root_index == 2:
            heading("股指基差计算明细 续", page=True)
            front_note(f"行情日期：{end:%Y-%m-%d}。分红快照日期：{dividend_info.get('snapshot_date', '未记录')}。"
                       "年化基差＝（期货收盘价＋G列分红−现货收盘值）÷现货收盘值×365÷剩余自然日×100%。")
        latest = panel[panel.root.eq(root) & panel.trade_date.eq(end)].sort_values("slot")
        rows = [[f"{r.continuous_code.replace('.CFE', '')}/{r.instrument_id.replace('.CFE', '')}",
                 fmt(r.futures_close, 2), f"{r.expiry_date:%Y-%m-%d}", str(r.days_to_expiry),
                 fmt(r.dividend_points, 5), fmt(r.basis_points, 4), fmt(r.annualized_basis_pct, 3)]
                for r in latest.itertuples()]
        proportions = [2.3, 2.1, 2.4, 1.0, 2.5, 2.5, 1.8]
        table(f"{NAMES[root]} {root} 现货收盘值 {fmt(latest.spot_close.iloc[0], 4)} 点",
              ["连续档位\n实际合约", "期货收盘\nF", "到期日", "天数\nD", "分红\nG", "调整基差\nB", "年化\n%"],
              rows, [width * part / sum(proportions) for part in proportions], font_size=9)
    front_note("表内数字按列显示精度四舍五入，计算使用原始精度。到期日取实际合约行情源；"
      "本报告生成日期不代表行情已更新到当天。")
    heading("晨会分析", page=True)
    actual_start = panel.loc[panel.trade_date.ge(start), "trade_date"].min()
    current = fixed[fixed.trade_date.eq(end)].set_index("root").annualized_basis_pct
    overview = "、".join(f"{r}为{fmt(current.get(r))}%" for r in roots)
    p(f"本期覆盖{'、'.join(NAMES[r] for r in roots)}，行情截至{end:%Y年%m月%d日}。"
      f"时序图展示{start:%Y年%m月%d日}至{end:%Y年%m月%d日}最近三个日历月，分位图比较近3、6、12个月，"
      f"首个实际交易日为{actual_start:%m月%d日}。分红调整后的60天固定期限年化基差，{overview}。")
    kind = provenance.get("dividends", {}).get("historical_type", "user_supplied_unverified")
    if kind == "fixed_snapshot_scenario":
        snapshot = provenance["dividends"]["snapshot_date"]
        p(f"分红统一采用{snapshot}所提供表格G列的固定点数，按各品种00—03连续档位应用于整个窗口。"
          "这是一组固定分红假设下的调整结果，并非历史逐日可得的分红预测；跨期比较的两端均使用各自档位分红。")
    p("先将每个交易日、每个合约对应的剩余分红点数加到期货收盘价，再计算基差、年化值与跨期价差。"
      "绝对基差＝调整后期货价−现货指数收盘值，负值为贴水，正值为升水；年化基差＝绝对基差"
      "÷现货指数收盘值×365÷剩余自然日数×100%。跨期价差取远月调整价减当月调整价。当前未另扣融资成本。")
    picture("60天固定期限比较", "各品种近三个月60天固定期限年化基差")
    rows = []
    for root in roots:
        g = fixed[fixed.root.eq(root) & fixed.trade_date.between(start, end)].sort_values("trade_date")
        v = g.annualized_basis_pct
        rows.append([root, fmt(v.iloc[-1]), fmt(delta(v, 1), signed=True), fmt(delta(v, 5), signed=True),
                     fmt(delta(v, 20), signed=True), fmt(v.iloc[-1] - v.iloc[0], signed=True)])
    table("60天固定期限年化基差及变化", ["品种", "当前\n%", "较1日前\n百分点", "较5日前\n百分点", "较20日前\n百分点", "较窗口首日\n百分点"], rows)
    p("固定期限先在相邻合约间插值分红调整后的基差率，再按60天年化；不足以包围60天的日期留空。"
      "连续合约到期不足5个自然日时，年化图沿用本合约最近有效值画水平线，换月后恢复新合约数据；"
      "填充值仅用于作图，不进入表格或历史统计。")
    if kind == "point_in_time_forecast":
        p("分红输入标记为历史时点预计剩余分红，按交易日和实际合约或连续档位精确匹配。")
    elif kind == "realized":
        p("分红输入为事后实际分红，调整序列用于事后解释；历史分位不代表当时可观察的估值信号。")
    elif preview:
        p("本次使用合成演示数据；正式报告必须使用实际行情与分红输入。")
    elif kind == "fixed_snapshot_scenario":
        pass  # The source and assumption are prominent above the first chart.
    else:
        p("分红采用所提供表格的历史列值。其是否属于当时可得预测尚未核实，相关历史分位只作该输入口径下的描述。")

    for root in roots:
        lead, holdings, curve, interpretation = root_analysis(root, panel, fixed, spread, totals, start, end)
        heading(f"{NAMES[root]} {root} 连续合约年化基差", page=True)
        p(lead)
        picture(f"{root}_连续年化基差", f"{root}各期限年化基差分图")
        picture(f"{root}_合并年化基差", f"{root}各期限年化基差合并比较")
        p("00、01、02、03依次对应当月、次月、季月一和半年度连续档位。半年度沿用原报告称谓，"
          "实际是随后第二个季月档位，剩余期限并不固定为半年。分图纵轴独立，合并图纵轴统一。")
        heading(f"{NAMES[root]} {root} 年化基差分位", page=True)
        picture(f"{root}_年化基差分位蜡烛图", f"{root}各连续合约近3个月6个月12个月分位蜡烛图")
        p(quantile_guide(provenance, end, annualized=True))
        for slot in [1, 2, 3, 4]:
            g = annual_quantiles[annual_quantiles.root.eq(root) & annual_quantiles.slot.eq(slot)]
            p(quantile_description(f"{root}0{slot-1}", g, "%"))
        p("实体越高表示年化基差整体越高，负值区间内对应贴水较浅；实体或上下须越长表示历史分布越分散。"
          "红点高于P95表示高于约95%的历史有效样本，并不意味着未来有95%的概率回落。"
          "不同连续档位的剩余期限会随换月变化，比较时仍需结合60天固定期限图。")
        heading(f"{NAMES[root]} {root} 绝对基差", page=True)
        picture(f"{root}_连续绝对基差", f"{root}各期限绝对基差分图")
        picture(f"{root}_合并绝对基差", f"{root}各期限绝对基差合并比较")
        latest = panel[panel.root.eq(root) & panel.trade_date.eq(end)].sort_values("slot")
        rows = [[r.continuous_code.replace(".CFE", ""), r.instrument_id.replace(".CFE", ""),
                 str(r.days_to_expiry), fmt(r.dividend_points), fmt(r.basis_points), fmt(r.annualized_basis_pct)] for r in latest.itertuples()]
        table(f"{root}最新期限与分红调整结果", ["连续档位", "实际合约", "剩余天数", "分红\n点", "绝对基差\n点", "年化基差\n%"], rows)
        heading(f"{NAMES[root]} {root} 持仓与跨期价差", page=True)
        picture(f"{root}_基差与持仓变化", f"{root}绝对基差与持仓日变化双坐标轴")
        p(holdings)
        picture(f"{root}_跨期价差", f"{root}01 {root}02 {root}03分别减{root}00")
        p(curve)
        p(interpretation)

    if not cross.empty:
        heading("同期限 IM与IC 绝对基差比较", page=True)
        picture("IM_IC_绝对基差差分图", "IM减IC同到期日绝对基差差分图")
        picture("IM_IC_绝对基差差合并图", "IM减IC四期限绝对基差差合并比较")
        p("图中数值＝（IM期货价＋IM分红−中证1000现货）−（IC期货价＋IC分红−中证500现货）。"
          "每个交易日核验两边实际到期日相同；分图便于观察各档变化，合并图便于比较期限之间的差异。")
        recent = cross[cross.trade_date.between(start, end)]
        parts = []
        for slot, g in recent.groupby("slot"):
            g = g.sort_values("trade_date")
            parts.append(f"{slot-1:02d}档当前{fmt(g.basis_difference_points.iloc[-1], signed=True)}点，"
                         f"较窗口首日{fmt(g.basis_difference_points.iloc[-1]-g.basis_difference_points.iloc[0], signed=True)}点")
        p("；".join(parts) + "。")
        p("正值表示IM的点数基差高于IC，负值表示IM的点数基差低于IC；上行表示IM基差相对IC抬升。"
          "两指数点位尺度不同，该差值不是百分比相对估值，也不是中性组合收益或直接套利利润。")
        heading("IM与IC 基差差历史分位", page=True)
        picture("IM_IC_绝对基差差分位蜡烛图", "同期限IM减IC绝对基差差近3个月6个月12个月分位")
        p(quantile_guide(provenance, end, annualized=False))
        for slot in [1, 2, 3, 4]:
            p(quantile_description(f"IM0{slot-1}−IC0{slot-1}", cross_quantiles[cross_quantiles.slot.eq(slot)], "点"))
        p("红点接近上须表示IM相对IC的点数基差差处于历史较高区域，接近下须则处于较低区域。"
          "不同窗口位置不一致，说明短期和长期参照分布不同；仅凭分位高低不能判断后续必然收敛。")

    heading("最新持仓与历史比较", page=True)
    rows = []
    for root in roots:
        g = totals[totals.root.eq(root) & totals.trade_date.between(start, end)].sort_values("trade_date")
        v = g.open_interest
        rows.append([root, fmt(v.iloc[-1]/10000), fmt(delta(v, 1), 0, True), fmt(delta(v, 5), 0, True),
                     fmt((v.iloc[-1]/v.iloc[0]-1)*100, signed=True) if v.iloc[0] else "—"])
    summary_font = 10 if len(roots) > 3 else None
    table("品种总持仓变化", ["品种", "持仓\n万手", "较1日前\n手", "较5日前\n手", "较窗口首日\n%"], rows,
          font_size=summary_font)
    rows = []
    for root in roots:
        for slot in [1, 2, 3, 4]:
            g = history[(history.root == root) & (history.slot == slot)].set_index("months")
            rows.append([f"{root}0{slot-1}", fmt(g.loc[3, "current"]), fmt(g.loc[3, "mean"]),
                         fmt(g.loc[3, "percentile"], 1), fmt(g.loc[6, "mean"]), fmt(g.loc[12, "mean"])])
    table("固定分红假设下的历史水平" if kind == "fixed_snapshot_scenario" else "同品种同期限档位的调整后历史水平",
          ["连续档位", "当前\n%", "三月均值\n%", "三月分位\n%", "半年均值\n%", "一年均值\n%"], rows,
          font_size=summary_font)
    p("历史比较不包含当日，且剔除剩余期限不足5天的年化样本；分位采用并列值半权重。"
      "窗口内分红或行情覆盖不足时显示横线，不以未调整数据或短窗口冒充完整历史分布。"
      "不同档位仍受剩余期限变化影响，宜结合60天固定期限比较。")
    if kind == "fixed_snapshot_scenario":
        p("表中历史值均加回同一组固定档位分红，用于该假设下的横向比较，不代表当时可观察的估值分位。")

    fixed90, fixed90_quantiles, calendar, calendar_quantiles = terms if terms is not None else term_analysis(panel, end)
    for days, series in [(60, fixed), (90, fixed90)]:
        heading(f"{days}日固定期限年化基差", page=True)
        p(f"对{'、'.join(roots)}分别使用当日00—03四档实际合约，在包围{days}日的相邻到期日之间，"
          f"先线性插值分红调整后的基差率，再乘365/{days}年化；每个品种得到一条固定期限曲线。"
          "剩余不足5日的合约不参与插值，无法包围目标期限时留空，不作外推。")
        for months in (3, 1):
            picture(f"{days}天固定期限_近{months}个月", f"{days}日固定期限年化基差近{months}个月")
        current = series[series.trade_date.eq(end)].set_index("root").annualized_basis_pct
        p("当前水平：" + "；".join(f"{root} {fmt(current.get(root))}%" for root in roots) + "。"
          "三个月图用于观察阶段变化，一个月图突出近期变化；两图使用相同算法和同一分红快照。")

    def term_table(caption, quantiles):
        windows = sorted(quantiles.months.unique())
        headers = ["品种", "当前\n%"]
        for months in windows:
            headers += [f"{months}月当前\n百分位", f"{months}月中位\n%"]
        rows = []
        for root in roots:
            q = quantiles[quantiles.root.eq(root)].set_index("months")
            row = [root, fmt(q.iloc[0].current)]
            for months in windows:
                row += [fmt(q.loc[months, "percentile"], 1), fmt(q.loc[months, "p50"])]
            rows.append(row)
        table(caption, headers, rows, font_size=9.5 if len(windows) == 3 else None)

    term_ends = "P5—P95" if provenance.get("configuration", {}).get("quantile_whiskers", "p05_p95") == "p05_p95" else "最小—最大值"
    term_guide = (f"箱体为P25—P75，中线P50，上下须{term_ends}，"
                  "红菱形为当前年化值；横轴同时列出有效样本数和当前百分位，并列值按半权重计入。"
                  "各品种纵轴独立，覆盖不足时不显示分位统计。")
    heading("90日固定期限当前水平与历史分位", page=True)
    picture("90天固定期限_分位蜡烛图", "90日固定期限年化基差6个月和12个月分位")
    term_table("90日固定期限当前水平及历史位置", fixed90_quantiles)
    p("历史窗口为截止日前6、12个自然月，不含当日。" + term_guide)
    p("固定期限分位剔除无法包围90日的数学不可用日期；有插值合约但缺少分红等输入时，不当作可剔除样本。"
      "当前百分位衡量其在固定分红假设下的历史位置，不能据此判断未来必然回归。")

    for pair in PAIRS:
        heading(f"{pair}跨期年化基差", page=True)
        picture(f"跨期{pair}_年化基差", f"{pair}跨期年化基差近3个月")
        term_table(f"{pair}跨期年化基差当前水平及历史位置", calendar_quantiles[calendar_quantiles.pair.eq(pair)])
        p("跨期年化＝［远月期货收盘价＋远月分红−近月期货收盘价−近月分红］÷现货指数收盘值"
          "×365÷两合约实际到期日的自然日间隔×100%。02−00是季月一减当月，03−02是最远月减季月一。")
        p("取同日实际合约配对；符号保留，负值表示调整后远月价格更低。该指标按两到期日间隔年化，"
          "不是两条单合约年化基差相减，近月不足5日也不沿用旧值。虚线表示任一端换月，换月会改变实际配对和期限。")
        current = calendar[calendar.pair.eq(pair) & calendar.trade_date.eq(end)]
        for row in current.itertuples():
            p(f"{row.root}当前配对：{row.far_contract}减{row.near_contract}，到期日间隔{row.expiry_gap_days}天，"
              f"调整价差{fmt(row.spread_points)}点，年化{fmt(row.annualized_basis_pct)}%。")

    for pair in PAIRS:
        heading(f"{pair}跨期年化基差3个月6个月12个月分位", page=True)
        picture(f"跨期{pair}_年化分位蜡烛图", f"{pair}跨期年化基差3个月6个月12个月分位")
        p("历史窗口为截止日前3、6、12个自然月，不含当日。" + term_guide)
        p("上述固定期限和跨期统计均沿用本次G列固定分红快照；跨期用远、近两档各自的分红，"
          "不能用当前假设下的历史分位代替当时可观察的分红预测。")
    doc.core_properties.title = "晨会分析"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(target)
    verified = Document(target)
    if len(verified.inline_shapes) != len(charts):
        raise ValueError("Word 图片数量核验失败")
    return {"figures": len(verified.inline_shapes), "tables": len(verified.tables),
            "paragraphs": len(verified.paragraphs)}


def chart_catalog():
    names = []
    for root in ROOTS:
        names += [f"{root}_{kind}" for kind in ("连续年化基差", "合并年化基差", "年化基差分位蜡烛图",
                  "连续绝对基差", "合并绝对基差", "基差与持仓变化", "跨期价差")]
    return names + ["60天固定期限比较", "IM_IC_绝对基差差分图", "IM_IC_绝对基差差合并图",
                    "IM_IC_绝对基差差分位蜡烛图"] + [f"IM0{i}_IC0{i}_绝对基差差" for i in range(4)] + term_chart_catalog()
