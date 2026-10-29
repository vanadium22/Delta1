"""Downloaded-file opt-in import and actual dated source evidence, no interfaces."""
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from dashboard.dividend_check import main,audit
from dashboard.marketdata.overview_dividends import DividendStore,source_directory,DEFAULT_SOURCE,LEGACY_DEFAULT_SOURCE
from dashboard.marketdata.overview_dividend_sync import source_directories
from dashboard.native.overview_dividend_job import DividendJob
from tests.test_overview_dated import source,snapshot
from tests.test_overview_percentiles import metadata,DAY


class DividendInputTests(unittest.TestCase):
    def test_confirmed_y_source_replaces_only_old_builtin_default_and_retains_custom_or_disabled(self):
        expected=r'Y:\jrcx\for_jrcx\OTC-Quant\pybct\TZQ\日终基差结果\2026'
        self.assertEqual(DEFAULT_SOURCE,expected)
        config=dict(overview_dividend_source_dir=LEGACY_DEFAULT_SOURCE,password='retained');before=dict(config)
        self.assertEqual(source_directory(config),expected);self.assertEqual(config,before)
        self.assertEqual(source_directory({}),expected)
        self.assertEqual(source_directory(dict(overview_dividend_source_dir='')), '')
        self.assertEqual(source_directory(dict(overview_dividend_source_dir='custom')), 'custom')
        self.assertEqual(source_directory(dict(overview_dividend_source_dir=LEGACY_DEFAULT_SOURCE.upper().replace('\\','/')+'/')),expected)
        other=r'\\configured-host\kept-share\dividends'
        self.assertEqual(source_directory(dict(overview_dividend_source_dir=other)),other)
        self.assertEqual(source_directories({},'2026-07-01','2026-10-29'),[expected])
        self.assertEqual(source_directories({},'2025-12-01','2026-01-05'),[expected.replace('2026','2025'),expected])

    def profile(self,root):
        path=root/'native.json';path.write_text('{"password":"DO_NOT_EXPORT"}',encoding='utf8')
        return path,path.read_bytes()

    def test_old_successful_slot_is_rechecked_after_confirmed_source_policy_changes(self):
        from dashboard.marketdata.dividends import digest
        for config in ({},{'overview_dividend_source_dir':LEGACY_DEFAULT_SOURCE}):
            with self.subTest(config=config),tempfile.TemporaryDirectory() as tmp:
                root=Path(tmp);directory=root/'overview_dividends';directory.mkdir()
                old=digest({k:config.get(k) for k in ('overview_dividend_source_dir','overview_dividend_pattern','overview_dividend_file',
                    'dividend_root','overview_dividend_check_am','overview_dividend_check_pm')})
                schedule=directory/'schedule.json'
                schedule.write_text(json.dumps(dict(day='2026-10-29',slot='am',profile=old,complete=True,attempts=1)),encoding='utf8')
                with patch('dashboard.native.overview_dividend_job.subprocess.Popen') as spawn:
                    DividendJob(root).tick(config,enabled=True,now=datetime.fromisoformat('2026-10-29T10:00:00+08:00'))
                spawn.assert_called_once();self.assertNotEqual(json.loads(schedule.read_text())['profile'],old)

    def test_settings_displays_effective_source_without_saving_personal_profile(self):
        from types import SimpleNamespace
        from dashboard.native.bootstrap import configure_qt
        configure_qt()
        from PySide6.QtWidgets import QApplication
        from dashboard.native.inline_settings import InlineSettings
        app=QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);native,before=self.profile(root)
            config=dict(overview_dividend_source_dir=LEGACY_DEFAULT_SOURCE)
            pipeline=SimpleNamespace(model=SimpleNamespace(config=config,local_root=root))
            settings=InlineSettings(pipeline)
            try:
                self.assertEqual(settings.fields['overview_dividend_source_dir'].text(),DEFAULT_SOURCE)
                self.assertEqual(config['overview_dividend_source_dir'],LEGACY_DEFAULT_SOURCE)
                self.assertEqual(native.read_bytes(),before)
            finally:
                settings.close();settings.deleteLater();app.processEvents()

    def test_explicit_download_import_retains_source_and_does_not_query_share_or_edit_config(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);path,before=self.profile(root)
            downloaded=source(root,'misleading_2026-09-18.xlsx',DAY,2)
            original=downloaded.read_bytes();out=root/'import.json'
            with patch('dashboard.marketdata.overview_dividends.read_snapshot',side_effect=snapshot),\
                    patch('dashboard.marketdata.contract_reference.latest_contracts',return_value=metadata()),\
                    patch('dashboard.marketdata.overview_dividend_sync.run_sync',side_effect=AssertionError('file import must not read share')):
                exit_code=main(['--local-root',str(root),'--stage','import','--input-file',str(downloaded),'--day',DAY,'--output',str(out)])
                value=json.loads(out.read_text(encoding='utf8'))
                self.assertEqual(exit_code,2) # current file cannot fill old history
                self.assertEqual(value['file_import']['inserted'],16)
                self.assertEqual(value['file_import']['business_dates'],[DAY])
                self.assertFalse(value['network_requested']);self.assertNotIn('sync',value)
                self.assertEqual(value['audit']['current_dividends']['IF00.CFE']['points'],2)
                self.assertEqual(value['audit']['current_dividends']['IF00.CFE']['source_business_date'],DAY)
                main(['--local-root',str(root),'--stage','import','--input-file',str(downloaded),'--output',str(out)])
                self.assertEqual(json.loads(out.read_text(encoding='utf8'))['file_import']['status'],'unchanged')
            self.assertEqual(path.read_bytes(),before);self.assertEqual(downloaded.read_bytes(),original)
            self.assertNotIn('DO_NOT_EXPORT',out.read_text(encoding='utf8'))

    def test_local_reports_only_configured_inputs_and_dated_contract_selection(self):
        with tempfile.TemporaryDirectory() as tmp,patch('dashboard.marketdata.overview_dividends.read_snapshot',side_effect=snapshot),\
                patch('dashboard.marketdata.contract_reference.latest_contracts',return_value=metadata()):
            root=Path(tmp);self.profile(root)
            old=source(root,'old.xlsx',DAY,1);new=source(root,'new.xlsx','2026-10-28',3)
            for item in (old,new):DividendStore(root/'overview_dividends').import_file(item,now=datetime.fromisoformat('2026-10-29T10:00:00+08:00'))
            now=datetime.fromisoformat('2026-10-29T10:00:00+08:00')
            current=audit(dict(overview_dividend_file=str(new)),root,now=now)
            past=audit({},root,now=now,day=DAY)
            self.assertFalse(current['sources']['downloads_auto_scanned'])
            self.assertEqual(current['sources']['local_files'][0]['business_dates'],['2026-10-28'])
            self.assertEqual(current['oldest_business_day'],DAY)
            self.assertEqual(current['latest_business_day'],'2026-10-28')
            self.assertEqual(current['current_dividends']['IF00.CFE']['points'],3)
            self.assertEqual(past['current_dividends']['IF00.CFE']['points'],1)
            self.assertEqual(current['sources']['store_directory'],str((root/'overview_dividends').resolve()))

    def test_input_file_requires_explicit_import_and_invalid_file_preserves_store(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);self.profile(root)
            with self.assertRaises(SystemExit):main(['--local-root',str(root),'--input-file',str(root/'missing.xlsx')])
            out=root/'bad.json'
            result=main(['--local-root',str(root),'--stage','import','--input-file',str(root/'missing.xlsx'),'--output',str(out)])
            self.assertEqual(result,2);self.assertEqual(json.loads(out.read_text())['error_type'],'ValueError')
            self.assertFalse((root/'overview_dividends/records.sqlite').exists())

    @unittest.skipUnless(os.name=='nt','Windows command launcher')
    def test_real_cmd_forwards_chinese_space_download_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);profile=root/'company profile';profile.mkdir();path,before=self.profile(profile)
            # A genuine minimal workbook exercises the existing cached-value
            # parser through CMD/PowerShell, with no SDK or network involved.
            import openpyxl
            book=openpyxl.Workbook();sheet=book.active;sheet.title='Sheet1';sheet['B1']=datetime(2026,10,27)
            for i,(product,month) in enumerate((p,m) for p in ('IF','IH','IC','IM') for m in ('2611','2612','2703','2706')):
                slot=('2611','2612','2703','2706').index(month)
                row=i+3;sheet.cell(row,2,product+'0'+str(slot)+'.CFE');sheet.cell(row,3,product+month+'.CFE');sheet.cell(row,7,0)
            downloaded=root/'下载 新分红.xlsx';book.save(downloaded);book.close()
            sha=hashlib.sha256(downloaded.read_bytes()).hexdigest();out=root/'report output.json'
            command=Path(__file__).resolve().parents[1]/'dashboard/diagnose_dividends.cmd'
            wrapper=root/'invoke.ps1'
            wrapper.write_text('param($Command,$ProfilePath,$InputPath,$OutputPath)\n& $Command -LocalRoot $ProfilePath -Stage Import -InputFile $InputPath -Output $OutputPath -NoPause\n',encoding='utf8')
            result=subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(wrapper),
                str(command),str(profile),str(downloaded),str(out)],capture_output=True,text=True,encoding='utf8',errors='replace',
                timeout=25,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            self.assertTrue(out.is_file(),result.stdout+'\n'+result.stderr)
            value=json.loads(out.read_text(encoding='utf8'))
            self.assertEqual(value['file_import']['sha256'],sha);self.assertEqual(value['file_import']['records'],16)
            self.assertEqual(value['file_import']['business_dates'],['2026-10-27']);self.assertFalse(value['network_requested'])
            self.assertEqual(path.read_bytes(),before)
            year=root/'2026';year.mkdir();(year/'指数基差情况_2026-01-02.xlsx').write_bytes(downloaded.read_bytes())
            wrapper.write_text('param($Command,$ProfilePath,$SourcePath,$OutputPath)\n& $Command -LocalRoot $ProfilePath -Stage Share -SourceDirectory $SourcePath -Output $OutputPath -NoPause\n',encoding='utf8')
            subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(wrapper),
                str(command),str(profile),str(year),str(out)],capture_output=True,timeout=25,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            shared=json.loads(out.read_text(encoding='utf8'))
            self.assertEqual(shared['sync']['source_directories'],[str(year)])
            self.assertEqual(shared['sync']['validated_files'],1);self.assertFalse(shared['sync']['published'])
            self.assertEqual(shared['sync']['business_dates'],['2026-10-27'])
            self.assertEqual(path.read_bytes(),before)

    def test_temporary_directory_override_reads_requested_year_without_editing_settings(self):
        with tempfile.TemporaryDirectory() as tmp,patch('dashboard.marketdata.overview_dividends.read_snapshot',side_effect=snapshot),\
                patch('dashboard.marketdata.contract_reference.latest_contracts',return_value=metadata()):
            root=Path(tmp);native,before=self.profile(root);share=root/'2026';share.mkdir()
            source(share,'指数基差情况_2026-10-27.xlsx',DAY,2)
            from dashboard.marketdata.overview_dividend_sync import run_sync,stage_inputs
            def local_reader(config,profile):return run_sync(config,profile,reader=stage_inputs,now=datetime.fromisoformat('2026-10-29T10:00:00+08:00'))
            with patch('dashboard.marketdata.overview_dividend_sync.run_sync',side_effect=local_reader):
                out=root/'override.json'
                main(['--local-root',str(root),'--stage','import','--source-directory',str(share),'--output',str(out)])
            value=json.loads(out.read_text(encoding='utf8'))
            self.assertEqual(value['sync']['source_directories'],[str(share)])
            self.assertEqual(value['sync']['matched_names'],['指数基差情况_2026-10-27.xlsx'])
            self.assertEqual(value['sync']['inserted'],16);self.assertEqual(native.read_bytes(),before)
            self.assertIsNone(value['audit']['sources']['configured_share_directory'])
            self.assertEqual(value['audit']['sources']['share_directory'],str(share))
            self.assertEqual(value['diagnostic_profile']['native_file'],str(native.resolve()))

    def test_share_uses_business_date_even_when_filename_date_is_outside_window(self):
        with tempfile.TemporaryDirectory() as tmp,patch('dashboard.marketdata.overview_dividends.read_snapshot',side_effect=snapshot) as parse:
            from dashboard.marketdata.overview_dividend_sync import run_sync,stage_inputs
            root=Path(tmp);share=root/'2026';share.mkdir()
            current=source(share,'指数基差情况_2026-01-02.xlsx',DAY,2)
            (share/'指数基差情况_2026-01-01.xlsx').write_bytes(current.read_bytes())
            source(share,'指数基差情况_2026-10-28.xlsx','2026-01-02',1)
            source(share,'指数基差情况_2026-10-29.xlsx','2026-10-30',3)
            config=dict(overview_dividend_source_dir=str(share))
            now=datetime.fromisoformat('2026-10-29T10:00:00+08:00')
            result=run_sync(config,root,reader=stage_inputs,now=now)
            self.assertEqual(result['status'],'complete');self.assertEqual(result['inserted'],16)
            self.assertEqual(result['outside_range_files'],2);self.assertEqual(result['business_dates'],[DAY])
            self.assertEqual(parse.call_count,3)
            self.assertEqual({r['source_business_date'] for r in DividendStore(root/'overview_dividends').read()},{DAY})
            with patch.object(Path,'read_bytes',side_effect=AssertionError('unchanged share file read')):
                again=run_sync(config,root,reader=stage_inputs,now=now)
            self.assertEqual(again['unchanged'],4);self.assertEqual(again['inserted'],0)

    def test_source_override_and_input_file_are_explicit_mutually_exclusive_stages(self):
        with tempfile.TemporaryDirectory() as tmp:
            for extra in (['--source-directory',tmp],['--stage','share','--input-file',tmp],
                    ['--stage','import','--input-file',tmp,'--source-directory',tmp]):
                with self.subTest(args=extra),self.assertRaises(SystemExit):main(['--local-root',tmp,*extra])


if __name__=='__main__':unittest.main()
