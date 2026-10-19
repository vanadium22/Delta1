from datetime import date
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
import openpyxl
import pandas as pd

from dashboard.marketdata.bars import BarDataset, normalize_bars
from dashboard.marketdata.ddb_import import import_runs, merge_checked, validate_standard
from dashboard.marketdata.dividends import DividendBook, adjusted_basis, make_scenario, read_snapshot
from dashboard.marketdata.quote_quality import aligned_quotes, quote_quality
from dashboard.marketdata.repository import ReferenceBook
from dashboard.native.nonregression import statistics, efficiency, weighted, single, minute_diagnostics, minute_daily, collection_center, MULTIPLIERS
from index_strategy.instant_dld.quotes import normalize_quote
from index_strategy.instant_dld.realtime_store import RealtimeStore


def quote(total=100,time=100000,**values):
    return dict(latestPrice=4000,openingPrice=3990,highPrice=4010,lowPrice=3980,yesterdayClosePrice=3988,
        tradedQuantities=total,quoteTimestamp=time,openInterest=0,closingPrice=0,final=False,
        b1Price=3999.8,s1Price=4000.2,b1Stocks=12,s1Stocks=14,b2Price=9223372036854.775,**values)


class Quotes(unittest.TestCase):
    def test_fields_zero_oi_sentinel_and_final(self):
        r=normalize_quote('IF2612.CFE',quote(),'2026-10-16T10:00:00+08:00')
        self.assertEqual([r[k] for k in ('close','open','high','low','prev_close')],[4000,3990,4010,3980,3988])
        self.assertEqual(r['open_interest_raw'],0);self.assertIsNone(r['open_interest'])
        self.assertIsNone(r['bid_price_2']);self.assertFalse(r['is_final'])
        self.assertIn('unverified',r['date_status']);self.assertIsNone(r['volume'])

    def test_delta_duplicate_gap_rollback_day_and_reconnect(self):
        a=normalize_quote('IF2612.CFE',quote(),'2026-10-16T10:00:00+08:00',expected_interval=2)
        b=normalize_quote('IF2612.CFE',quote(105,100002),'2026-10-16T10:00:02+08:00',a,expected_interval=2)
        self.assertEqual(b['volume'],5)
        dup=normalize_quote('IF2612.CFE',quote(105,100002),'2026-10-16T10:00:04+08:00',b,expected_interval=2)
        self.assertEqual(dup['volume'],0);self.assertEqual(dup['snapshot_status'],'duplicate_source_snapshot')
        for ts,total in [('2026-10-16T10:00:30+08:00',150),('2026-10-17T10:00:04+08:00',200),('2026-10-16T10:00:04+08:00',90)]:
            self.assertIsNone(normalize_quote('IF2612.CFE',quote(total,100004),ts,b,expected_interval=2)['volume'])
        with TemporaryDirectory() as directory:
            with RealtimeStore(Path(directory)) as store:
                store.interval_seconds=2
                def batch(poll,total,time,status='success'):
                    return dict(received_at=f'2026-10-16T10:00:{time:02d}+08:00',symbols=['IF2612.CFE'],status=status,poll=poll,batch=1,batch_count=1,response={'data':{'IF2612.CFE':quote(total,100000+time)}})
                store.batch(batch(1,100,0));store.batch(batch(2,200,2,'timeout'));store.batch(batch(3,250,4))
                self.assertIsNone(store.last_rows[0]['volume'])
                store.batch(batch(4,260,6));self.assertEqual(store.last_rows[0]['volume'],10)
                self.assertEqual(store.connection.execute('select open_interest_raw,open_interest,open from quotes order by sequence desc').fetchone()[:],(0,None,3990))

    def test_receipt_is_not_fresh_source_or_same_time(self):
        q=normalize_quote('IF2612.CFE',quote(),'2026-10-16T10:00:00+08:00')
        q['observed_at']=pd.Timestamp(q['received_at'])
        self.assertTrue(quote_quality(q,'2026-10-16T10:00:02+08:00')['usable'])
        cached={**q,'received_at':'2026-10-16T10:05:00+08:00','observed_at':pd.Timestamp('2026-10-16T10:05:00+08:00')}
        self.assertFalse(quote_quality(cached,'2026-10-16T10:05:01+08:00')['usable'])
        self.assertFalse(aligned_quotes([q,{**q,'source_event_time':'2026-10-16T09:59:55+08:00'}],'2026-10-16T10:00:02+08:00')[0])
        self.assertFalse(quote_quality(q,'2026-10-16T12:00:00+08:00')['usable'])
        close=normalize_quote('IF2612.CFE',quote(time=150000),'2026-10-16T15:00:00+08:00')
        self.assertFalse(quote_quality(close,'2026-10-16T15:00:01+08:00')['usable'])

    def test_observation_center_uses_deltas_and_resets_gaps(self):
        previous={};samples=[];live={}
        for poll,(f,total) in enumerate([(90,100),(92,110),(96,130)],1):
            for symbol,price in [('IF2612.CFE',f),('000300.SH',100)]:
                r=normalize_quote(symbol,{**quote(total,100000+poll*2),'latestPrice':price},f'2026-10-16T10:00:{poll*2:02d}+08:00',previous.get(symbol),expected_interval=2)
                previous[symbol]=r
                r.update(instrument_id=symbol,observed_at=pd.Timestamp(r['received_at']),run_id='test',poll=poll,batch=1)
                samples.append(r);live[symbol]=r
        book=DividendBook(reason='',info={'mapping':{'IF2612.CFE':{'points':5,'expiry':'2026-12-18'}}})
        c,gap,_,note=collection_center(pd.DataFrame(samples),'IF2612.CFE','000300.SH','basis','basis',book,{},live,'2026-10-16T10:00:07+08:00')
        self.assertAlmostEqual(c,-1000/30);self.assertAlmostEqual(gap,100-c);self.assertIn('采集区间',note)
        samples[-2]['quality_flags']|=128
        self.assertIsNone(collection_center(pd.DataFrame(samples),'IF2612.CFE','000300.SH','basis','basis',book,{},live,'2026-10-16T10:00:07+08:00')[0])


class MathTests(unittest.TestCase):
    def test_dividend_sign_simple_annual_expiry(self):
        r=adjusted_basis(90,100,5,10)
        self.assertEqual(r['points'],-5);self.assertEqual(r['bp'],-500);self.assertEqual(r['annual_pct'],-182.5)
        for days in (-1,0,4):self.assertIsNone(adjusted_basis(90,100,5,days)['annual_pct'])
        self.assertIsNone(adjusted_basis(90,100,None,10)['bp'])
        self.assertEqual((90+5)-(80+8),7)
        self.assertEqual(MULTIPLIERS,{'IF':300,'IH':300,'IC':200,'IM':200})

    def test_signed_sequence_prior_only_missing_and_zero_weights(self):
        s=pd.Series(np.arange(-20.,2.));z,p,_=statistics(s)
        self.assertAlmostEqual(z.iloc[-1],(1-np.arange(-19.,1.).mean())/np.arange(-19.,1.).std(ddof=1))
        self.assertEqual(p.iloc[-1],100);self.assertEqual(efficiency(s).iloc[-1],1)
        self.assertTrue(statistics(pd.Series([0.]*21))[0].isna().all())
        s.iloc[-3]=np.nan;self.assertTrue(pd.isna(statistics(s)[0].iloc[-1]));self.assertTrue(pd.isna(efficiency(s).iloc[-1]))
        self.assertTrue(weighted(pd.Series(range(10)),pd.Series([0.]*10)).isna().all())

    def test_missing_dividend_preserves_unrelated_price_volume(self):
        dates=pd.date_range('2026-10-01 07:00',periods=8,tz='UTC')
        f=pd.DataFrame(dict(close=np.arange(90.,98.),open=90,high=100,low=80,volume=100,open_interest=100),index=dates)
        s=pd.DataFrame(dict(close=100),index=dates)
        r=single(f,s,'IF2610.CFE','2026-10-08',DividendBook(),dates)
        self.assertTrue(r.basis.isna().all());self.assertTrue(r.future.notna().all())
        self.assertTrue(pd.isna(r.doi.iloc[-1]));self.assertEqual(r.open_interest.iloc[-1],100)
        # Missing middle source date invalidates adjacent OI changes, not a false decline.
        r=single(f.drop(dates[3]),s,'IF2610.CFE','2026-10-08',DividendBook(),dates)
        self.assertTrue(r.doi.iloc[3:5].isna().all())

    def test_minute_gap_session_and_sample_ohlc(self):
        t=pd.date_range('2026-10-09 01:31Z',periods=70,freq='min')
        f=pd.DataFrame(dict(x=np.linspace(-2,2,70),volume=1,session_id='am',trade_date=date(2026,10,9)),index=t)
        rv,_=minute_diagnostics(f,'basis');self.assertTrue(pd.notna(rv.iloc[-1]))
        rv,_=minute_diagnostics(f.drop(t[50]),'basis');self.assertTrue(pd.isna(rv.iloc[-1]))
        self.assertTrue(minute_daily(f).empty)


class DividendTests(unittest.TestCase):
    def workbook(self,path,value=5):
        w=openpyxl.Workbook();s=w.active;s.title='Sheet1';s['B1']=date(2026,9,18)
        for i,(p,k) in enumerate((p,k) for p in ('IF','IH','IC','IM') for k in range(4)):
            s.cell(i+4,2,f'{p}0{k}.CFE');s.cell(i+4,3,f'{p}9901.CFE');s.cell(i+4,5,'2099-01-01');s.cell(i+4,7,value+k)
        w.create_sheet('Sheet2')['B1']=0;w.save(path);w.close()

    def test_latest_date_dedup_same_day_conflict_and_no_roll(self):
        with TemporaryDirectory() as d:
            path=Path(d);self.workbook(path/'a.xlsx');(path/'b.xlsx').write_bytes((path/'a.xlsx').read_bytes())
            (path/'~$tmp.xlsx').write_bytes(b'ignored')
            snap=read_snapshot(path);self.assertEqual(len(snap['equivalent_files']),2)
            ref=ReferenceBook();ref.version='verified'
            ref.data['contracts']=[dict(instrument_id=p+m+'.CFE',listing_date='2026-01-01',expiry_date=e,known_at='2026-10-15T00:00:00Z') for p in ('IF','IH','IC','IM') for m,e in [('2610','2026-10-16'),('2611','2026-11-20'),('2612','2026-12-18'),('2703','2027-03-19')]]
            scenario=make_scenario(snap,ref,'2026-10-16');file=path/'scenario.json';file.write_text(json.dumps(scenario),encoding='utf-8')
            book=DividendBook.load(path,file)
            self.assertEqual(book.points('IF2610.CFE','2026-10-14'),5)
            self.assertIsNone(book.points('IF2610.CFE','2026-10-17'));self.assertIsNone(book.points('IF2706.CFE','2026-10-17'))
            self.assertEqual(book.points('IF2611.CFE','2026-10-20'),6)
            self.workbook(path/'conflict.xlsx',9)
            with self.assertRaisesRegex(ValueError,'同日冲突'):read_snapshot(path)


class Imports(unittest.TestCase):
    def fixture(self,root,bulk=True,close=4000):
        root.mkdir(parents=True);f=root/'daily.parquet'
        pd.DataFrame([dict(source_symbol='IF2612.CFE',source_code='IF2612.CFE',symbol='IF2612.CFE',trading_date=pd.Timestamp('2026-10-09'),timestamp=pd.Timestamp('2026-10-09'),open=3990.,high=4010.,low=3980.,close=float(close),volume=10.,amount=12000000.,oi=0.,quality='complete')]).to_parquet(f,index=False)
        item=dict(file=str(f),rows=1,bytes=f.stat().st_size,sha256=hashlib.sha256(f.read_bytes()).hexdigest(),frequency='1d',kind='normalized',schema_version=2,source_symbol='IF2612.CFE',downloaded_at='2026-10-15T00:00:00Z',units={'amount':'元','volume':'手'},start='2026-10-09',end='2026-10-09')
        manifest=dict(state='complete',files=[item],created_at='2026-10-15T00:00:00Z')
        manifest.update(schema='basis_research_bulk_v1') if bulk else manifest.update(version=2)
        run=root/'run.json';run.write_text(json.dumps(manifest),encoding='utf-8');return run

    def test_formats_relocation_idempotency_oi_conflict_and_hash(self):
        with TemporaryDirectory() as d:
            root=Path(d);ref=root/'ref.json';ref.write_text(json.dumps({'schema_version':'cn_index_reference_v1','contracts':[dict(instrument_id='IF2612.CFE',listing_date='2026-04-20',expiry_date='2026-12-18',known_at='2026-10-15T00:00:00Z')],'aliases':[]}),encoding='utf-8')
            a=self.fixture(root/'batch');out=root/'normalized'
            first=import_runs([a],out,ref);second=import_runs([a],out,ref)
            self.assertEqual(first['datasets'][0]['inserted'],1);self.assertEqual(second['datasets'][0]['inserted'],0)
            self.assertEqual(first['datasets'][0]['version'],second['datasets'][0]['version'])
            b=self.fixture(root/'older_v2',bulk=False);third=import_runs([b],out,ref);self.assertEqual(third['datasets'][0]['duplicates'],1)
            store=BarDataset(out/'cn_index_future');self.assertEqual(store._rows(store.snapshot())[0]['open_interest'],0)
            validate_standard(store.root)
            c=self.fixture(root/'conflict',close=4001)
            with self.assertRaisesRegex(ValueError,'来源冲突'):import_runs([c],out,ref)
            self.assertEqual(store.snapshot()['version'],first['datasets'][0]['version'])
            (a.parent/'daily.parquet').write_bytes(b'bad hash')
            with self.assertRaisesRegex(ValueError,'哈希'):import_runs([a],out,ref)


if __name__=='__main__':unittest.main()
