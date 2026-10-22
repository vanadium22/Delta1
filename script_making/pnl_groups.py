"""Conservative group discovery from saved workbook styles, never rendered HTML."""
from decimal import Decimal
from openpyxl import load_workbook
from .query import HEADER_ROWS
from .storage import ReportError
from .pnl_sectors import option_sectors


def discover(path, grid, focus):
    book = load_workbook(path, data_only=True, keep_links=False)
    try:
        ws = book['累积盈亏统计']
        header = HEADER_ROWS[ws.title]
        columns = {str(c.value).strip(): c.column for c in ws[header] if c.value}
        name_col, value_col = columns['交易台'], columns['净利润']
        merges = list(ws.merged_cells.ranges)
        sectors = option_sectors(ws, grid, name_col, value_col)
        diagnostics, groups, totals, candidates = [], [], [], []
        def label(row):
            cell = ws.cell(row, name_col)
            merged = next((m for m in merges if m.min_row <= row <= m.max_row and m.min_col <= name_col <= m.max_col), None)
            if merged:
                cell = ws.cell(merged.min_row, merged.min_col)
            return cell, merged
        anchor, _ = label(focus)
        level = (anchor.alignment.indent or 0, ws.row_dimensions[focus].outlineLevel)
        if not anchor.font.bold:
            diagnostics.append(f'第{focus}行交易台名称未在原Excel加粗，无法确认汇总层级。')
        for row in range(header + 1, ws.max_row + 1):
            if sectors and row in sectors['child_rows']:
                continue
            cell, merged = label(row)
            if merged and merged.min_row != row:
                continue
            name = str(cell.value or '').strip()
            if not name or name.startswith(('注', '说明', '单位')) or 'T-2' in name or name == '交易台':
                continue
            if name == '场外业务合计':
                totals.append(row)
                continue
            if not cell.font.bold:
                continue
            if any(word in name for word in ('合计', '总计', '小计')):
                diagnostics.append(f'第{row}行存在额外汇总层级“{name}”，排名暂停。')
                continue
            if merged and ((merged.max_row != row and not (sectors and row == sectors['parent']))
                           or merged.min_col <= value_col <= merged.max_col):
                diagnostics.append(f'第{row}行合并范围{merged}无法区分组汇总与明细。')
                continue
            candidates.append(row)
            row_level = (cell.alignment.indent or 0, ws.row_dimensions[row].outlineLevel)
            if row_level != level:
                diagnostics.append(f'第{row}行“{name}”与目标台缩进/折叠层级不同，排名暂停。')
                continue
            if focus < row <= focus + 4:
                continue
            raw = ws.cell(row, value_col).value
            try:
                value = Decimal(str(raw))
                if not value.is_finite():
                    raise ValueError()
            except Exception:
                diagnostics.append(f'第{row}行“{name}”净利润缓存缺失或无效，未按零处理。')
                continue
            groups.append({'name': name, 'row': row, 'value': str(value)})
        if sectors:
            parent = sectors['parent']
            if any(focus <= r <= focus + 4 for r in sectors['child_rows'] | {parent}):
                raise ReportError('行业明细与境内权益互换台指定五行重叠，未混用两组来源。')
            replacement = [{k: v for k, v in item.items() if k != 'values'} for item in sectors['entries']]
            if not any(g['row'] == parent for g in groups):
                diagnostics.append('境内个股期权台未通过原Excel组汇总样式识别。')
            groups = [entry for g in groups for entry in (replacement if g['row'] == parent else [g])]
            candidates = [entry for r in candidates for entry in
                          ([g['row'] for g in replacement] if r == parent else [r])]
            diagnostics.extend(sectors['notes'])
        if len(totals) != 1:
            diagnostics.append('场外业务合计不是唯一一行，无法确认完整分组边界。')
        if len({g['name'] for g in groups}) != len(groups):
            diagnostics.append('组名称重复，无法排除重复汇总。')
        if not any(g['row'] == focus for g in groups):
            diagnostics.append('目标台未通过原Excel汇总识别。')
        if totals and any(g['row'] > totals[0] for g in groups):
            diagnostics.append('总计之后仍存在加粗组，无法确认同层级范围。')
        if len(totals) == 1:
            try:
                total = Decimal(str(ws.cell(totals[0], value_col).value))
                if not total.is_finite() or abs(sum((Decimal(g['value']) for g in groups), Decimal(0)) - total) > Decimal('0.01'):
                    diagnostics.append('同层级组净利润之和与场外业务合计不一致，可能遗漏或重复层级，排名暂停。')
            except Exception:
                diagnostics.append('场外业务合计数值缓存缺失，无法校验分组完整性。')
        selected = list(range(focus, focus + 5))
        selected += [r for r in candidates if r not in selected and r not in totals]
        selected = [r for r in selected if r not in totals] + totals[:1]
        lookup = dict(zip(grid['source_rows'], grid['rows']))
        source_map = {}
        if sectors:
            for item in sectors['entries']:
                lookup[item['row']] = item['values']
                source_map[item['row']] = item['source_rows']
        empty = [None] * len(grid['columns'])
        return {'rows': selected, 'values': [lookup.get(r, empty) for r in selected],
                'row_sources': [source_map.get(r, [r]) for r in selected],
                'headers': lookup.get(header, empty), 'total_rows': totals[:1],
                'groups': groups, 'diagnostics': diagnostics,
                'otc': {'row': totals[0], 'value': str(ws.cell(totals[0], value_col).value)}
                       if len(totals) == 1 and ws.cell(totals[0], value_col).value is not None else None,
                'ranking': [] if diagnostics else sorted(groups, key=lambda g: -Decimal(g['value']))}
    finally:
        book.close()
