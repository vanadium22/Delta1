"""Local configuration and verified, dated report copies. No business API calls."""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from threading import RLock
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from shared_data import Catalog, DataPaths, Dataset
from shared_data.io import file_lock, portable_manifest, resolved_manifest
from shared_data.migration import ensure_migrated

CHINA = timezone(timedelta(hours=8))
REPORTS = {
    "swap": {"label": "收益互换日报", "sheet": "存续合约估值"},
    "review": {"label": "日报复核表", "sheet": "簿记要素"},
}
DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}\Z")
FILE_RE = re.compile(r"(收益互换日报|日报复核表)\((\d{4}-\d{2}-\d{2})\)-D1\.xlsx\Z")


class ReportError(ValueError):
    pass


def today() -> date:
    return datetime.now(CHINA).date()


def timestamp() -> str:
    return datetime.now(CHINA).isoformat(timespec="seconds")


def valid_date(value: str) -> str:
    if not isinstance(value, str) or not DATE_RE.fullmatch(value):
        raise ReportError("日期格式应为 YYYY-MM-DD。")
    try:
        date.fromisoformat(value)
    except ValueError as exc:
        raise ReportError("日期无效。") from exc
    return value


def filename(kind: str, day: str) -> str:
    if kind not in REPORTS:
        raise ReportError("报表类型无效。")
    return f"{REPORTS[kind]['label']}({valid_date(day)})-D1.xlsx"


def identify(name: str) -> tuple[str, str]:
    match = FILE_RE.fullmatch(name)
    if not match:
        raise ReportError("文件名应为 收益互换日报(YYYY-MM-DD)-D1.xlsx 或 日报复核表(YYYY-MM-DD)-D1.xlsx。")
    kind = next(k for k, v in REPORTS.items() if v["label"] == match[1])
    return kind, valid_date(match[2])


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


class Settings:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.paths = DataPaths.for_module(self.root)
        ensure_migrated("script_making", self.root, self.paths)
        self.path = self.root / "config.local.json"
        self.lock = RLock()

    def get(self) -> dict:
        with self.lock:
            values = {"source_root": "", "data_dir": "@shared", "epsilon": 1e-8}
            if self.path.exists():
                try:
                    values.update(json.loads(self.path.read_text(encoding="utf-8-sig")))
                except (ValueError, OSError) as exc:
                    raise ReportError("本机配置无法读取，请检查 config.local.json。") from exc
            values["source_root"] = os.environ.get("SCRIPT_MAKING_SOURCE", values["source_root"])
            values["data_dir"] = os.environ.get("SCRIPT_MAKING_DATA", values["data_dir"])
            return values

    def data_path(self, values: dict | None = None) -> Path:
        value = (values or self.get())["data_dir"]
        path = self.paths.resolve_module_setting(self.root, value, self.paths.reports, ("data/reports",))
        if path == self.paths.reports:
            Catalog(self.paths).register(Dataset("swap_reports", "company_share", "snapshot"), path)
        return path

    def update(self, incoming: dict) -> dict:
        from .query import parse_epsilon
        with self.lock:
            values = self.get()
            for key in ("source_root", "data_dir"):
                if key in incoming:
                    value = str(incoming[key]).strip()
                    if "\x00" in value or len(value) > 2048:
                        raise ReportError("目录无效。")
                    if key == "data_dir" and not value:
                        raise ReportError("请填写本地保存目录。")
                    values[key] = value
            if "epsilon" in incoming:
                values["epsilon"] = parse_epsilon(incoming["epsilon"])
            atomic_json(self.path, {k: values[k] for k in ("source_root", "data_dir", "epsilon")})
            return self.get()


def validate_workbook(path: Path, kind: str) -> None:
    try:
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 1024 * 1024 * 1024:
                raise ReportError("工作簿解压后超过 1 GB，请先缩减文件。")
            if archive.testzip():
                raise ReportError("Excel 文件损坏。")
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if REPORTS[kind]["sheet"] not in book.sheetnames:
                raise ReportError(f"文件缺少默认工作表「{REPORTS[kind]['sheet']}」。")
        finally:
            book.close()
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        if isinstance(exc, ReportError):
            raise
        raise ReportError(f"无法读取 Excel：{path.name}。") from exc


def manifest_for(folder: Path) -> dict:
    path = folder / "manifest.json"
    if path.is_file():
        try:
            return resolved_manifest(json.loads(path.read_text(encoding="utf-8")), folder.parent)
        except (ValueError, OSError):
            pass
    return {"date": folder.name, "files": {}}


def save_files(data_dir: Path, sources: dict[str, Path], day: str, origin: str) -> dict:
    try:
        with file_lock(data_dir / ".writer.lock"):
            return _save_files(data_dir, sources, day, origin)
    except TimeoutError as exc:
        raise ReportError(str(exc)) from exc


def _save_files(data_dir: Path, sources: dict[str, Path], day: str, origin: str) -> dict:
    """Stage and validate every input before touching the destination."""
    valid_date(day)
    if not sources or any(kind not in REPORTS for kind in sources):
        raise ReportError("未选择有效报表。")
    data_dir.mkdir(parents=True, exist_ok=True)
    folder = data_dir / day
    with tempfile.TemporaryDirectory(prefix=".incoming-", dir=data_dir) as temporary:
        staged = Path(temporary)
        entries = {}
        for kind, source in sources.items():
            target = staged / filename(kind, day)
            before = source.stat()
            shutil.copyfile(source, target)
            after = source.stat()
            if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
                raise ReportError("来源文件仍在生成，请稍后重试。")
            if target.stat().st_size != before.st_size:
                raise ReportError("文件复制不完整，请重试。")
            validate_workbook(target, kind)
            digest = sha256(target)
            if sha256(source) != digest:
                raise ReportError("来源与本地文件校验不一致，请重试。")
            destination = folder / target.name
            if destination.exists() and sha256(destination) != digest:
                raise ReportError(f"{day} 已有内容不同的同名文件，已保留原文件。可更换保存目录后重新下载。")
            entries[kind] = {"name": target.name, "size": before.st_size,
                             "sha256": digest, "source": str(source), "origin": origin,
                             "saved_at": timestamp(), "path": str(destination.resolve())}
        folder.mkdir(exist_ok=True)
        manifest = manifest_for(folder)
        for kind, entry in entries.items():
            destination = folder / entry["name"]
            if not destination.exists():
                os.replace(staged / entry["name"], destination)
            if not destination.is_file() or sha256(destination) != entry["sha256"]:
                raise ReportError("保存后的文件核验失败。")
            previous = manifest["files"].get(kind)
            if previous and origin == "共享目录下载":
                entry["original_origin"] = previous.get("original_origin", previous["origin"])
                entry["first_saved_at"] = previous.get("first_saved_at", previous["saved_at"])
                manifest["files"][kind] = entry
            else:
                manifest["files"].setdefault(kind, entry)
        manifest["updated_at"] = timestamp()
        atomic_json(folder / "manifest.json", portable_manifest(manifest, data_dir))
    return {"date": day, "files": manifest["files"], "saved_dir": str(folder.resolve())}


def sync_reports(values: dict, data_dir: Path, as_of: date | None = None, *, only_missing=False) -> dict:
    if not values.get("source_root"):
        raise ReportError("请先在「数据与同步」中设置共享来源目录。")
    root = Path(values["source_root"])
    cutoff = as_of or today()
    try:
        days = []
        for child in root.iterdir():
            if DATE_RE.fullmatch(child.name):
                try:
                    parsed = date.fromisoformat(child.name)
                except ValueError:
                    continue
                if parsed < cutoff and child.is_dir():
                    days.append(parsed)
        if not days:
            raise ReportError("共享目录内没有早于今天的日期文件夹。")
        day = max(days).isoformat()
        if only_missing:
            manifest = manifest_for(data_dir / day)
            valid = True
            for kind in REPORTS:
                path = data_dir / day / filename(kind, day)
                entry = manifest.get("files", {}).get(kind, {})
                try:
                    if not path.is_file() or not entry.get("sha256") or sha256(path) != entry["sha256"]:
                        valid = False
                        break
                    validate_workbook(path, kind)
                except (ReportError, OSError):
                    valid = False
                    break
            if valid:
                return {**manifest, "reused": True,
                        "selection_rule": "早于今日的最新日期目录（依赖上游按交易日生成）"}
        sources = {kind: root / day / filename(kind, day) for kind in REPORTS}
        missing = [p.name for p in sources.values() if not p.is_file()]
        if missing:
            raise ReportError(f"最新日期 {day} 的两份文件尚未齐全：{'、'.join(missing)}。未回退到更旧日期。")
        result = save_files(data_dir, sources, day, "共享目录下载")
        result["selection_rule"] = "早于今日的最新日期目录（依赖上游按交易日生成）"
        return result
    except OSError as exc:
        raise ReportError("无法访问共享来源目录，请连接公司网络并检查路径或 Windows 访问权限；已保存的数据仍可查询。") from exc


def catalog(settings: Settings) -> dict:
    values = settings.get()
    base = settings.data_path(values)
    dates = []
    if base.is_dir():
        for folder in base.iterdir():
            if not DATE_RE.fullmatch(folder.name) or not folder.is_dir():
                continue
            try:
                valid_date(folder.name)
            except ReportError:
                continue
            manifest = manifest_for(folder)
            files = {}
            for kind in REPORTS:
                path = folder / filename(kind, folder.name)
                if path.is_file():
                    entry = manifest.get("files", {}).get(kind, {})
                    files[kind] = {**entry, "name": path.name, "size": path.stat().st_size,
                                   "path": str(path.resolve()), "origin": entry.get("origin", "本地文件")}
            if files:
                dates.append({"date": folder.name, "complete": len(files) == len(REPORTS), "files": files})
    dates.sort(key=lambda item: item["date"], reverse=True)
    return {"today": today().isoformat(), "settings": values, "saved_dir": str(base),
            "dates": dates, "reports": REPORTS}


def report_path(settings: Settings, kind: str, day: str) -> Path:
    path = settings.data_path() / valid_date(day) / filename(kind, day)
    if not path.is_file():
        raise ReportError("本地尚未保存所选日期的报表，请同步或导入文件。")
    return path
