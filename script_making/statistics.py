"""Read-only, traceable report statistics. No recalculation or trading actions."""
from __future__ import annotations

import json
import hashlib
import re
import unicodedata
from copy import deepcopy
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from .downloads import validate_report
from .query import read_table
from .storage import ReportError, atomic_json, sha256, timestamp, today

FIELDS = ("标的代码", "标的", "总盈亏", "当日盈亏", "Delta总敞口", "Delta总敞口金额")
INDEX_CODES = {"000016.SH", "000300.SH", "000905.SH", "000852.SH"}
GROUPS = ("指数、指数ETF及股指期货", "个股和其他ETF")
# Name-based classification requested by the user, restricted to fund codes.
INDEX_ETF_NAME = re.compile(
    r"(?<!\d)(?:50|100|180|200|300|500|800|1000|2000)(?!\d)"
    r"|中证|上证|沪深|深证|国证|创业板|科创|恒生|纳指|纳斯达克|标普|道琼斯")
# Fund-manager/exchange verification is documented in README. Keep OF/SH aliases
# explicit, since replacing every .OF code with .SH would misclassify other funds.
DEFAULT_EXCLUSIONS = ["159001.SZ", "511660.OF", "511660.SH", "511850.SH", "511990.SH"]
LEGACY_DEFAULT_TEMPLATES = {
    "尊敬的领导、同事：\n\n现将本次日报统计汇报如下。\n\n"
    "一、组内收益（pnl1）\n{{pnl1}}\n\n二、风险观察\n{{risk}}\n\n" + attachment + "\n\n顺颂商祺。"
    for attachment in (
        "附件为 pnl1 和风险观察两张完整统计图片。",
        "附件为 pnl1、风险观察完整统计图片及两份日报今年的盈亏双轴图。",
        "附件为 pnl1 完整统计图片（含当年最近30份报表的收益图）及风险观察汇总图片。")
}
DEFAULT_TEMPLATE = ("尊敬的领导、同事：\n\n"
                    "【组内收益】\n{{pnl1}}\n\n【风险观察】\n{{risk}}\n\n"
                    "详情见附件 pnl1.png（含当年收益图）和 risk.png。\n\n顺颂\n  时祺")
DEFAULT_CONFIG = {"metric": "Delta总敞口金额", "threshold": "500000",
                  "exclusions": DEFAULT_EXCLUSIONS, "to": "", "cc": "",
                  "subject": "日报统计 {{dates}}", "template": DEFAULT_TEMPLATE}


def decimal_value(value, context):
    if value is None or isinstance(value, bool):
        raise ReportError(f"{context}为空或不是数值，未按零处理。")
    try:
        result = Decimal(unicodedata.normalize("NFKC", str(value)).strip().replace(",", ""))
        if not result.is_finite():
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        raise ReportError(f"{context}不是有效数值：{value}") from None


def normalize_code(value):
    return unicodedata.normalize("NFKC", str(value or "")).strip().upper()


def validated_config(values):
    result = {**deepcopy(DEFAULT_CONFIG), **values}
    if result["metric"] not in FIELDS[2:]:
        raise ReportError("请选择有效的观察阈值字段。")
    threshold = decimal_value(result["threshold"], "观察阈值")
    if threshold < 0:
        raise ReportError("绝对值观察阈值不能小于零。")
    result["threshold"] = str(threshold)
    excluded = result["exclusions"]
    if isinstance(excluded, str):
        excluded = re.split(r"[\s,，;；]+", excluded)
    if not isinstance(excluded, list):
        raise ReportError("永久剔除项应为代码列表。")
    result["exclusions"] = sorted({normalize_code(code) for code in excluded if str(code).strip()})
    for code in result["exclusions"]:
        if not re.fullmatch(r"[A-Z0-9]+\.[A-Z]+", code):
            raise ReportError(f"剔除代码格式无效：{code}，请填写完整 Wind 代码。")
    for field in ("to", "cc", "subject", "template"):
        if not isinstance(result[field], str):
            raise ReportError("邮件设置应为文本。")
    if result["template"] in LEGACY_DEFAULT_TEMPLATES:
        result["template"] = DEFAULT_TEMPLATE
    if any("\n" in result[field] or "\r" in result[field] for field in ("to", "cc", "subject")):
        raise ReportError("邮件地址和主题不能包含换行。")
    return result


def load_config(root):
    path = Path(root) / "statistics.local.json"
    if not path.exists():
        return deepcopy(DEFAULT_CONFIG)
    try:
        return validated_config(json.loads(path.read_text(encoding="utf-8")))
    except (OSError, ValueError, TypeError) as exc:
        raise ReportError("统计本机配置无法读取，请检查 statistics.local.json。") from exc


def save_config(root, values):
    config = validated_config(values)
    atomic_json(Path(root) / "statistics.local.json", config)
    return config


def pnl_value(path, day):
    validate_report(path, "pnl", date.fromisoformat(day))
    table = read_table(path, "累积盈亏统计")
    headers = [column["name"] for column in table["columns"]]
    if headers.count("交易台") != 1 or headers.count("净利润") != 1:
        raise ReportError("pnl1 缺少唯一的交易台/净利润列。")
    index = headers.index("交易台")
    from openpyxl import load_workbook
    from .query import HEADER_ROWS
    book = load_workbook(path, data_only=True, keep_links=False)
    try:
        ws = book['累积盈亏统计']
        header = HEADER_ROWS[ws.title]
        columns = {str(c.value).strip(): c.column for c in ws[header] if c.value}
        matches = []
        for line in table['source_rows']:
            cell = ws.cell(line, columns['交易台'])
            merged = next((m for m in ws.merged_cells.ranges if cell.coordinate in m), None)
            if merged:
                if merged.min_row != line:
                    continue
                cell = ws.cell(merged.min_row, merged.min_col)
            if str(cell.value or '').strip() == '境内权益互换台':
                matches.append((ws.cell(line, columns['净利润']).value, line))
        if len(matches) != 1:
            raise ReportError('pnl1 未找到唯一的境内权益互换台行。')
        value, line = matches[0]
        return decimal_value(value, '境内权益互换台净利润'), line
    finally:
        book.close()


def pnl_statistics(path, day, previous=None):
    current, line = pnl_value(path, day)
    grid = read_table(path, "累积盈亏统计", full_grid=True)
    result = {"date": day, "current": str(current), "previous_date": None,
              "previous": None, "change": None, "row": line, "grid": grid,
              "notes": [str(v) for row in grid["rows"] for v in row
                        if isinstance(v, str) and ("T-2" in v or v.startswith("注"))],
              "warning": "未取得更早报表，上期净利润和两期变动暂缺。"}
    from .pnl_groups import discover
    result["presentation"] = discover(path, grid, line)
    if previous:
        previous_path, previous_day = previous
        if previous_day >= day:
            raise ReportError("上期报表日期必须早于本期报表日期。")
        prior, _ = pnl_value(previous_path, previous_day)
        result.update(previous_date=previous_day, previous=str(prior))
        if date.fromisoformat(day).year != date.fromisoformat(previous_day).year:
            result["warning"] = "两份报表跨年，累计口径可能重置，未计算差额。"
        else:
            result["change"] = str(current - prior)
            result["warning"] = ("两期净利润变动 = 本期 − 上期。仅使用所选两份报表；"
                                 "未独立校验交易日连续性，缺日时代表区间变动。原表 T-2 注释保留，报表日不等于收益归属日。")
    return result


def classify(code, name=""):
    if code in INDEX_CODES or re.fullmatch(r"(?:IF|IH|IC|IM)\d+\.CFE", code):
        return GROUPS[0]
    if (re.fullmatch(r"(?:5\d{5}\.SH|1[56]\d{4}\.SZ)", code)
            and INDEX_ETF_NAME.search(unicodedata.normalize("NFKC", str(name or "")))):
        return GROUPS[0]
    # Recognise mainland stocks/ETFs, BSE stocks and HK stocks. Other indices,
    # bonds and open-end funds must not silently enter the equity bucket.
    if (re.fullmatch(r"[56]\d{5}\.SH", code)
            or re.fullmatch(r"[03]\d{5}\.SZ", code)
            or re.fullmatch(r"\d{6}\.BJ", code)
            or re.fullmatch(r"\d{4,5}\.HK", code)
            or re.fullmatch(r"1[56]\d{4}\.SZ", code)):
        return GROUPS[1]
    return None


def risk_statistics(path, day, kind, config):
    if kind == "strategy":
        validate_report(path, kind, date.fromisoformat(day))
    table = read_table(path, "风险盈亏统计")
    names = [c["name"] for c in table["columns"]]
    for field in FIELDS:
        if names.count(field) != 1:
            raise ReportError(f"风险盈亏统计缺少唯一的「{field}」列。")
    positions = [names.index(field) for field in FIELDS]
    metric_index = FIELDS.index(config["metric"])
    threshold = Decimal(config["threshold"])
    excluded = set(config["exclusions"])
    groups = {group: {"name": group, "rows": [], "count": 0, "total": Decimal(0),
                      "daily": Decimal(0), "net": Decimal(0), "gross": Decimal(0)} for group in GROUPS}
    skipped, unknown, seen = [], [], set()
    for source, line in zip(table["rows"], table["source_rows"]):
        row = [source[index] for index in positions]
        code = normalize_code(row[0])
        if not code:
            if any(v not in (None, "") for v in row):
                raise ReportError(f"风险表第 {line} 行缺少标的代码，请核对。")
            continue
        if code in seen:
            raise ReportError(f"标的代码 {code} 重复，未自动合并，以免重复统计。")
        seen.add(code)
        if code in excluded:
            skipped.append(code)
            continue
        metric = decimal_value(row[metric_index], f"第 {line} 行 {config['metric']}")
        if abs(metric) < threshold:
            continue
        group = classify(code, row[1])
        if group is None:
            unknown.append(code)
            continue
        values = [decimal_value(row[i], f"第 {line} 行 {FIELDS[i]}") for i in range(2, 6)]
        item = groups[group]
        item["rows"].append([code, str(row[1] or ""), *map(str, values)])
        item["count"] += 1
        item["total"] += values[0]
        item["daily"] += values[1]
        item["net"] += values[3]
        item["gross"] += abs(values[3])
    if unknown:
        raise ReportError("达到阈值的标的无法归入两类，未发布部分汇总：" + "、".join(unknown[:20]))
    for item in groups.values():
        item["rows"].sort(key=lambda row: (bool(re.fullmatch(r"(?:IF|IH|IC|IM)\d+\.CFE", row[0])), -abs(Decimal(row[metric_index]))))
        for field in ("total", "daily", "net", "gross"):
            item[field] = str(item[field])
    return {"kind": kind, "date": day, "groups": list(groups.values()), "excluded": skipped,
            "source_count": len(seen)}


def combined_risks(risks):
    """Totals within each source; never combine different reports."""
    return [{"kind": kind, "name": "观察项合计", "date": risk["date"],
             "count": sum(g["count"] for g in risk["groups"]),
             **{key: str(sum((Decimal(g[key]) for g in risk["groups"]), Decimal(0)))
                for key in ("total", "daily", "net")}}
            for kind, risk in risks.items()]


def stock_subset(risk):
    rows = [row for group in risk["groups"] for row in group["rows"]
            if re.fullmatch(r"(?:6\d{5}\.SH|[03]\d{5}\.SZ|\d{6}\.BJ|\d{4,5}\.HK)", row[0])
            and "ETF" not in unicodedata.normalize("NFKC", row[1]).upper()]
    return {"date": risk["date"], "rows": rows, "count": len(rows),
            **{key: str(sum((Decimal(row[index]) for row in rows), Decimal(0)))
               for key, index in (("total", 2), ("daily", 3), ("net", 5))}}


def build_snapshot(service, selections, values):
    config = validated_config(values)
    snapshot = {"id": uuid4().hex, "created_at": timestamp(), "config": config,
                "pnl": None, "risks": {}, "errors": {}, "sources": {}, "history": {}}

    def capture(key, kind, day):
        path = service.path_for(kind, day)
        snapshot["sources"][key] = {"path": str(path), "name": path.name, "date": day,
                                    "sha256": sha256(path)}
        return path

    for kind in ("pnl", "swap", "strategy"):
        day = selections.get(kind)
        if not day:
            snapshot["errors"][kind] = "尚未选择本地报表。"
            continue
        try:
            path = capture(kind, kind, day)
            if kind == "pnl":
                prior_day = selections.get("previous")
                prior = (capture("previous", kind, prior_day), prior_day) if prior_day else None
                snapshot["pnl"] = pnl_statistics(path, day, prior)
            else:
                snapshot["risks"][kind] = risk_statistics(path, day, kind, config)
        except (OSError, ValueError, KeyError, BadZipFile, ParseError) as exc:
            snapshot["errors"][kind] = str(exc)
    snapshot["risk_totals"] = combined_risks(snapshot["risks"])
    snapshot["swap_stocks"] = stock_subset(snapshot["risks"]["swap"]) if "swap" in snapshot["risks"] else None
    catalog = service.query_catalog()
    year = today().year
    cache = service.settings.data_path() / "statistics" / "history_cache"
    history = {"year": year, "through": selections.get("pnl"), "limit": None, "points": [], "errors": {}}
    snapshot["history"]["pnl"] = history
    if history["through"]:
        entries = [entry for entry in catalog["dates"] if "pnl" in entry["files"]
                   and f"{year}-01-01" <= entry["date"] <= min(history["through"], f"{year}-12-31")]
        prior = None
        for entry in reversed(entries):
            day = entry["date"]
            point = {"date": day, "current": None, "previous_date": None, "previous": None, "change": None}
            history["points"].append(point)
            key = f"history:pnl:{day}"
            try:
                path = capture(key, "pnl", day)
                digest = snapshot["sources"][key]["sha256"]
                cache_key = hashlib.sha256(json.dumps(["pnl1-value-v1", day, digest]).encode()).hexdigest()
                cached = cache / (cache_key + ".json")
                current = None
                if cached.exists():
                    try:
                        value = json.loads(cached.read_text(encoding="utf-8"))
                        current = decimal_value(value["current"], "pnl1 历史缓存")
                    except (OSError, ValueError, KeyError, TypeError):
                        current = None
                if current is None:
                    current, _ = pnl_value(path, day)
                    if sha256(path) != digest:
                        raise ReportError("读取期间文件更新，请重试。")
                    atomic_json(cached, {"current": str(current)})
                point["current"] = str(current)
                if prior is not None:
                    point.update(previous_date=prior["date"], previous=prior["current"],
                                 change=str(current - Decimal(prior["current"])))
                prior = point
            except (OSError, ValueError, KeyError, BadZipFile, ParseError) as exc:
                history["errors"][day] = str(exc)
                prior = None  # Never bridge a failed report and call it a one-period change.
    for source in snapshot["sources"].values():
        if sha256(Path(source["path"])) != source["sha256"]:
            raise ReportError("计算期间来源文件发生变化，请重新生成统计。")
    if not snapshot["pnl"] and not snapshot["risks"]:
        raise ReportError("；".join(snapshot["errors"].values()))
    folder = service.settings.data_path() / "statistics" / "records"
    atomic_json(folder / (snapshot["id"] + ".json"), snapshot)
    return snapshot
