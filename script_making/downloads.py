"""Independent dated report archives. No report analytics or formula calculation."""
from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from shared_data.io import file_lock
from .storage import ReportError, atomic_json, sha256, timestamp, today

DOWNLOADS = {
    "pnl": {"label": "盈亏统计", "pattern": r"盈亏统计_(\d{4}-\d{2}-\d{2})_new_adjust\.xlsx",
            "relative": "OTC-Quant/pybct/TZQ/盈亏统计_new", "sheet": "累积盈亏统计", "cell": "B3"},
    "strategy": {"label": "D1策略台日报", "pattern": r"D1策略台日报_(\d{4}-\d{2}-\d{2})\.xlsx",
                 "relative": "OTC-Quant/pybct/LH/D1策略台日报/report", "sheet": "总表", "cell": "J5"},
}


def source_path(values: dict, kind: str) -> Path:
    spec = DOWNLOADS[kind]
    explicit = values.get(kind + "_source", "").strip()
    if explicit:
        root = Path(explicit)
        # Accept a pasted year directory, but keep year rollover automatic.
        return root.parent if re.fullmatch(r"20\d{2}", root.name) else root
    existing = str(values.get("source_root", "")).replace("\\", "/")
    match = re.match(r"(?i)^(.*?/for_jrcx)(?:/|$)", existing)
    if not match:
        raise ReportError("无法从原日报来源推导服务器前缀。请在「数据与同步」填写该日报来源根目录。")
    return Path(match[1]) / spec["relative"]


def archive_path(settings, kind: str) -> Path:
    if kind not in DOWNLOADS:
        raise ReportError("未知下载来源。")
    return settings.data_path() / "downloads" / kind


def file_date(name: str, kind: str) -> date | None:
    match = re.fullmatch(DOWNLOADS[kind]["pattern"], name)
    if not match:
        return None
    try:
        return date.fromisoformat(match[1])
    except ValueError:
        return None


def read_manifest(base: Path) -> dict:
    path = base / "manifest.json"
    if not path.exists():
        return {"files": {}}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(value.get("files"), dict):
            raise ValueError()
        return value
    except (ValueError, OSError, AttributeError) as exc:
        raise ReportError("下载清单无法读取，已保留文件。请检查该来源的 manifest.json。") from exc


def local_files(base: Path, kind: str) -> list[dict]:
    result = []
    if base.is_dir():
        for folder in base.iterdir():
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", folder.name) or not folder.is_dir():
                continue
            for path in folder.glob("*.xlsx"):
                day = file_date(path.name, kind)
                if day and day.isoformat() == folder.name:
                    result.append({"date": day.isoformat(), "name": path.name,
                                   "path": str(path.resolve()), "size": path.stat().st_size})
    return sorted(result, key=lambda item: (item["date"], item["name"]), reverse=True)


def validate_report(path: Path, kind: str, day: date) -> None:
    try:
        if not 0 < path.stat().st_size <= 150 * 1024 * 1024:
            raise ReportError("日报文件应大于 0 且不超过 150 MB。")
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 1024**3 or archive.testzip():
                raise ReportError("日报文件损坏或解压后超过 1 GB。")
        spec = DOWNLOADS[kind]
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if spec["sheet"] not in book.sheetnames:
                raise ReportError(f"缺少日期核验工作表「{spec['sheet']}」。")
            value = book[spec["sheet"]][spec["cell"]].value
            if isinstance(value, datetime):
                value = value.date()
            if isinstance(value, str):
                text = value.strip().replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-")
                match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
                value = date(*map(int, match.groups())) if match else None
            if value != day:
                raise ReportError(f"{path.name} 的表内报表日期与文件名不一致或无法识别。")
        finally:
            book.close()
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        if isinstance(exc, ReportError):
            raise
        raise ReportError(f"无法校验日报 {path.name}（{type(exc).__name__}）。") from exc


def _folder_parts(name: str, parent: tuple[int, ...]) -> tuple[int, ...] | None:
    """Only traverse date partitions, never arbitrary trees on the server."""
    normalized = name.replace("年", "-").replace("月", "-").replace("日", "").rstrip("-")
    if re.fullmatch(r"20\d{2}(?:-\d{1,2}){0,2}", normalized):
        parts = tuple(map(int, normalized.split("-")))
    elif re.fullmatch(r"20\d{4}(?:\d{2})?", normalized):
        parts = tuple(int(normalized[i:j]) for i, j in ((0, 4), (4, 6), (6, 8)) if i < len(normalized))
    elif parent and len(parent) < 3 and re.fullmatch(r"\d{1,2}", normalized):
        parts = (*parent, int(normalized))
    else:
        return None
    if parent and parts[:len(parent)] != parent:
        return None
    try:
        date(*(parts + (1,) * (3 - len(parts))))
    except ValueError:
        return None
    return parts


def discover(root: Path, kind: str, start: date, cutoff: date) -> tuple[dict, list[str]]:
    candidates, errors = {}, []
    def visit(folder, parts=()):
        for path in sorted(folder.iterdir()):
            day = file_date(path.name, kind)
            if day and start <= day < cutoff and path.is_file():
                if parts and (day.year, day.month, day.day)[:len(parts)] != parts:
                    errors.append(f"{path.name} 与来源日期目录不一致，未下载。")
                    continue
                candidates.setdefault(day, []).append(path)
            elif len(parts) < 3:
                child_parts = _folder_parts(path.name, parts)
                if child_parts and path.is_dir():
                    lo = (start.year, start.month, start.day)[:len(child_parts)]
                    hi = (cutoff.year, cutoff.month, cutoff.day)[:len(child_parts)]
                    if lo <= child_parts <= hi:
                        visit(path, child_parts)
    visit(root)
    return candidates, errors


def sync_download(settings, kind: str, as_of: date | None = None) -> dict:
    cutoff = as_of or today()
    root = source_path(settings.get(), kind)
    base = archive_path(settings, kind)
    base.mkdir(parents=True, exist_ok=True)
    try:
        with file_lock(base / ".writer.lock"):
            manifest = read_manifest(base)
            # Advance only from verified local files; a failed/empty first check
            # never marks initialization complete or skips its seven-day window.
            verified = []
            damaged = []
            for name, entry in manifest["files"].items():
                day = file_date(name, kind)
                if not day or day >= cutoff:
                    continue
                path = base / day.isoformat() / name
                if path.is_file() and sha256(path) == entry.get("sha256"):
                    verified.append(day)
                else:
                    damaged.append(day)
            pending = set()
            for value in manifest.get("pending_dates", []):
                try:
                    pending_day = date.fromisoformat(value)
                    if pending_day < cutoff:
                        pending.add(pending_day)
                except (TypeError, ValueError):
                    continue
            recent = cutoff - timedelta(days=7)
            start = min([recent, *damaged, *pending, *([max(verified)] if verified else [])])
            candidates, errors = discover(root, kind, start, cutoff)
            pending.update(candidates)
            pending.update(damaged)
            manifest["pending_dates"] = sorted(day.isoformat() for day in pending)
            atomic_json(base / "manifest.json", manifest)
            saved, reused = [], []
            for day, sources in sorted(candidates.items()):
                try:
                    with tempfile.TemporaryDirectory(prefix=".incoming-", dir=base) as temporary:
                        staged = Path(temporary) / sources[0].name
                        source = sources[0]
                        before = source.stat()
                        destination = base / day.isoformat() / staged.name
                        if destination.is_file() and sha256(destination) == sha256(source):
                            checked = destination
                        else:
                            shutil.copyfile(source, staged)
                            checked = staged
                        digest = sha256(checked)
                        for duplicate in sources:
                            first = duplicate.stat()
                            source_digest = sha256(duplicate)
                            last = duplicate.stat()
                            if (first.st_size, first.st_mtime_ns) != (last.st_size, last.st_mtime_ns):
                                raise ReportError("来源文件仍在生成，请稍后重试。")
                            if source_digest != digest:
                                raise ReportError("多个来源位置的同日期文件内容不同，未覆盖本地。")
                        after = source.stat()
                        if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns) or checked.stat().st_size != before.st_size:
                            raise ReportError("来源文件仍在生成或复制不完整，请稍后重试。")
                        validate_report(checked, kind, day)
                        existed = destination.exists()
                        if existed and sha256(destination) != digest:
                            raise ReportError("已有同名不同内容的本地文件，已保留原文件。请核对后移走冲突文件再同步。")
                        destination.parent.mkdir(parents=True, exist_ok=True)
                        if not existed:
                            os.replace(staged, destination)
                        if sha256(destination) != digest:
                            raise ReportError("保存后文件哈希核验失败。")
                        previous = manifest["files"].get(destination.name, {})
                        manifest["files"][destination.name] = {
                            "date": day.isoformat(), "path": f"{day.isoformat()}/{destination.name}",
                            "source": str(source), "size": destination.stat().st_size, "sha256": digest,
                            "saved_at": previous.get("saved_at", timestamp()), "checked_at": timestamp()}
                        pending.discard(day)
                        manifest["pending_dates"] = sorted(value.isoformat() for value in pending)
                        # Commit each successful file so a later failure is retryable.
                        atomic_json(base / "manifest.json", manifest)
                        (reused if existed else saved).append(day.isoformat())
                except (ReportError, OSError) as exc:
                    errors.append(f"{day}: {exc}")
            for day in sorted(pending - candidates.keys()):
                errors.append(f"{day}: 待补文件在来源检查范围内尚未找到，将继续重试。")
            result = {"saved": saved, "reused": reused, "errors": errors,
                      "source_latest": max(candidates).isoformat() if candidates else None,
                      "checked_at": timestamp(), "window_start": start.isoformat(),
                      "message": "检查完成" if candidates else "检查范围内没有可下载的日报；已保留本地文件。"}
            manifest["last_check"] = result
            atomic_json(base / "manifest.json", manifest)
            return result
    except (OSError, TimeoutError) as exc:
        raise ReportError("无法访问日报来源或保存目录，请检查公司网络及目录权限；已保存文件仍可使用。") from exc
