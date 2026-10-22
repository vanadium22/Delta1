"""Five-to-three sector rollups and the isolated self-only history task."""
from contextlib import redirect_stdout
from datetime import date, timedelta
from decimal import Decimal
from email import policy
from email.parser import BytesParser
import json
import io
import os
from pathlib import Path
import subprocess
import sys
import threading
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from shared_data.io import file_lock
from script_making.downloads import archive_path
from script_making.pnl_history import build_history, self_config, mail_manifest
from script_making.pnl_history_job import run_once, main, history_diagnostics
from script_making.statistics import pnl_statistics
from script_making.statistics_export import render_pages
from script_making.storage import ReportError, atomic_json, sha256
from script_making.tests import test_statistics as samples


def sector_book(path, day, scale=1, vertical=False):
    book = Workbook()
    ws = book.active
    ws.title = '累积盈亏统计'
    ws['B3'] = day
    for c, title in enumerate(('交易台', '行业', '净利润', '当日收入'), 1):
        ws.cell(5, c, title)
    values = {6: ('境内权益互换台', 100), 7: ('原始明细一', 10), 9: ('原始明细三', 30),
              10: ('原始明细四', 40), 11: ('境内个股期权台', 62),
              12: ('TMT', 10), 13: ('消费', 20), 14: ('周期-消费代管', -5),
              15: ('周期', 30), 16: ('周期-其他', 7),
              17: ('用于验证长名称换行的其他业务汇总组', -25), 18: ('场外业务合计', 137)}
    for row, (name, value) in values.items():
        ws.cell(row, 2 if vertical and 12 <= row <= 16 else 1, name)
        ws.cell(row, 3, value * scale)
        ws.cell(row, 4, value * scale / 10)
    for row in (6, 11, 17, 18):
        ws.cell(row, 1).font = Font(bold=True)
    for row in range(12, 17):
        cell = ws.cell(row, 2 if vertical else 1)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(indent=1)
    if vertical:
        ws.merge_cells('A11:A16')
    ws['A20'] = '注：境内权益互换台取T-2日数据'
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()
    return path


class PnlHistoryTests(unittest.TestCase):
    setUpClass = classmethod(samples.StatisticsTests.setUpClass.__func__)
    setUp = samples.StatisticsTests.setUp
    tearDown = samples.StatisticsTests.tearDown
    snapshot = samples.StatisticsTests.snapshot

    def prepare(self):
        sector_book(self.previous_path, '2026-10-16', 1)
        sector_book(self.pnl_path, '2026-10-19', 2)
        atomic_json(self.service.settings.root / 'pnl_history.local.json', {'to': 'self@example.invalid'})

    def test_five_to_three_and_parent_removed_with_original_provenance(self):
        for vertical in (False, True):
            sector_book(self.pnl_path, '2026-10-19', vertical=vertical)
            before = sha256(self.pnl_path)
            view = pnl_statistics(self.pnl_path, '2026-10-19')['presentation']
            self.assertEqual(view['diagnostics'], [])
            self.assertEqual([g['name'] for g in view['groups']],
                             ['境内权益互换台', 'TMT', '消费', '周期', '用于验证长名称换行的其他业务汇总组'])
            self.assertEqual([g['value'] for g in view['groups']], ['100', '10', '15', '37', '-25'])
            self.assertEqual(view['rows'], [6, 7, 8, 9, 10, 12, 13, 15, 17, 18])
            self.assertEqual(view['row_sources'][6:8], [[13, 14], [15, 16]])
            self.assertEqual(view['values'][2], [None] * 4)
            self.assertEqual(sha256(self.pnl_path), before)
            html = render_pages(self.snapshot())[0]
            self.assertIn('13 + 14', html)
            self.assertIn('15 + 16', html)
            self.assertNotIn('<b>境内个股期权台</b>', html)

    def test_missing_duplicate_unrecognized_and_inconsistent_sectors_fail(self):
        for mode in ('missing', 'duplicate', 'cache', 'total'):
            sector_book(self.pnl_path, '2026-10-19')
            book = load_workbook(self.pnl_path)
            ws = book.active
            if mode == 'missing': ws['A14'] = '未知行业'
            if mode == 'duplicate': ws['A14'] = '消费'
            if mode == 'cache': ws['C14'] = '=SUM(C1:C2)'
            if mode == 'total': ws['C11'] = 999
            book.save(self.pnl_path)
            book.close()
            with self.subTest(mode=mode), self.assertRaises(ReportError):
                pnl_statistics(self.pnl_path, '2026-10-19')

    def test_otc_and_group_history_gaps_full_year_and_next_year(self):
        self.prepare()
        for i in range(40):
            day = (date(2026, 1, 1) + timedelta(days=i)).isoformat()
            sector_book(archive_path(self.service.settings, 'pnl') / day / f'盈亏统计_{day}_new_adjust.xlsx', day)
        snap = build_history(self.service, '2026-10-19', {'to': 'self@example.invalid'})
        self.assertEqual(len(snap['otc']), 42)
        self.assertEqual(snap['otc'][-1]['current'], '274')
        self.assertEqual(snap['otc'][-1]['change'], '137')
        consumer = next(g for g in snap['groups'] if g['name'] == '消费')
        self.assertEqual(consumer['points'][-1]['change'], '15')
        self.assertEqual(consumer['points'][-1]['source_rows'], [13, 14])
        self.previous_path.write_bytes(b'bad')
        broken = build_history(self.service, '2026-10-19', {'to': 'self@example.invalid'})
        self.assertIn('2026-10-16', broken['errors'])
        self.assertIsNone(broken['otc'][-2]['current'])
        self.assertIsNone(broken['otc'][-1]['change'])
        for day in ('2027-01-04', '2027-01-05'):
            sector_book(archive_path(self.service.settings, 'pnl') / day / f'盈亏统计_{day}_new_adjust.xlsx', day)
        with patch('script_making.pnl_history.today', return_value=date(2027, 1, 6)):
            new = build_history(self.service, '2027-01-05', {'to': 'self@example.invalid'})
        self.assertEqual([p['date'] for p in new['otc']], ['2027-01-04', '2027-01-05'])
        self.assertIsNone(new['otc'][0]['change'])

    def test_missing_group_is_not_zero_or_bridged(self):
        self.prepare()
        middle = archive_path(self.service.settings, 'pnl') / '2026-10-18/盈亏统计_2026-10-18_new_adjust.xlsx'
        sector_book(middle, '2026-10-18')
        book = load_workbook(middle)
        ws = book.active
        ws['A17'], ws['C17'], ws['D17'] = None, None, None
        ws['C18'] = 162
        book.save(middle); book.close()
        snap = build_history(self.service, '2026-10-19', {'to': 'self@example.invalid'})
        other = snap['groups'][-1]['points']
        self.assertIsNone(other[1]['current'])
        self.assertIsNone(other[2]['change'])
        self.assertEqual(snap['errors'], {})

    def test_export_self_only_mime_reuse_and_revision_preserve_old(self):
        self.prepare()
        original = self.service.settings.root / 'statistics.local.json'
        atomic_json(original, {'to': 'team@example.invalid', 'cc': 'manager@example.invalid'})
        main_config_hash = sha256(original)
        code, first = run_once(self.service.settings.root, date(2026, 10, 21), local_only=True)
        self.assertEqual(code, 0, first)
        folder = Path(first['export_directory'])
        self.assertEqual({p.name for p in folder.glob('*.png')}, {'otc_pnl.png', 'groups_pnl.png'})
        manifest = mail_manifest(folder)
        self.assertEqual(manifest['mail']['to'], ['self@example.invalid'])
        self.assertEqual(manifest['mail']['cc'], [])
        self.assertEqual(manifest['mail']['bcc'], [])
        msg = BytesParser(policy=policy.default).parsebytes((folder / '邮件草稿.eml').read_bytes())
        self.assertEqual(len([p for p in msg.walk() if p.get_filename()]), 2)
        self.assertIsNotNone(msg.get_body(('plain',)))
        self.assertIsNotNone(msg.get_body(('html',)))
        code, reused = run_once(self.service.settings.root, date(2026, 10, 21), local_only=True)
        self.assertEqual(reused['state'], 'reused')
        self.assertEqual(reused['export_directory'], str(folder))
        sector_book(self.pnl_path, '2026-10-19', 3)
        code, revised = run_once(self.service.settings.root, date(2026, 10, 21), local_only=True)
        self.assertEqual(code, 0, revised)
        self.assertNotEqual(revised['export_directory'], str(folder))
        self.assertTrue((folder / '邮件草稿.eml').is_file())
        self.assertEqual(sha256(original), main_config_hash)
        (Path(revised['export_directory']) / 'otc_pnl.png').write_bytes(b'changed')
        with self.assertRaises(ReportError): mail_manifest(revised['export_directory'])

    def test_failure_incomplete_busy_and_no_sendable_manifest(self):
        self.prepare()
        with patch('script_making.daily_job.sync_once', return_value={'errors': ['synthetic']}):
            code, record = run_once(self.service.settings.root, date(2026, 10, 21))
        self.assertEqual(code, 2)
        self.assertNotIn('export_directory', record)
        with patch('script_making.daily_job.sync_once', side_effect=OSError('synthetic source offline')):
            code, record = run_once(self.service.settings.root, date(2026, 10, 21))
        self.assertEqual(code, 2)
        self.assertNotIn('export_directory', record)
        ready, release = threading.Event(), threading.Event()
        def hold():
            with file_lock(self.service.settings.data_path() / 'statistics/.daily_job.lock'):
                ready.set(); release.wait(10)
        thread = threading.Thread(target=hold)
        thread.start(); self.assertTrue(ready.wait(5))
        try:
            code, _ = run_once(self.service.settings.root, date(2026, 10, 21), local_only=True)
            self.assertEqual(code, 3)
        finally:
            release.set(); thread.join()
        with patch('script_making.statistics_export.save_image', side_effect=OSError('synthetic render failure')):
            code, record = run_once(self.service.settings.root, date(2026, 10, 21), local_only=True)
        self.assertEqual(code, 1, record)
        self.assertFalse(list((self.service.settings.data_path() / 'statistics/pnl_history_exports').rglob('*.eml')))
        self.previous_path.unlink()
        code, record = run_once(self.service.settings.root, date(2026, 10, 21), local_only=True)
        self.assertEqual(code, 2, record)

    def test_single_self_address_and_direct_work_automation_cli(self):
        for bad in ({'to': 'a@example.invalid,b@example.invalid'}, {'to': 'a@example.invalid', 'cc': ''},
                    {'to': ''}, {'to': 'Name <a@example.invalid>'}, {'to': 'a@example.invalid\nBcc: b@example.invalid'}):
            with self.assertRaises(ReportError): self_config(bad)
        self.prepare()
        output = self.root / 'wa-result.json'
        program = Path(__file__).resolve().parents[1] / 'pnl_history_job.py'
        result = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(program), '--root',
                                 str(self.service.settings.root), '--run', '--local-only', '--as-of', '2026-10-21',
                                 '--wa-output', str(output)], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        value = json.loads(output.read_text(encoding='utf-8'))
        self.assertEqual(value['mail']['to'], ['self@example.invalid'])
        self.assertEqual(len(value['mail']['attachments']), 2)
        fresh = self.root / 'never-created'
        self.assertEqual(main(['--root', str(fresh), '--inspect']), 0)
        self.assertFalse(fresh.exists())

    def test_frozen_catalog_paths_and_changed_sources_are_rejected(self):
        self.prepare()
        with patch.object(self.service, 'path_for', side_effect=AssertionError('must use frozen catalog paths')):
            snap = build_history(self.service, '2026-10-19', {'to': 'self@example.invalid'})
        self.assertEqual(snap['errors'], {})
        real_sha, counts = sha256, {}
        def changed(path):
            path = str(path)
            counts[path] = counts.get(path, 0) + 1
            return real_sha(Path(path)) if counts[path] == 1 else 'changed'
        with patch('script_making.pnl_history.sha256', side_effect=changed), self.assertRaises(ReportError):
            build_history(self.service, '2026-10-19', {'to': 'self@example.invalid'})

    def test_only_pnl_sync_zero_values_and_incomplete_has_no_manifest(self):
        self.prepare()
        sector_book(self.previous_path, '2026-10-16', 0)
        sector_book(self.pnl_path, '2026-10-19', 0)
        later = archive_path(self.service.settings, 'pnl') / '2026-10-20/盈亏统计_2026-10-20_new_adjust.xlsx'
        sector_book(later, '2026-10-20', 9)
        def sync(root, as_of, kind):
            self.assertEqual(kind, 'pnl')
            frozen = json.loads((root / 'config.local.json').read_text(encoding='utf-8'))
            self.assertEqual(Path(frozen['data_dir']), self.service.settings.data_path())
            return {'source_latest': '2026-10-19', 'errors': []}
        with patch('script_making.daily_job.sync_once', side_effect=sync) as mocked:
            code, record = run_once(self.service.settings.root, date(2026, 10, 21))
        self.assertEqual(code, 0, record)
        self.assertEqual(mocked.call_count, 1)
        snapshot = json.loads(Path(record['statistics_record']).read_text(encoding='utf-8'))
        self.assertEqual(snapshot['through'], '2026-10-19')
        self.assertEqual(snapshot['otc'][-1]['current'], '0')
        self.assertEqual(snapshot['otc'][-1]['change'], '0')
        self.previous_path.unlink()
        output = self.root / 'must-not-exist.json'
        self.assertEqual(main(['--root', str(self.service.settings.root), '--run', '--local-only',
                               '--as-of', '2026-10-20', '--wa-output', str(output)]), 2)
        self.assertFalse(output.exists())

    def test_incomplete_cli_reports_actual_dates_and_errors_without_manifest(self):
        self.prepare()
        self.previous_path.write_bytes(b'bad')
        output, stdout = self.root / 'must-not-exist.json', io.StringIO()
        with redirect_stdout(stdout):
            code = main(['--root', str(self.service.settings.root), '--run', '--local-only',
                         '--as-of', '2026-10-21', '--wa-output', str(output)])
        record = json.loads(stdout.getvalue())
        self.assertEqual(code, 2)
        self.assertEqual(record['state'], 'incomplete')
        diagnostics = record['diagnostics']
        snapshot = json.loads(Path(record['statistics_record']).read_text(encoding='utf-8'))
        self.assertEqual(diagnostics['errors'], snapshot['errors'])
        self.assertEqual(diagnostics['error_count'], 1)
        self.assertIn('2026-10-16', record['reason'])
        self.assertIn(snapshot['errors']['2026-10-16'], record['reason'])
        self.assertFalse(diagnostics['recent_otc'][0]['has_current'])
        self.assertTrue(diagnostics['recent_otc'][1]['has_current'])
        self.assertFalse(diagnostics['recent_otc'][1]['has_change'])
        self.assertNotIn('export_directory', record)
        self.assertFalse(output.exists())
        self.assertFalse(list(self.service.settings.data_path().rglob('*.eml')))

    def test_existing_record_diagnostic_cli_is_read_only_and_omits_mail_and_amounts(self):
        self.prepare()
        snapshot = build_history(self.service, '2026-10-19', {'to': 'self@example.invalid'})
        path = self.root / 'saved-record.json'
        atomic_json(path, snapshot)
        digest, files = sha256(path), set(self.root.rglob('*'))
        fresh = self.root / 'never-created'
        stdout = io.StringIO()
        with patch('script_making.storage.Settings', side_effect=AssertionError('must not load settings')), \
                patch('script_making.daily_job.sync_once', side_effect=AssertionError('must not sync')), \
                patch('script_making.pnl_history.build_history', side_effect=AssertionError('must not rebuild')), \
                redirect_stdout(stdout):
            code = main(['--root', str(fresh), '--diagnose-record', str(path)])
        result = json.loads(stdout.getvalue())
        self.assertEqual(code, 0)
        self.assertEqual(result['state'], 'diagnostic_only')
        self.assertFalse(result['send_email'])
        self.assertTrue(result['diagnostics']['complete'])
        self.assertNotIn('self@example.invalid', stdout.getvalue())
        self.assertEqual(set(result['diagnostics']['recent_otc'][-1]),
                         {'date', 'previous_date', 'has_current', 'has_change'})
        self.assertEqual(sha256(path), digest)
        self.assertEqual(set(self.root.rglob('*')), files)
        self.assertFalse(fresh.exists())
        # A record without an earlier interval remains incomplete; actual zero is present.
        snapshot['otc'] = [{'date': '2026-10-19', 'current': '0', 'change': None}]
        snapshot['reports'] = snapshot['reports'][-1:]
        diagnostics, reason = history_diagnostics(snapshot)
        self.assertFalse(diagnostics['complete'])
        self.assertTrue(diagnostics['recent_otc'][0]['has_current'])
        self.assertIn('仅取得1份', reason)
        snapshot['otc'][0]['change'] = '0'
        self.assertTrue(history_diagnostics(snapshot)[0]['complete'])

    def test_diagnostic_cli_rejects_wrong_or_missing_record_without_creating_output(self):
        path = self.root / 'invalid-record.json'
        atomic_json(path, {'kind': 'daily_report', 'errors': {}, 'otc': [], 'reports': []})
        fresh = self.root / 'never-created'
        for record in (path, self.root / 'missing.json'):
            with self.subTest(record=record), redirect_stdout(io.StringIO()) as stdout:
                self.assertEqual(main(['--root', str(fresh), '--diagnose-record', str(record)]), 1)
            result = json.loads(stdout.getvalue())
            self.assertEqual(result['state'], 'failed')
            self.assertFalse(result['send_email'])
            self.assertNotIn('export_directory', result)
        self.assertFalse(fresh.exists())
