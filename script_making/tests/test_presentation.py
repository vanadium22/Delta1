import unittest
from copy import copy
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from script_making.tests import test_statistics as samples
from script_making.tests.test_statistics import risk_book
from script_making.statistics import pnl_statistics, risk_statistics, stock_subset
from script_making.statistics_export import render_pages, export_bundle


def styled(path, day='2026-10-19'):
    book=Workbook(); ws=book.active; ws.title='累积盈亏统计'
    ws['B3']=day
    for col,value in enumerate(['交易台','说明','净利润'],1): ws.cell(5,col,value)
    rows={6:('境内权益互换台',-30000),7:('明细一',100),9:('明细三',200),10:('明细四',300),11:('一个非常长的合成业务组名称用于验证图片换行不会裁切',50000),12:('普通明细',20),13:('其他亏损组',-90000),14:('场外业务合计',-70000)}
    for r,(name,value) in rows.items():
        ws.cell(r,1,name); ws.cell(r,3,value)
    for r in (6,11,13,14): ws.cell(r,1).font=Font(bold=True)
    ws['C12'].font=Font(bold=True)
    ws.merge_cells('A11:B11')
    ws['A16']='注：境内权益互换台取T-2日数据'; ws['A16'].font=Font(bold=True)
    path.parent.mkdir(parents=True,exist_ok=True); book.save(path); book.close()

class PresentationTests(unittest.TestCase):
    setUp=samples.StatisticsTests.setUp
    tearDown=samples.StatisticsTests.tearDown
    snapshot=samples.StatisticsTests.snapshot
    setUpClass=classmethod(samples.StatisticsTests.setUpClass.__func__)

    def test_original_styles_merge_blank_and_signed_ranking(self):
        styled(self.pnl_path)
        pnl=pnl_statistics(self.pnl_path,'2026-10-19')
        view=pnl['presentation']
        self.assertEqual(view['diagnostics'],[])
        self.assertEqual(view['rows'],[6,7,8,9,10,11,13,14])
        self.assertEqual([g['row'] for g in view['ranking']],[11,6,13])
        self.assertEqual(view['values'][2],[None,None,None])
        snap=self.snapshot(); html, risk_html=render_pages(snap)
        summary_title = '<h2>风险观察 · 分日报汇总</h2>'
        self.assertLess(html.index('alt="pnl1 当年累计报表收益图"'), html.index(summary_title))
        self.assertLess(html.index(summary_title), html.index('<h2>pnl1 · 组汇总及重点明细</h2>'))
        summary_table = html.split(summary_title, 1)[1].split('</table>', 1)[0]
        self.assertEqual(summary_table.count('<tr>'), 7)  # Header plus the original six summary rows.
        self.assertEqual(summary_table.count('观察项合计'), 2)
        self.assertIn('2026-10-14', summary_table)
        self.assertIn('2026-10-19', summary_table)
        self.assertIn('净敞口金额 Σx', summary_table)
        self.assertNotIn(summary_title, risk_html)
        self.assertNotIn('净敞口金额 Σx', risk_html)
        self.assertEqual(risk_html.count('—'), 4)  # All four detail totals remain.
        self.assertNotIn('上期累计净利润',html)
        self.assertNotIn('普通明细',html)
        self.assertEqual(html.count('场外业务合计'),1)
        self.assertNotIn('各组累计净利润排名',html)
        self.assertEqual(html.count('data:image/png;base64,'), 1)

    def test_ambiguous_hierarchy_and_missing_cache_stop_ranking(self):
        for mode in ('indent','merge','missing','duplicate'):
            styled(self.pnl_path); book=load_workbook(self.pnl_path); ws=book.active
            if mode=='indent': ws['A13'].alignment=Alignment(indent=1)
            if mode=='merge': ws.merge_cells('A13:A14')
            if mode=='missing': ws['C13']='=SUM(C1:C2)'
            if mode=='duplicate': ws['A13']=ws['A11'].value
            book.save(self.pnl_path); book.close()
            view=pnl_statistics(self.pnl_path,'2026-10-19')['presentation']
            self.assertTrue(view['diagnostics'],mode); self.assertEqual(view['ranking'],[])

    def test_futures_last_stock_subset_and_empty_totals(self):
        rows=[['IF2612.CFE','股指',1,2,3,9000000],['000300.SH','指数',4,5,6,-500000],['600001.SH','股票',7,8,9,500000],['512200.SH','其他ETF',10,11,12,800000]]
        risk_book(self.swap_path,'2026-10-14',rows)
        snap=self.snapshot(); risk=snap['risks']['swap']
        self.assertEqual([r[0] for r in risk['groups'][0]['rows']],['000300.SH','IF2612.CFE'])
        self.assertEqual(snap['swap_stocks']['count'],1)
        self.assertEqual(snap['swap_stocks']['net'],'500000')
        self.assertEqual(render_pages(snap)[1].count('—'),4)
        risk_book(self.swap_path,'2026-10-14',[])
        snap=self.snapshot(); self.assertEqual(snap['swap_stocks']['net'],'0')
        self.assertIn('无达到阈值',render_pages(snap)[1])
        self.swap_path.write_bytes(b'bad')
        snap=self.snapshot(); self.assertIsNone(snap['swap_stocks'])
        with self.assertRaises(ValueError): export_bundle(self.root/'exports',snap,render_pages(snap),self.config)
