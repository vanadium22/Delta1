"""Bounded Excel reading, Unicode abbreviation matching and tolerant filters."""
from __future__ import annotations

import math
import posixpath
import re
import unicodedata
from datetime import date, datetime, time
from functools import lru_cache
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter

from .storage import ReportError

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
HEADER_ROWS = {"存续合约估值": 2, "合约交易记录": 2, "场内交易流水": 2,
               "风险盈亏统计": 6, "风险盈亏统计汇总": 6, "DMA盈亏统计": 6,
               "簿记要素": 1, "汇总": 1, "累积盈亏统计": 5,
               "每日盈亏统计": 2, "合约信息": 2}


def normalized(value) -> str:
    return unicodedata.normalize("NFKC", "" if value is None else str(value)).casefold().strip()


def fuzzy_match(value, query: str) -> bool:
    """Slash/comma = OR; spaces = AND; each term allows ordered abbreviations."""
    text = normalized(value)
    alternatives = re.split(r"[/|,，；;]+", normalized(query))
    alternatives = [alt.strip() for alt in alternatives if alt.strip()]
    if not alternatives:
        return True
    def term_matches(term):
        if term in text:
            return True
        iterator = iter(text)
        return all(any(char == candidate for candidate in iterator) for char in term)
    return any(all(term_matches(term) for term in alt.split()) for alt in alternatives)


def number(value):
    if value is None or isinstance(value, (bool, date, datetime, time)):
        return None
    if isinstance(value, str):
        text = unicodedata.normalize("NFKC", value).strip().replace(",", "")
        if not text:
            return None
        percentage = text.endswith("%")
        text = text[:-1] if percentage else text
    else:
        text, percentage = value, False
    try:
        result = float(text)
        if percentage:
            result /= 100
        return result if math.isfinite(result) else None
    except (ValueError, TypeError, OverflowError):
        return None


def parse_epsilon(value) -> float:
    result = number(value)
    if result is None or result < 0:
        raise ReportError("ε 必须是大于或等于 0 的有限数值。")
    return result


def is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def serial_value(value):
    if isinstance(value, (datetime, date)):
        if isinstance(value, datetime) and value.time() != time(0):
            return value.isoformat(sep=" ")
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def sheet_xml(path: Path, sheet: str) -> str:
    with ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        mapping = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
        for entry in workbook.find(NS + "sheets"):
            if entry.attrib["name"] == sheet:
                target = mapping[entry.attrib[REL]]
                return target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    raise ReportError("所选工作表不存在。")


def data_bounds(path: Path, sheet: str) -> tuple[int, int]:
    """Ignore formatting-only rows, including Excel's 1,048,576-row dimension."""
    last_row = last_col = 0
    with ZipFile(path) as archive, archive.open(sheet_xml(path, sheet)) as stream:
        for _, element in ET.iterparse(stream, events=("end",)):
            if element.tag == NS + "row":
                for cell in element:
                    if cell.tag != NS + "c":
                        continue
                    if any(cell.find(NS + tag) is not None for tag in ("v", "is", "f")):
                        ref = cell.attrib.get("r", "")
                        match = re.fullmatch(r"([A-Z]+)(\d+)", ref)
                        if match:
                            last_col = max(last_col, column_index_from_string(match[1]))
                            last_row = max(last_row, int(match[2]))
                element.clear()
    if last_row > 250_000 or last_col > 512:
        raise ReportError("工作表实际数据范围过大（超过 25 万行或 512 列），请先整理后读取。")
    return last_row, last_col


@lru_cache(maxsize=8)
def _sheet_names(path: str, size: int, modified: int) -> tuple:
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        return tuple(book.sheetnames)
    finally:
        book.close()


def sheet_names(path: Path) -> list:
    info = path.stat()
    return list(_sheet_names(str(path), info.st_size, info.st_mtime_ns))


@lru_cache(maxsize=4)
def _read_table(path_text: str, size: int, modified: int, sheet: str, full_grid=False) -> dict:
    path = Path(path_text)
    last_row, last_col = data_bounds(path, sheet)
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = book[sheet]
        header_row = HEADER_ROWS.get(sheet)
        if header_row is None:
            for row_number, cells in enumerate(ws.iter_rows(min_row=1, max_row=min(15, last_row), max_col=last_col), 1):
                populated = [c for c in cells if not is_blank(c.value)]
                if len(populated) >= 3 and all(isinstance(c.value, str) for c in populated):
                    header_row = row_number
                    break
        # Summary sheets are read as a cell grid, retaining every column and row
        # position rather than guessing a header and dropping source content.
        grid = full_grid or header_row is None or sheet in ("收益总览", "总表")
        if grid:
            header_row = 0
            selected = list(range(last_col))
            columns = [{"id": f"c{i}", "name": get_column_letter(i + 1), "type": "text"}
                       for i in selected]
        else:
            headers = next(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=last_col))
            selected = [i for i, cell in enumerate(headers) if not is_blank(cell.value)]
            columns = [{"id": f"c{i}", "name": str(headers[i].value).strip(), "type": "text"}
                       for i in selected]
        raw_rows = []
        source_rows = []
        number_formats = {}
        for row_number, cells in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=last_row,
                                                       max_col=last_col), header_row + 1):
            row = [cells[i].value for i in selected]
            if all(is_blank(v) for v in row):
                continue
            raw_rows.append(row)
            source_rows.append(row_number)
            for idx, value in enumerate(row):
                if not is_blank(value) and idx not in number_formats:
                    number_formats[idx] = cells[selected[idx]].number_format
        for idx, column in enumerate(columns):
            values = [row[idx] for row in raw_rows if not is_blank(row[idx])]
            identifier = column["name"].endswith(("代码", "编号", "账号", "帐号")) or column["name"] in ("标的", "操作人员", "开仓操作人员")
            if values and all(isinstance(v, (date, datetime)) for v in values):
                column["type"] = "date"
            elif values and not identifier and all(number(v) is not None for v in values):
                column["type"] = "number"
                numeric = [number(v) for v in values]
                column["min"], column["max"] = min(numeric), max(numeric)
                column["percent"] = "%" in number_formats.get(idx, "") or all(
                    isinstance(v, str) and v.strip().endswith("%") for v in values)
            column["missing"] = len(raw_rows) - len(values)
        rows = [[serial_value(v) for v in row] for row in raw_rows]
        return {"columns": columns, "rows": rows, "source_rows": source_rows,
                "sheet": sheet, "header_row": header_row, "layout": "grid" if grid else "table"}
    finally:
        book.close()


def read_table(path: Path, sheet: str, *, full_grid=False) -> dict:
    if sheet not in sheet_names(path):
        raise ReportError("所选工作表不存在。")
    info = path.stat()
    return _read_table(str(path), info.st_size, info.st_mtime_ns, sheet, full_grid)


def build_predicate(column: dict, condition: dict, epsilon: float):
    op = condition.get("op", "fuzzy")
    raw = condition.get("value", "")
    if op == "blank":
        return is_blank
    if op == "not_blank":
        return lambda x: not is_blank(x)
    if column["type"] == "number":
        allowed = {"range", "gt", "gte", "lt", "lte", "eq", "ne", "zero", "nonzero"}
        if op not in allowed:
            raise ReportError("数值列的筛选方式无效。")
        if op in ("zero", "nonzero"):
            return lambda x: number(x) is not None and (
                abs(number(x)) <= epsilon if op == "zero" else abs(number(x)) > epsilon)
        lower = upper = None
        if op == "range":
            if not is_blank(condition.get("min")):
                lower = number(condition["min"])
                if lower is None:
                    raise ReportError("范围下限不是有效数值。")
            if not is_blank(condition.get("max")):
                upper = number(condition["max"])
                if upper is None:
                    raise ReportError("范围上限不是有效数值。")
            if lower is None and upper is None:
                return lambda x: True
            if lower is not None and upper is not None and lower > upper:
                raise ReportError("范围下限不能大于上限。")
            return lambda x: number(x) is not None and (
                (lower is None or number(x) >= lower - epsilon) and
                (upper is None or number(x) <= upper + epsilon))
        if is_blank(raw):
            return lambda x: True
        bound = number(raw)
        if bound is None:
            raise ReportError("请输入有效数值，可使用千分位、科学计数法或百分号。")
        comparisons = {
            "gt": lambda x: x > bound + epsilon, "gte": lambda x: x >= bound - epsilon,
            "lt": lambda x: x < bound - epsilon, "lte": lambda x: x <= bound + epsilon,
            "eq": lambda x: abs(x - bound) <= epsilon, "ne": lambda x: abs(x - bound) > epsilon,
        }
        return lambda x: number(x) is not None and comparisons[op](number(x))
    if column["type"] == "date":
        if op not in ("range", "eq", "gte", "lte"):
            raise ReportError("日期列的筛选方式无效。")
        bounds = (condition.get("min"), condition.get("max")) if op == "range" else (raw, raw)
        for bound in bounds:
            if bound:
                try:
                    date.fromisoformat(bound)
                except (TypeError, ValueError) as exc:
                    raise ReportError("请输入有效日期。") from exc
        lower, upper = bounds
        if op == "range" and lower and upper and lower > upper:
            raise ReportError("开始日期不能晚于结束日期。")
        if op == "gte":
            upper = None
        if op == "lte":
            lower = None
        def date_matches(value):
            if is_blank(value):
                return False
            try:
                parsed = date.fromisoformat(str(value)[:10]).isoformat()
            except (ValueError, TypeError):
                return False
            return (not lower or parsed >= lower) and (not upper or parsed <= upper)
        return date_matches
    if op not in ("fuzzy", "contains", "eq", "ne"):
        raise ReportError("文本列的筛选方式无效。")
    if is_blank(raw):
        return lambda x: True
    if op == "fuzzy":
        return lambda x: not is_blank(x) and fuzzy_match(x, raw)
    if op == "contains":
        return lambda x: not is_blank(x) and normalized(raw) in normalized(x)
    return lambda x: not is_blank(x) and ((normalized(x) == normalized(raw)) if op == "eq" else (normalized(x) != normalized(raw)))


def query_table(table: dict, request: dict) -> dict:
    epsilon = parse_epsilon(request.get("epsilon", 1e-8))
    columns = table["columns"]
    index = {column["id"]: i for i, column in enumerate(columns)}
    conditions = request.get("filters", [])
    if not isinstance(conditions, list) or len(conditions) > 30:
        raise ReportError("最多可组合 30 个列条件。")
    predicates = []
    for condition in conditions:
        if not isinstance(condition, dict) or condition.get("column") not in index:
            raise ReportError("筛选列不存在，请重新选择。")
        idx = index[condition["column"]]
        value_type = condition.get("value_type", columns[idx]["type"])
        if value_type not in ("text", "number", "date"):
            raise ReportError("内容格式应为文本、数字或日期。")
        effective_column = {**columns[idx], "type": value_type}
        predicates.append((idx, build_predicate(effective_column, condition, epsilon)))
    query = str(request.get("search", ""))[:500]
    matches = []
    for idx, row in enumerate(table["rows"]):
        if all(predicate(row[col]) for col, predicate in predicates) and (
                not query or any(fuzzy_match(value, query) for value in row if not is_blank(value))):
            matches.append(idx)
    sort = request.get("sort") or {}
    if sort.get("column"):
        if sort["column"] not in index or sort.get("direction") not in ("asc", "desc"):
            raise ReportError("排序设置无效。")
        idx = index[sort["column"]]
        numeric = columns[idx]["type"] == "number"
        populated = [i for i in matches if not is_blank(table["rows"][i][idx])]
        blanks = [i for i in matches if is_blank(table["rows"][i][idx])]
        populated.sort(key=lambda i: number(table["rows"][i][idx]) if numeric else normalized(table["rows"][i][idx]),
                       reverse=sort["direction"] == "desc")
        matches = populated + blanks
    try:
        size = min(200, max(1, int(request.get("page_size", 50))))
        pages = max(1, math.ceil(len(matches) / size))
        page = min(pages, max(1, int(request.get("page", 1))))
    except (TypeError, ValueError, OverflowError) as exc:
        raise ReportError("分页参数无效。") from exc
    chosen = matches[(page - 1) * size: page * size]
    return {"columns": columns, "layout": table.get("layout", "table"), "total": len(table["rows"]), "matched": len(matches),
            "page": page, "pages": pages, "page_size": size, "epsilon": epsilon,
            "rows": [{"excel_row": table["source_rows"][i], "values": table["rows"][i]} for i in chosen]}
