from __future__ import annotations

import re
import threading
from pathlib import Path

from . import config as settings
from .config import atomic_json, load_config, source_directory, timestamp
from .calculation import decimal_value, size_order
from .inventory import MissingSheetError, normalize_code, parse_workbook
from .sync import cached_manifest, save_manifest, synchronize
from .wind import IndustryProvider


def demo_source(root: Path, day: str):
    from openpyxl import Workbook
    folder = root / day
    folder.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "库存券单"
    sheet.append(["合成演示券单 · 不代表真实可融库存"])
    sheet.append(["证券代码", "证券名称", "可融数量（股）", "名本（万元）", "出借费率", "所属行业", "期限", "实际可融股数", "来源"])
    for row in [
        ["600519.SH", "贵州茅台", 12000, 1680, .025, "食品饮料", "28天"],
        ["600519.SH", "贵州茅台", 8000, 1120, .028, "食品饮料", "14天"],
        ["300750.SZ", "宁德时代", 160000, 3200, .032, "电力设备", "28天"],
        ["601318.SH", "中国平安", 420000, 2100, .022, "非银金融", "28天"],
        ["000001.SZ", "平安银行", 800000, 960, .018, "银行", "14天"],
        ["600036.SH", "招商银行", 300000, 1050, .021, "银行", "28天"],
        ["688981.SH", "中芯国际", 50000, 450, .04, "电子", "14天"],
        ["000858.SZ", "五粮液", 0, 0, .03, "食品饮料", "28天"],
    ]:
        source = "互换" if row[0] == "600519.SH" and row[6] == "28天" else "外部券源（演示）" if row[0] == "600519.SH" else "合成券源"
        sheet.append([*row, row[2] + 37 if row[2] else 0, source])
        sheet.cell(sheet.max_row, 5).number_format = "0.00%"
    book.save(folder / "合成演示_库存券单.xlsx")
    book.close()


def summarize(code: str, rows: list[dict]):
    quantities = [r["quantity"] for r in rows if r["quantity"] is not None]
    nominals = [r["nominal"] for r in rows if r["nominal"] is not None]
    rates = sorted(set(r["rate"] for r in rows if r["rate"] is not None))
    raw_rates = sorted(set(r["rate_raw"] for r in rows if r["rate"] is None and r["rate_raw"]))
    industries = sorted(set(r["industry"] for r in rows if r["industry"]))
    return {
        "code": code, "name": next((r["name"] for r in rows if r["name"]), "未提供名称"),
        "quantity": sum(quantities) if quantities else None,
        "quantity_complete": len(quantities) == len(rows),
        "quantity_primary": sum(r["quantity_primary"] for r in rows if r.get("quantity_primary") is not None) if any(r.get("quantity_primary") is not None for r in rows) else None,
        "quantity_actual": sum(r["quantity_actual"] for r in rows if r.get("quantity_actual") is not None) if any(r.get("quantity_actual") is not None for r in rows) else None,
        "quantity_policy": rows[0].get("quantity_policy", "primary"),
        "nominal": sum(nominals) if nominals else None,
        "nominal_complete": len(nominals) == len(rows),
        "nominal_estimated": any(r["nominal_source"].startswith("估算") for r in rows),
        "rate_min": min(rates) if rates else None, "rate_max": max(rates) if rates else None,
        "rate_raw": raw_rates, "rate_complete": all(r["rate"] is not None or r["rate_raw"] for r in rows),
        "industry": " / ".join(industries) or None,
        "industry_source": " / ".join(sorted(set(r["industry_source"] for r in rows if r["industry"]))) or None,
        "rows": rows,
    }


class InventoryService:
    def __init__(self, root: Path = settings.ROOT, demo=False, clock=None):
        self.root, self.demo = root, demo
        self.clock = clock or settings.today
        self.config_path = root / "data" / "demo" / "config.local.json" if demo else root / "config.local.json"
        self.config = load_config(self.config_path)
        self.data_root = root / "data" / ("demo" if demo else "live")
        self.provider = IndustryProvider(self.data_root)
        self.records = []
        self.lock = threading.RLock()
        self.state = {"date": self.clock(), "mode": "demo" if demo else "live", "state": "loading", "warnings": [], "files": [], "record_count": 0, "security_count": 0}

    def active_config(self):
        config = dict(self.config)
        if self.demo:
            config.update(source_root=str(self.root / "data" / "demo_source"), wind_enabled=False, file_glob="*.xlsx", file_selection="latest", sheet_name="库存券单", columns={}, header_row=None, quantity_policy="minimum")
        return config

    def update_config(self, patch, progress=None):
        unknown = set(patch) - settings.DEFAULTS.keys()
        if unknown:
            raise ValueError(f"未知配置：{', '.join(sorted(unknown))}")
        config = {**self.config, **patch}
        settings.validate(config)
        atomic_json(self.config_path, config)
        self.config = config
        return self.refresh(progress)

    def _parse(self, manifest, config, progress=None):
        records, warnings, files = [], list(manifest["warnings"]), []
        for entry in manifest["files"]:
            if progress:
                progress(f"正在读取本地文件：{entry['name']} → {config['sheet_name']}")
            try:
                rows, notices, mapping = parse_workbook(Path(entry["path"]), config)
            except Exception as exc:
                entry["read_error"] = str(exc)
                hint = ""
                if isinstance(exc, MissingSheetError):
                    entry["available_sheets"] = exc.available_sheets
                    hint = "\n请在 Excel 中打开下方本地文件，核对文件及工作表。"
                    if any(char in config["file_glob"] for char in "*?["):
                        hint += (f"当前筛选「{config['file_glob']}」可能匹配其他 Excel；"
                                 "默认券单请在设置中恢复「申万内部券单{date}.xlsx」和「库存券单」。")
                raise ValueError(f"{entry['name']}：{exc}\n共享源：{entry['source']}\n本地文件：{entry['path']}{hint}") from exc
            for row in rows:
                row.update(local_path=entry["path"], source_path=entry["source"], sha256=entry["sha256"])
            records.extend(rows)
            warnings.extend(notices)
            files.append({**entry, **mapping, "records": len(rows)})
        return records, warnings, files

    def load_today(self, progress=None):
        """Use a verified local snapshot before considering a shared-drive download."""
        return self.refresh(progress, prefer_local=True)

    def refresh(self, progress=None, *, prefer_local=False):
        with self.lock:
            day, config = self.clock(), self.active_config()
            self.provider.failed_keys.clear()
            manifest, parsed, issue, downloaded = None, None, None, None
            local_warnings, using_local = [], False
            if prefer_local:
                if progress:
                    progress("正在检查当日本地券单…")
                try:
                    manifest = cached_manifest(self.data_root, day, config)
                    if manifest:
                        parsed = self._parse(manifest, config, progress)
                        using_local = True
                except Exception as exc:
                    local_warnings.append(f"当日本地券单不可用，已尝试重新下载：{exc}")
                    manifest, parsed = None, None
            if not using_local:
                try:
                    if self.demo:
                        demo_source(self.root / "data" / "demo_source", day)
                    manifest = synchronize(config, day, self.data_root, progress)
                    manifest["file_glob"] = config["file_glob"]
                    downloaded = manifest
                    parsed = self._parse(manifest, config, progress)
                    save_manifest(self.data_root, manifest)
                except Exception as exc:
                    issue = str(exc)
                    manifest = None
                    try:
                        manifest = cached_manifest(self.data_root, day, config)
                        if manifest:
                            parsed = self._parse(manifest, config, progress)
                    except Exception as cache_exc:
                        issue += f"；缓存也无法使用：{cache_exc}"
                        manifest = None
            self.records = parsed[0] if manifest and parsed else []
            try:
                source = str(source_directory(config, day))
            except ValueError:
                source = "未配置"
            self.state = {
                "date": day, "mode": "demo" if self.demo else "live",
                "state": "local" if using_local else "cached" if manifest and issue else "ready" if manifest else "unavailable",
                "source_dir": source, "refreshed_at": timestamp(),
                "downloaded_at": manifest["downloaded_at"] if manifest else None,
                # A copied workbook can fail validation. Keep its provenance separate
                # from the accepted inventory/cache so it can still be inspected.
                "downloaded_files": downloaded["files"] if downloaded else [],
                "attempt_downloaded_at": downloaded["downloaded_at"] if downloaded else None,
                "sheet_name": config["sheet_name"],
                "warnings": local_warnings + ([issue] if issue else []) + (parsed[1] if manifest and parsed else []),
                "files": parsed[2] if manifest and parsed else [],
                "record_count": len(self.records), "security_count": len(set(r["code"] for r in self.records)),
                "selection": config["file_selection"],
            }
            return self.state

    def status(self, progress=None):
        with self.lock:
            if self.state["state"] == "loading" or self.state["date"] != self.clock():
                self.load_today(progress)
            return self.state

    def query(self, text: str, targets=None, input_targets=None, progress=None):
        with self.lock:
            status = self.status(progress)
            tokens = list(dict.fromkeys(x for x in re.split(r"[\s,，;；、]+", text.strip()) if x))
            if not tokens:
                raise ValueError("请输入至少一个证券代码或名称")
            if len(tokens) > 100 or len(text) > 10000:
                raise ValueError("每次最多查询 100 个标的")
            targets = targets or {}
            if not isinstance(targets, dict):
                raise ValueError("目标名本应按证券代码提供")
            groups, codes, quote_codes = [], set(), set()
            for token in tokens:
                normalized = normalize_code(token)
                if normalized:
                    quote_codes.add(normalized)
                rows = [dict(r) for r in self.records if (r["code"] == normalized if normalized else token.casefold() in r["name"].casefold())]
                if normalized and not rows:
                    codes.add(normalized)
                grouped = {}
                for row in rows:
                    grouped.setdefault(row["code"], []).append(row)
                    quote_codes.add(row["code"])
                    if not row["industry"]:
                        codes.add(row["code"])
                groups.append((token, grouped))
            resolved_targets = dict(targets)
            batch_codes = {}
            for token, grouped in groups:
                if token not in (input_targets or {}):
                    continue
                code = normalize_code(token)
                if not code:
                    if len(grouped) != 1:
                        raise ValueError(f"批量名本未应用：名称「{token}」匹配 {len(grouped)} 只股票，请改用完整证券代码后重试")
                    code = next(iter(grouped))
                value = input_targets[token]
                if code in batch_codes and decimal_value(batch_codes[code]) != decimal_value(value):
                    raise ValueError(f"批量名本冲突：{code} 对应多个不同金额，请核对后重试")
                batch_codes[code] = value
            resolved_targets.update(batch_codes)
            if progress:
                progress("正在获取本次 Wind 最新价与行业…" if not self.demo else "正在生成合成演示价格…")
            if self.demo:
                prices = {"600519.SH": "1400.25", "300750.SZ": "200.18", "601318.SH": "50.36", "000001.SZ": "12.17", "600036.SH": "35.28", "688981.SH": "90.45", "000858.SZ": "130.52"}
                quotes = {code: {"price": prices[code], "retrieved_at": timestamp(), "wind_returned_at": None, "source": "合成演示价格"} for code in quote_codes if code in prices}
                industries, notices = {}, []
            else:
                industries, quotes, notices = self.provider.market(quote_codes, codes, status["date"], self.active_config())
            results, unique = [], {}
            for token, grouped in groups:
                matches = []
                for code, rows in grouped.items():
                    for row in rows:
                        if not row["industry"] and code in industries:
                            row["industry"] = industries[code]
                            row["industry_source"] = "Wind · " + self.config["industry_label"]
                    item = summarize(code, rows)
                    item["quote"] = quotes.get(code)
                    price = decimal_value((quotes.get(code) or {}).get("price"))
                    item["inventory_market_value"] = item["quantity"] * price if item["quantity_complete"] and item["quantity"] is not None and price is not None and price > 0 else None
                    item["calculation"] = size_order(resolved_targets.get(code), (quotes.get(code) or {}).get("price"), item["quantity"], item["quantity_complete"])
                    matches.append(item)
                    unique[code] = item
                results.append({"input": token, "matches": matches})
            values = list(unique.values())
            quantities = [v["quantity"] for v in values if v["quantity"] is not None]
            nominals = [v["nominal"] for v in values if v["nominal"] is not None]
            return {"status": status, "results": results, "quotes": quotes, "industries": industries, "warnings": notices, "targets": resolved_targets, "summary": {
                "inputs": len(tokens), "matched": len(values), "unmatched": sum(not g["matches"] for g in results),
                "quantity": sum(quantities) if quantities else None,
                "quantity_complete": bool(values) and all(v["quantity_complete"] for v in values),
                "nominal": sum(nominals) if nominals else None,
                "nominal_complete": bool(values) and all(v["nominal_complete"] for v in values),
                "nominal_estimated": any(v["nominal_estimated"] for v in values),
            }}
