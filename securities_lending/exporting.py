"""Traceable native exports; numeric business values remain exact until display."""
import csv
import io
import os
import tempfile
from decimal import Decimal
from pathlib import Path

from .calculation import decimal_value


def excel_value(value, heading):
    """Use Excel numeric cells where its 15 significant digits are sufficient."""
    numeric = isinstance(value, Decimal) or any(word in heading for word in ("（元", "（股", "股数", "股数合计", "费率（%"))
    amount = decimal_value(value) if numeric and value is not None else None
    if amount is None:
        return value
    if len(amount.normalize().as_tuple().digits) <= 15:
        return int(amount) if amount == amount.to_integral_value() else float(amount)
    return str(amount)


def safe_text(value):
    text = "" if value is None else str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r")) else text


def rate_text(item):
    def fmt(value):
        return format(value, ".4f").rstrip("0").rstrip(".")
    low, high = item["rate_min"], item["rate_max"]
    rate = "" if low is None else f"{fmt(low)}%" if low == high else f"{fmt(low)}–{fmt(high)}%"
    return " / ".join(value for value in [rate, *item["rate_raw"]] if value) or "—"


def quality(item):
    notes = []
    if not item.get("quantity_complete", False):
        notes.append("库存数量缺失/不完整，不能确认上限")
    if not item.get("nominal_complete", False):
        notes.append("券单名本缺失/仅合计已知值")
    if item.get("nominal_estimated", False):
        notes.append("券单名本含估算")
    if not item.get("rate_complete", False):
        notes.append("费率缺失/不完整")
    return "；".join(notes) or "完整"


def source_description(row):
    rate = f"{row['rate']}%" if row["rate"] is not None else row["rate_raw"] or "缺失"
    return f"{row['file']} 第 {row['row']} 行 | 有效 {row['quantity']} 股（实际 {row.get('quantity_actual')} / 可融 {row.get('quantity_primary')}） | 原始市值 {row['nominal']} 元（{row['nominal_source']}，仅参考） | 费率 {rate} | 期限 {row['term'] or '未提供'} | 来源 {row.get('source') or '未提供'}"


def export_tables(items, status):
    headers = ["数据日期", "数据模式", "库存状态", "证券代码", "名称", "行业", "最新价（元）", "价格来源", "价格获取时间", "目标名本（元）", "换算需求（股）", "校验通过股数", "最大可借（股）", "取整名本（元）", "未用名本（元）", "校验结果", "券单原始数量（股）", "券单名义本金（元）", "出借费率", "数据完整性", "库存下载时间", "Wind返回时间", "行业来源", "券源明细", "本地快照文件"]
    detail_headers = ["数据日期", "数据模式", "库存状态", "库存下载时间", "证券代码", "名称", "可融数量（股）", "券单名义本金（元）", "名本来源", "出借费率（%）", "原始费率", "期限", "行业", "行业来源", "文件名", "Excel行号", "本地快照文件", "共享源文件", "SHA256", "券源来源"]
    rows, details = [], []
    for item in items:
        calc, quote = item["calculation"], item["quote"] or {}
        rows.append([status["date"], status["mode"], status["state"], item["code"], item["name"], item["industry"], quote.get("price"), quote.get("source"), quote.get("retrieved_at"), calc.get("target_nominal"), calc.get("requested_shares"), calc.get("executable_shares"), calc.get("maximum_shares"), calc.get("actual_nominal"), calc.get("unused_nominal"), calc["message"], item["quantity"], item["nominal"], rate_text(item), quality(item), status.get("downloaded_at"), quote.get("wind_returned_at"), item.get("industry_source"), "\n".join(source_description(row) for row in item["rows"]), "\n".join(dict.fromkeys(row.get("local_path", "") for row in item["rows"]))])
        for row in item["rows"]:
            details.append([status["date"], status["mode"], status["state"], status.get("downloaded_at"), item["code"], item["name"], row["quantity"], row["nominal"], row["nominal_source"], row["rate"], row["rate_raw"], row["term"], row["industry"], row["industry_source"], row["file"], row["row"], row.get("local_path"), row.get("source_path"), row.get("sha256"), row.get("source")])
    headers[17] = "券单原始市值（元，仅参考）"
    headers.extend(["库存市值（元，本次最新价）", "实际可融股数合计", "可融股数合计", "库存取值口径"])
    for row, item in zip(rows, items):
        row.extend([item.get("inventory_market_value"), item.get("quantity_actual"), item.get("quantity_primary"), "逐笔取较小值再合计" if item.get("quantity_policy") == "minimum" else "可融数量列"])
    detail_headers[7] = "券单原始市值（元，仅参考）"
    detail_headers.extend(["实际可融股数", "可融股数", "券单市值上限（元，仅参考）"])
    sources = [row for item in items for row in item["rows"]]
    for row, source in zip(details, sources):
        row.extend([source.get("quantity_actual"), source.get("quantity_primary"), source.get("nominal_limit")])
    return [("查询汇总", headers, rows), ("券源明细", detail_headers, details)]


def clipboard_text(items, status):
    _, headers, rows = export_tables(items, status)[0]
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter="\t", lineterminator="\r\n")
    # Preserve one physical line per security when pasted into Excel.
    for row in [headers, *rows]:
        writer.writerow(safe_text(value).replace("\r", " ").replace("\n", " | ").replace("\t", " ") for value in row)
    return stream.getvalue()


def write_export(path, items, status, detail_only=False):
    path = Path(path)
    tables = export_tables(items, status)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(prefix=".delta1_export_", suffix=path.suffix, dir=path.parent)
    os.close(handle)
    try:
        if path.suffix.lower() == ".xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
            book = Workbook()
            book.remove(book.active)
            for name, headers, rows in tables:
                sheet = book.create_sheet(name)
                sheet.append(headers)
                for row in rows:
                    sheet.append([excel_value(value, heading) for value, heading in zip(row, headers)])
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                        elif isinstance(cell.value, (int, float)):
                            cell.number_format = "#,##0.####"
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                sheet.sheet_view.showGridLines = False
                for cell in sheet[1]:
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="193E4A")
                sheet.freeze_panes = "F2"
                sheet.auto_filter.ref = sheet.dimensions
                for index, heading in enumerate(headers, 1):
                    sheet.column_dimensions[get_column_letter(index)].width = 32 if any(key in heading for key in ("时间", "文件", "明细", "结果", "完整性")) else 20
            note = book.create_sheet("数据说明")
            for row in [("模式", "合成演示数据，不代表真实库存或行情" if status["mode"] == "demo" else "正式券单"), ("库存状态", status["state"]), ("日期（北京时间）", status["date"]), ("下载时间", status.get("downloaded_at")), ("本次同步检查时间", status.get("refreshed_at")), ("来源目录", status.get("source_dir")), ("口径", "库存市值=有效股数×本次最新价；原始券单市值仅参考；目标名本单独输入。股数按100股向下取整。"), ("精度", "不超过15位有效数字时保存为Excel数值；超出时用精确文本保留，避免截断。代码始终为文本。"), ("价格", "本次WSQ RT_LAST；返回时间不是逐笔成交时间保证。"), ("提示", "\n".join(status.get("warnings", [])))]:
                note.append(row)
                for cell in note[note.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            note.column_dimensions["A"].width = 25
            note.column_dimensions["B"].width = 100
            try:
                book.save(temporary)
            finally:
                book.close()
        else:
            _, headers, rows = tables[1 if detail_only else 0]
            with open(temporary, "w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                for row in [headers, *rows]:
                    writer.writerow(safe_text(value) for value in row)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
