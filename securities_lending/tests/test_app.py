import json
import os
import shutil
import tempfile
import threading
import unittest
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook

from securities_lending.config import DEFAULTS, ROOT, atomic_json, source_directory
from securities_lending.inventory import normalize_code, parse_workbook
from securities_lending.server import make_server
from securities_lending.service import InventoryService
from securities_lending.wind import IndustryProvider
from securities_lending.wind_worker import fetch


class AppTests(unittest.TestCase):
    def setUp(self):
        temp_root = ROOT / "data" / "test_tmp"
        temp_root.mkdir(parents=True, exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=temp_root)
        self.root = Path(self.tmp.name)
        self.day = "2026-10-14"
        self.config = {**DEFAULTS, "source_root": str(self.root / "source"), "wind_enabled": False, "file_glob": "*.xlsx", "columns": {}, "quantity_policy": "primary", "numeric_rate_unit": "percent"}
        atomic_json(self.root / "config.local.json", self.config)

    def tearDown(self):
        self.tmp.cleanup()

    def workbook(self, name="券单.xlsx", rows=None, headers=None, sheet="库存券单", preamble=True):
        path = Path(self.config["source_root"]) / self.day / name
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        ws = book.active
        ws.title = sheet
        if preamble:
            ws.append(["公司库存券单"])
        ws.append(headers or ["证券代码", "证券名称", "可融数量（万股）", "名本（万元）", "出借费率", "所属行业"])
        for row in rows or [[1,"平安银行",12,120,2.5,"银行"],["000001.SZ","平安银行",3,30,"3%","银行"],["600519","贵州茅台",1,1400,2.2,""]]:
            ws.append(row)
        book.save(path)
        book.close()
        return path

    def service(self):
        return InventoryService(self.root, clock=lambda:self.day)

    def test_codes_and_leading_zero(self):
        for value in (1, 1.0, "000001", "SZ000001", "000001.SZ", "1.0"):
            self.assertEqual(normalize_code(value), "000001.SZ")
        self.assertEqual(normalize_code("920001"), "920001.BJ")
        self.assertEqual(normalize_code("688981"), "688981.SH")
        self.assertEqual(normalize_code("SH000001.SZ"), "")

    def test_runtime_date_and_dated_source(self):
        config = {**self.config, "source_root": str(self.root / "2026-10-14")}
        self.assertEqual(source_directory(config,"2026-10-15"), self.root / "2026-10-15")
        config["source_root"] = str(self.root / "{date}")
        self.assertEqual(source_directory(config,"2026-10-16"), self.root / "2026-10-16")

    def test_first_status_initializes_inventory(self):
        self.workbook()
        self.assertEqual(self.service().status()["state"],"ready")

    def test_header_units_aggregation_and_multi_input(self):
        self.workbook()
        service = self.service()
        self.assertEqual(service.refresh()["state"],"ready")
        response = service.query("000001，平安银行\n600519;不存在")
        bank = response["results"][0]["matches"][0]
        self.assertEqual(bank["quantity"],150000)
        self.assertEqual(bank["nominal"],1500000)
        self.assertEqual((bank["rate_min"],bank["rate_max"]),(2.5,3))
        self.assertEqual(len(bank["rows"]),2)
        self.assertEqual(bank["rows"][0]["row"],3)
        self.assertEqual(response["summary"]["matched"],2)
        self.assertEqual(response["summary"]["unmatched"],1)
        self.assertEqual(response["summary"]["quantity"],160000)

    def test_percent_cell_and_decimal_config(self):
        path = self.workbook(rows=[["000001","平安银行",1,10,.025,"银行"]])
        from openpyxl import load_workbook
        book = load_workbook(path)
        book.active["E3"].number_format = "0.00%"
        book.save(path)
        book.close()
        records, _, _ = parse_workbook(path,self.config)
        self.assertEqual(records[0]["rate"],2.5)
        path = self.workbook(rows=[["000001","平安银行",1,10,.032,"银行"]])
        records, _, _ = parse_workbook(path,{**self.config,"numeric_rate_unit":"decimal"})
        self.assertEqual(records[0]["rate"],Decimal("3.2"))

    def test_missing_quantity_formula_and_nominal_estimate(self):
        path = self.workbook(headers=["证券代码","证券名称","可融数量","参考价","出借费率"], rows=[["000001","平安银行",100,12,"协商"],["600519","贵州茅台","=2*100",1400,.03]])
        records, warnings, _ = parse_workbook(path,self.config)
        self.assertEqual(records[0]["nominal"],1200)
        self.assertTrue(records[0]["nominal_source"].startswith("估算"))
        self.assertIsNone(records[1]["quantity"])
        self.assertTrue(any("公式没有缓存" in x for x in warnings))

    def test_units_in_cells_override_headers(self):
        path = self.workbook(rows=[["000001","平安银行","100股","2亿元","2.8%","银行"]])
        records, _, _ = parse_workbook(path,self.config)
        self.assertEqual(records[0]["quantity"],100)
        self.assertEqual(records[0]["nominal"],200000000)

    def test_ambiguous_header_requires_mapping(self):
        path=self.workbook(headers=["证券代码","可融数量（股）","可融数量（万股）"], rows=[["000001",100,1]])
        with self.assertRaisesRegex(ValueError,"歧义"):
            parse_workbook(path,self.config)
        records,_,_=parse_workbook(path,{**self.config,"columns":{"quantity":"可融数量（股）"}})
        self.assertEqual(records[0]["quantity"],100)

    def test_latest_version_and_lock_file(self):
        old=self.workbook("old.xlsx",rows=[["000001","平安银行",1,1,2,"银行"]])
        new=self.workbook("new.xlsx",rows=[["000001","平安银行",4,1,2,"银行"]])
        self.workbook("~$lock.xlsx")
        os.utime(old,(100,100));os.utime(new,(200,200))
        service=self.service()
        status=service.refresh()
        self.assertEqual(status["files"][0]["name"],"new.xlsx")
        self.assertEqual(service.query("000001")["summary"]["quantity"],40000)

    def test_merge_skips_identical_file_content(self):
        path=self.workbook()
        shutil.copy2(path,path.with_name("duplicate.xlsx"))
        service=self.service();service.config["file_selection"]="merge"
        status=service.refresh()
        self.assertEqual(status["record_count"],3)
        self.assertTrue(any("完全相同" in x for x in status["warnings"]))

    def test_cache_and_rollover_never_use_previous_day(self):
        path=self.workbook();service=self.service();service.refresh();path.unlink()
        self.assertEqual(service.refresh()["state"],"cached")
        self.assertEqual(service.query("000001")["summary"]["quantity"],150000)
        self.day="2026-10-15"
        self.assertEqual(service.query("000001")["status"]["state"],"unavailable")
        self.assertEqual(service.records,[])

    def test_invalid_new_workbook_falls_back_to_valid_snapshot(self):
        self.workbook();service=self.service();original=service.refresh()
        self.workbook(sheet="其他Sheet")
        fallback=service.refresh()
        self.assertEqual(fallback["state"],"cached")
        self.assertEqual(fallback["files"][0]["path"],original["files"][0]["path"])

    def test_changed_source_does_not_reuse_cache(self):
        self.workbook();service=self.service();service.refresh()
        service.config["source_root"]=str(self.root/"other")
        self.assertEqual(service.refresh()["state"],"unavailable")

    def test_changed_glob_does_not_reuse_cache(self):
        self.workbook();service=self.service();service.refresh();service.config["file_glob"]="missing*.xlsx"
        self.assertEqual(service.refresh()["state"],"unavailable")

    def test_tampered_cache_is_rejected(self):
        source=self.workbook();service=self.service();status=service.refresh()
        Path(status["files"][0]["path"]).write_bytes(b"corrupt")
        source.unlink()
        self.assertEqual(service.refresh()["state"],"unavailable")

    def test_zero_is_distinct_from_missing(self):
        self.workbook(rows=[["000001","平安银行",0,0,2,"银行"],["600519","贵州茅台",None,None,None,""]])
        service=self.service();service.refresh();response=service.query("000001 600519 300750")
        bank=response["results"][0]["matches"][0]
        self.assertEqual(bank["quantity"],0)
        self.assertTrue(bank["quantity_complete"])
        self.assertFalse(response["summary"]["quantity_complete"])
        self.assertEqual(response["summary"]["unmatched"],1)

    def test_missing_sheet_diagnostic(self):
        path=self.workbook(sheet="发行券单")
        with self.assertRaisesRegex(ValueError,"已有：发行券单"):
            parse_workbook(path,self.config)

    def test_all_missing_values_do_not_display_zero_totals(self):
        self.workbook(rows=[["600519","贵州茅台",None,None,None,""]])
        service=self.service();service.refresh();response=service.query("600519")
        self.assertIsNone(response["summary"]["quantity"])
        self.assertIsNone(response["summary"]["nominal"])

    def test_wind_response_matches_by_code_and_errors(self):
        fake=SimpleNamespace(start=lambda **k:SimpleNamespace(ErrorCode=0),wss=lambda *a:SimpleNamespace(ErrorCode=0,Codes=["600519.SH","000001.SZ"],Data=[["食品饮料","银行"]]))
        request={"codes":["000001.SZ","600519.SH"],"field":"industry_sw","options":"industryType=1"}
        self.assertEqual(fetch(request,fake),{"600519.SH":"食品饮料","000001.SZ":"银行"})
        fake.wss=lambda *a:SimpleNamespace(ErrorCode=-1)
        with self.assertRaisesRegex(RuntimeError,"Wind 查询失败"):
            fetch(request,fake)

    def test_wind_timeout_keeps_inventory_available(self):
        from subprocess import TimeoutExpired
        provider=IndustryProvider(self.root/"data")
        with patch("securities_lending.wind.subprocess.run",side_effect=TimeoutExpired("wind",1)) as run:
            industries,warning=provider.lookup(["000001.SZ"],self.day,{**self.config,"wind_enabled":True})
            self.assertEqual(industries,{})
            self.assertIn("Wind",warning)
            provider.lookup(["000001.SZ"],self.day,{**self.config,"wind_enabled":True})
            self.assertEqual(run.call_count,1)

    def test_wind_cache_is_daily(self):
        provider=IndustryProvider(self.root/"data")
        result=SimpleNamespace(returncode=0,stdout=json.dumps({"industries":{"000001.SZ":"银行"},"error":None}))
        with patch("securities_lending.wind.subprocess.run",return_value=result) as run:
            config={**self.config,"wind_enabled":True}
            provider.lookup(["000001.SZ"],self.day,config)
            provider.lookup(["000001.SZ"],self.day,config)
            self.assertEqual(run.call_count,1)
            provider.lookup(["000001.SZ"],"2026-10-15",config)
            self.assertEqual(run.call_count,2)

    def test_demo_does_not_change_live_config(self):
        service=InventoryService(self.root,demo=True,clock=lambda:self.day)
        self.assertEqual(service.refresh()["state"],"ready")
        self.assertEqual(service.query("600519")["summary"]["quantity"],20000)
        self.assertEqual(json.loads((self.root/"config.local.json").read_text())["source_root"],self.config["source_root"])

    def test_http_query_config_validation_and_cross_origin(self):
        self.workbook();service=self.service();service.refresh()
        server=make_server(service,0)
        worker=threading.Thread(target=server.serve_forever,daemon=True);worker.start()
        base=f"http://127.0.0.1:{server.server_port}"
        try:
            with urlopen(base) as response:
                self.assertIn("券源查询",response.read().decode())
            def post(path,body,origin=None):
                headers={"Content-Type":"application/json"}
                if origin:headers["Origin"]=origin
                return urlopen(Request(base+path,data=json.dumps(body).encode(),headers=headers))
            with post("/api/query",{"symbols":"平安"}) as response:
                self.assertEqual(json.load(response)["summary"]["matched"],1)
            with self.assertRaises(HTTPError) as error:
                post("/api/refresh",{},"https://example.com")
            self.assertEqual(error.exception.code,403)
            with self.assertRaises(HTTPError) as error:
                post("/api/config",{"file_selection":"bad"})
            self.assertEqual(error.exception.code,400)
            self.assertEqual(service.config["file_selection"],"latest")
        finally:
            server.shutdown();server.server_close();worker.join()


if __name__ == "__main__":
    unittest.main()
