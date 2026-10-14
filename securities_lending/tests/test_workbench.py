import csv
import io
import json
import tempfile
import threading
import unittest
import subprocess
import sys
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest

from securities_lending.batch_input import parse_batch
from securities_lending.calculation import size_order
from securities_lending.config import DEFAULTS, ROOT, atomic_json, load_config
from securities_lending.exporting import source_note, write_export
from securities_lending.inventory import number, parse_workbook
from securities_lending.service import InventoryService
from securities_lending.sync import synchronize
from securities_lending.tests import test_desktop as native_tests


class InventoryRegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.day = "2030-02-07"
        self.config = {**DEFAULTS, "columns": dict(DEFAULTS["columns"]), "source_root": str(self.root / "source"), "wind_enabled": False}
        atomic_json(self.root / "config.local.json", self.config)

    def workbook(self, name=None, rows=None, headers=None):
        path = self.root / "source" / self.day / (name or f"申万内部券单{self.day}.xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        book.active.title = "库存券单"
        book.active.append(headers or ["证券代码", "证券简称", "实际可融股数", "可融股数", "可融市值", "参考费率", "来源", "期限"])
        for row in rows or [["000001.SZ", "测试银行", 27220, 27200, 9999999, .05, "甲", 14], ["000001.SZ", "测试银行", 123129, 123200, 123, .045, "乙", 28]]:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def service(self):
        return InventoryService(self.root, clock=lambda: self.day)

    def test_decimal_unit_conversion_survives_floor_boundary(self):
        quantity = number("0.57", 10000)
        self.assertEqual(quantity, Decimal(5700))
        self.assertEqual(size_order(57000, 10, quantity)["executable_shares"], 5700)
        path = self.workbook(headers=["证券代码", "可融数量（万股）"], rows=[["000001", .57]])
        rows, _, _ = parse_workbook(path, {**self.config, "columns": {}, "quantity_policy": "primary"})
        self.assertEqual(rows[0]["quantity"], Decimal(5700))

    def test_exact_filename_dynamic_date_copy_then_parse(self):
        source = self.workbook()
        self.workbook("other_newer.xlsx", rows=[["600000.SH", "不应读取", 100, 100, 1, .01, "甲", 7]])
        service = self.service()
        import securities_lending.service as module
        original = module.parse_workbook
        checked = []
        def local_only(path, config):
            self.assertTrue(path.resolve().is_relative_to((self.root / "data" / "live").resolve()))
            self.assertEqual(path.read_bytes(), source.read_bytes())
            checked.append(path)
            return original(path, config)
        with patch.object(module, "parse_workbook", side_effect=local_only):
            status = service.refresh()
        self.assertEqual(status["state"], "ready")
        self.assertEqual(len(checked), 1)
        self.assertEqual(status["files"][0]["name"], source.name)
        source.unlink()
        self.assertEqual(service.refresh()["state"], "cached")
        service.clock = lambda: "2030-02-08"
        self.assertEqual(service.refresh()["state"], "unavailable")

    def test_missing_sheet_keeps_download_and_never_accepts_sheet1(self):
        source = self.workbook()
        book = load_workbook(source)
        book.active.title = "Sheet1"
        book.save(source)
        book.close()
        service = self.service()
        status = service.refresh()
        self.assertEqual(status["state"], "unavailable")
        self.assertEqual(status["files"], [])
        self.assertEqual(service.records, [])
        self.assertIsNone(status["downloaded_at"])
        self.assertTrue(status["attempt_downloaded_at"])
        entry = status["downloaded_files"][0]
        self.assertEqual(entry["available_sheets"], ["Sheet1"])
        self.assertEqual(entry["source"], str(source))
        self.assertEqual(Path(entry["path"]).read_bytes(), source.read_bytes())
        self.assertEqual(entry["sha256"], entry["source_sha256"])
        self.assertIn(entry["path"], status["warnings"][0])
        self.assertFalse((service.data_root / self.day / "latest.json").exists())

    def test_wildcard_wrong_file_has_actionable_diagnostic_and_default_retry(self):
        correct = self.workbook()
        unrelated = correct.with_name("其他报表.xlsx")
        book = Workbook()
        book.active.title = "Sheet1"
        book.save(unrelated)
        book.close()
        import os
        os.utime(correct, (100, 100))
        os.utime(unrelated, (200, 200))
        service = self.service()
        service.config["file_glob"] = "*.xlsx*"
        status = service.refresh()
        self.assertEqual(status["state"], "unavailable")
        self.assertEqual(status["downloaded_files"][0]["name"], unrelated.name)
        self.assertIn("申万内部券单{date}.xlsx", status["warnings"][0])
        service.config["file_glob"] = DEFAULTS["file_glob"]
        status = service.refresh()
        self.assertEqual(status["state"], "ready")
        self.assertEqual(status["files"][0]["name"], correct.name)

    def test_failed_download_parse_remains_separate_from_valid_cache(self):
        source = self.workbook()
        service = self.service()
        accepted = service.refresh()
        manifest_path = service.data_root / self.day / "latest.json"
        manifest_bytes = manifest_path.read_bytes()
        book = load_workbook(source)
        book.active.title = "Sheet1"
        book.save(source)
        book.close()
        fallback = service.refresh()
        self.assertEqual(fallback["state"], "cached")
        self.assertEqual(fallback["files"], accepted["files"])
        self.assertNotEqual(fallback["downloaded_files"][0]["path"], fallback["files"][0]["path"])
        self.assertEqual(fallback["downloaded_files"][0]["available_sheets"], ["Sheet1"])
        self.assertEqual(manifest_path.read_bytes(), manifest_bytes)
        self.assertEqual(len(service.records), 2)

    def test_copy_preserves_all_sheets_in_single_workbook(self):
        source = self.workbook()
        book = load_workbook(source)
        book.create_sheet("Sheet1", 0)
        book.create_sheet("说明")
        book.save(source)
        book.close()
        status = self.service().refresh()
        self.assertEqual(status["state"], "ready")
        self.assertEqual(len(status["files"]), 1)
        local = Path(status["files"][0]["path"])
        self.assertEqual(local.read_bytes(), source.read_bytes())
        book = load_workbook(local, read_only=True)
        self.assertEqual(book.sheetnames, ["Sheet1", "库存券单", "说明"])
        book.close()
        self.assertEqual(status["record_count"], 2)

    def test_equal_size_copy_corruption_is_rejected(self):
        source = self.workbook()
        def corrupt_copy(original, local):
            data = original.read_bytes()
            local.write_bytes(bytes([data[0] ^ 1]) + data[1:])
        with patch("securities_lending.sync.shutil.copy2", side_effect=corrupt_copy):
            with self.assertRaisesRegex(ValueError, "内容不一致"):
                synchronize(self.config, self.day, self.root / "data" / "live")

    def test_minimum_is_per_source_then_aggregated_and_floored(self):
        self.workbook()
        service = self.service()
        service.refresh()
        item = service.query("000001")["results"][0]["matches"][0]
        self.assertEqual(item["quantity"], 150329)
        self.assertEqual(item["quantity_actual"], 150349)
        self.assertEqual(item["quantity_primary"], 150400)
        self.assertEqual(item["calculation"]["maximum_shares"], 150300)
        self.assertEqual((item["rate_min"], item["rate_max"]), (Decimal("4.5"), Decimal("5")))
        self.assertEqual([row["source"] for row in item["rows"]], ["甲", "乙"])

    def test_missing_either_quantity_never_confirms_capacity(self):
        self.workbook(rows=[["000001", "测试", None, 10000, 200, .05, "甲", 7], ["000001", "测试", 1000, 1000, 200, .05, "乙", 7]])
        service = self.service()
        service.refresh()
        quote = {"price": "10", "retrieved_at": "2030-02-07T10:00:00+08:00"}
        with patch.object(service.provider, "market", return_value=({}, {"000001.SZ": quote}, [])):
            item = service.query("000001", {"000001.SZ": "1000"})["results"][0]["matches"][0]
        self.assertEqual(item["quantity"], 1000)
        self.assertFalse(item["quantity_complete"])
        self.assertIsNone(item["inventory_market_value"])
        self.assertEqual(item["calculation"]["status"], "unknown_inventory")

    def test_market_value_uses_latest_quote_not_sheet_nominal(self):
        self.workbook()
        service = self.service()
        service.refresh()
        quote = {"price": "10.29"}
        with patch.object(service.provider, "market", return_value=({}, {"000001.SZ": quote}, [])):
            item = service.query("000001")["results"][0]["matches"][0]
        self.assertEqual(item["inventory_market_value"], Decimal(150329) * Decimal("10.29"))
        self.assertNotEqual(item["inventory_market_value"], item["nominal"])
        with patch.object(service.provider, "market", return_value=({}, {}, ["offline"])):
            item = service.query("000001")["results"][0]["matches"][0]
        self.assertIsNone(item["inventory_market_value"])

    def test_demo_save_and_invalid_live_config_are_isolated(self):
        live = self.root / "config.local.json"
        live.write_text("invalid live config", encoding="utf-8")
        demo = InventoryService(self.root, demo=True)
        demo.update_config({"source_root": "demo only", "sheet_name": "ignored in demo"})
        self.assertEqual(live.read_text(), "invalid live config")
        self.assertEqual(demo.state["state"], "ready")
        self.assertEqual(InventoryService(self.root, demo=True).config["source_root"], "demo only")

    def test_old_defaults_migrate_without_rewriting_file(self):
        path = self.root / "config.local.json"
        original = '{"source_root":"private source","file_glob":"*.xls*","columns":{},"numeric_rate_unit":"percent"}'
        path.write_text(original)
        config = load_config(path)
        self.assertEqual(config["source_root"], "private source")
        self.assertEqual(config["columns"]["quantity"], "可融股数")
        self.assertEqual(config["numeric_rate_unit"], "decimal")
        self.assertEqual(config["file_glob"], "申万内部券单{date}.xlsx")
        self.assertEqual(path.read_text(), original)

    def test_custom_sheet_filename_columns_remain_configurable(self):
        self.config.update(file_glob="custom_{date}.xlsx", sheet_name="自定Sheet", columns={"code": "代号", "quantity": "可用", "quantity_actual": "实际"})
        atomic_json(self.root / "config.local.json", self.config)
        path = self.workbook(f"custom_{self.day}.xlsx", headers=["代号", "可用", "实际"], rows=[["000001", 1000, 999]])
        book = load_workbook(path)
        book.active.title = "自定Sheet"
        book.save(path)
        book.close()
        service = self.service()
        self.assertEqual(service.refresh()["state"], "ready")
        self.assertEqual(service.records[0]["quantity"], 999)
        self.config["file_glob"] = "*.xls*"
        atomic_json(self.root / "config.local.json", self.config)
        self.assertEqual(load_config(self.root / "config.local.json")["file_glob"], "*.xls*")

    def test_batch_name_ambiguity_and_alias_conflicts_rejected_before_wind(self):
        self.workbook(rows=[["000001", "测试银行甲", 100, 100, 1, .05, "甲", 7], ["600000", "测试银行乙", 100, 100, 1, .05, "乙", 7]])
        service = self.service()
        service.refresh()
        with patch.object(service.provider, "market") as market:
            with self.assertRaisesRegex(ValueError, "匹配 2"):
                service.query("测试银行", input_targets={"测试银行": "10000"})
            with self.assertRaisesRegex(ValueError, "冲突"):
                service.query("000001 测试银行甲", input_targets={"000001": "10000", "测试银行甲": "20000"})
            market.assert_not_called()


class BatchTests(unittest.TestCase):
    def test_only_swap_source_is_labeled_as_own_inventory(self):
        self.assertEqual(source_note({"source": " 互换 "}), "自有券源，无需询其他交易台")
        for source in ("其他交易台", "非互换", "", None):
            self.assertEqual(source_note({"source": source}), "")

    def test_excel_two_columns_and_code_list(self):
        symbols, targets = parse_batch("证券代码\t目标名本（万元）\n000001\t1,000\n平安银行\t2.9\n600519")
        self.assertEqual(symbols, ["000001", "平安银行", "600519"])
        self.assertEqual(targets, {"000001": "10000000", "平安银行": "29000.0"})
        self.assertEqual(parse_batch("600519，000001\n平安银行;600519")[0], ["600519", "000001", "平安银行"])

    def test_invalid_input_has_no_partial_result(self):
        for text in ("000001\tNaN", "000001\t-1", "000001\t100\t100", "000001\t100\n000001\t200", ""):
            with self.assertRaises(ValueError):
                parse_batch(text)


class EntrypointTests(unittest.TestCase):
    def test_default_module_can_capture_native_demo(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "demo.png"
            result = subprocess.run([sys.executable, "-B", "-s", "-X", "utf8", "-m", "securities_lending", "--demo", "--capture", str(path)], cwd=ROOT.parent, capture_output=True, text=True, encoding="utf-8", timeout=25)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(path.is_file())
            self.assertEqual(path.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")


class WorkbenchTests(native_tests.DesktopCase):
    def test_multi_source_details_keep_rates_terms_and_sources_visible(self):
        self.query("600519 000001")
        self.assertEqual(len(self.window.items), 2)
        self.assertEqual(self.window.items[0]["quantity"], 20000)
        self.assertTrue(self.window.detail_panel.isVisible())
        detail = self.window.detail_table
        self.assertEqual(detail.rowCount(), 2)
        self.assertEqual([detail.item(row, 2).text() for row in range(2)], ["2.5%", "2.8%"])
        self.assertEqual([detail.item(row, 1).text() for row in range(2)], ["28天", "14天"])
        self.assertEqual([detail.item(row, 3).text() for row in range(2)], ["12,000", "8,000"])
        self.assertEqual(detail.item(0, 0).text(), "互换")
        self.assertEqual(detail.item(0, 6).text(), "自有券源，无需询其他交易台")
        self.assertEqual(detail.item(1, 6).text(), "—")
        self.assertIn("自有券源", self.window.detail.toPlainText())
        self.window.detail_button.setChecked(False)
        self.window.toggle_detail()
        self.assertFalse(self.window.detail_panel.isVisible())
        self.window.table.selectRow(1)
        self.assertEqual(detail.rowCount(), 1)
        self.window.table.selectRow(0)
        self.assertTrue(self.window.detail_panel.isVisible())

    def test_detail_copy_and_export_keep_individual_source_terms(self):
        self.query("600519")
        self.window.copy_detail()
        copied = list(csv.reader(io.StringIO(self.app.clipboard().text()), delimiter="\t"))
        self.assertEqual(len(copied), 3)
        headers = copied[0]
        self.assertEqual(copied[1][headers.index("期限")], "28天")
        self.assertEqual(copied[2][headers.index("期限")], "14天")
        self.assertEqual(copied[1][headers.index("券源来源")], "互换")
        self.assertIn("无需询其他交易台", copied[1][headers.index("券源说明")])
        self.window.detail_table.selectRow(1)
        QTest.keyClick(self.window.detail_table, Qt.Key.Key_C, Qt.KeyboardModifier.ControlModifier)
        copied = list(csv.reader(io.StringIO(self.app.clipboard().text()), delimiter="\t"))
        self.assertEqual(len(copied), 2)
        self.assertEqual(copied[1][headers.index("期限")], "14天")
        path = Path(self.tmp.name) / "sources.xlsx"
        self.window.export_to(path)
        book = load_workbook(path)
        self.addCleanup(book.close)
        summary = list(book["查询汇总"].values)
        self.assertEqual(summary[1][summary[0].index("券源笔数")], 2)
        self.assertEqual(summary[1][summary[0].index("期限（逐笔）")], "28天 / 14天")
        details = list(book["券源明细"].values)
        self.assertEqual(len(details), 3)
        self.assertEqual(details[1][details[0].index("券源来源")], "互换")
        self.assertEqual(details[1][details[0].index("券源说明")], "自有券源，无需询其他交易台")
        self.assertEqual([row[details[0].index("出借费率（%）")] for row in details[1:]], [2.5, 2.8])
        book.close()

    def test_export_failure_keeps_quote_and_can_retry_export(self):
        self.query("600519")
        market_status = self.window.wind_label.text()
        path = str(Path(self.tmp.name) / "retry.xlsx")
        with patch("securities_lending.desktop.QFileDialog.getSaveFileName", return_value=(path, "Excel 工作簿 (*.xlsx)")), patch("securities_lending.desktop.write_export", side_effect=[OSError("file locked"), None]) as writer:
            self.window.export()
            self.wait(lambda: not self.window.busy)
            self.assertEqual(self.window.wind_label.text(), market_status)
            self.assertIsNotNone(self.window.response)
            self.assertFalse(self.window.retry_button.isHidden())
            self.window.retry_action()
            self.wait(lambda: not self.window.busy)
            self.assertEqual(writer.call_count, 2)
            self.assertIn("已保存", self.window.statusBar().currentMessage())

    def test_extreme_manual_target_is_rejected_without_ui_overflow(self):
        self.query("600519")
        self.edit("600519.SH", "1e1000000")
        self.assertEqual(self.window.target_values()["600519.SH"], "invalid")
        self.calculate()
        self.assertEqual(self.window.items[0]["calculation"]["status"], "invalid")

    def test_editor_focus_tab_and_column_paste(self):
        self.query("600519 300750 000001")
        second = self.window.editors["300750.SZ"]
        second.setFocus()
        self.app.processEvents()
        self.assertEqual(self.window.table.currentRow(), 1)
        self.assertIn("300750.SZ", self.window.detail.toPlainText())
        self.window.paste_targets(1, "5000\n100")
        self.assertEqual(self.window.editors["000001.SZ"].text(), "100")
        QTest.keyClick(second, Qt.Key.Key_Tab)
        self.assertTrue(self.window.editors["000001.SZ"].hasFocus())
        before = dict(self.window.target_texts)
        self.window.paste_targets(1, "200\ninvalid")
        self.assertEqual(self.window.target_texts, before)
        self.assertFalse(self.window.export_button.isEnabled())

    def test_copy_selected_and_xlsx_include_trace_and_both_quantities(self):
        self.query("600519 300750")
        self.window.table.selectAll()
        self.window.copy_selected()
        copied = list(csv.reader(io.StringIO(self.app.clipboard().text()), delimiter="\t"))
        self.assertEqual(len(copied), 3)
        self.assertIn("库存市值（元，本次最新价）", copied[0])
        path = Path(self.tmp.name) / "result.xlsx"
        self.window.export_to(path)
        book = load_workbook(path)
        self.assertEqual(book.sheetnames, ["查询汇总", "券源明细", "数据说明"])
        self.assertEqual(book["查询汇总"].max_row, 3)
        self.assertEqual(book["券源明细"].max_row, 4)
        self.assertEqual(book["查询汇总"]["G2"].data_type, "n")
        self.assertEqual(book["查询汇总"]["D2"].data_type, "s")
        detail_headers = [cell.value for cell in book["券源明细"][1]]
        self.assertIn("实际可融股数", detail_headers)
        self.assertIn("SHA256", detail_headers)
        book.close()

    def test_open_folder_uses_downloaded_snapshot(self):
        with patch("securities_lending.desktop.QDesktopServices.openUrl", return_value=True) as opener:
            self.window.open_local_data()
        opened = Path(opener.call_args.args[0].toLocalFile())
        self.assertEqual(opened, Path(self.service.state["files"][0]["path"]).parent)
        self.assertTrue(opened.is_relative_to(self.service.data_root))

    def test_failed_sheet_opens_new_download_while_inventory_uses_cache(self):
        accepted = self.service.state["files"][0]["path"]
        source = Path(self.service.state["files"][0]["source"])
        book = load_workbook(source)
        book.active.title = "Sheet1"
        book.save(source)
        book.close()
        with patch("securities_lending.service.demo_source"):
            self.window.refresh()
            self.wait(lambda: not self.window.busy)
        status = self.service.state
        self.assertEqual(status["state"], "cached")
        self.assertEqual(status["files"][0]["path"], accepted)
        failed = Path(status["downloaded_files"][0]["path"])
        with patch("securities_lending.desktop.QDesktopServices.openUrl", return_value=True) as opener:
            self.window.open_local_data()
            self.assertEqual(Path(opener.call_args.args[0].toLocalFile()), failed.parent)
            self.window.open_local_excel()
            self.assertEqual(Path(opener.call_args.args[0].toLocalFile()), failed)
        self.assertIn("Sheet1", self.window.source_text())
        self.assertIn("当日缓存", self.window.source_text())
        self.assertIn(str(failed), self.window.source_text())
        self.assertIn("SHA256 一致", self.window.source_text())
        self.assertNotIn("\n", self.window.notice.text())
        # With no valid cache, the failed download remains inspectable, not usable inventory.
        with patch("securities_lending.service.demo_source"), patch("securities_lending.service.cached_manifest", return_value=None):
            self.window.refresh()
            self.wait(lambda: not self.window.busy)
        self.assertEqual(self.service.state["state"], "unavailable")
        self.assertIn("已下载", self.window.status_label.text())
        self.assertIn("未作为库存", self.window.source_label.text())
        self.assertEqual(self.window.local_files(), self.service.state["downloaded_files"])
        self.assertIn("读取失败", self.window.statusBar().currentMessage())

    def test_layout_keeps_table_visible_when_small_and_details_open(self):
        self.query("600519 300750 000001")
        for size in ((1920, 1000), (1280, 680), (1100, 740), (940, 560)):
            self.window.resize(*size)
            self.window.detail_button.setChecked(True)
            self.window.toggle_detail()
            self.app.processEvents()
            self.assertGreaterEqual(self.window.table.viewport().height(), 100, size)
            self.assertLessEqual(self.window.height(), size[1], size)

    def test_slow_background_query_keeps_event_loop_responsive(self):
        started, release = threading.Event(), threading.Event()
        original = self.service.query
        def slow(*args):
            started.set()
            release.wait(5)
            return original(*args)
        with patch.object(self.service, "query", side_effect=slow):
            self.window.symbols.setText("600519")
            self.window.query()
            self.wait(started.is_set)
            self.assertTrue(self.window.busy)
            self.window.resize(1100, 740)
            self.app.processEvents()
            self.assertFalse(self.window.progress.isHidden())
            release.set()
            self.wait(lambda: not self.window.busy)
        self.assertEqual(len(self.window.items), 1)

    def test_failed_query_retry_and_no_stale_market_value(self):
        self.window.symbols.setText("600519")
        with patch.object(self.service, "query", side_effect=RuntimeError("temporary failure")):
            self.window.query()
            self.wait(lambda: not self.window.busy)
        self.assertFalse(self.window.retry_button.isHidden())
        self.assertFalse(self.window.export_button.isEnabled())
        self.window.retry_action()
        self.wait(lambda: not self.window.busy)
        self.assertIsNotNone(self.window.items[0]["inventory_market_value"])

    def test_export_formula_like_source_text_is_literal(self):
        self.query("600519")
        self.window.items[0]["name"] = "=1+1"
        path = Path(self.tmp.name) / "literal.xlsx"
        self.window.export_to(path)
        book = load_workbook(path)
        self.assertEqual(book["查询汇总"]["E2"].data_type, "s")
        book.close()


if __name__ == "__main__":
    unittest.main()
