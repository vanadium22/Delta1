from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path

ALIASES = {
    "code": ["证券代码", "股票代码", "标的代码", "券代码", "代码", "Wind代码", "标的"],
    "name": ["证券名称", "证券简称", "股票名称", "股票简称", "标的名称", "名称", "简称"],
    "quantity": ["可融数量", "可融券数量", "可出借数量", "出借数量", "可用数量", "库存数量", "可借数量", "券源数量", "数量", "可融券余额"],
    "nominal": ["名本", "名义本金", "可融名本", "可融券名本", "可融金额", "名义金额"],
    "rate": ["出借费率", "融券费率", "借券费率", "出借利率", "费率", "参考费率", "年化费率"],
    "industry": ["所属行业", "所在行业", "申万一级行业", "申万行业", "行业名称", "行业"],
    "price": ["参考价格", "参考价", "收盘价", "最新价", "市价", "价格"],
    "term": ["期限", "出借期限", "融券期限", "期限天数"],
}


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def heading(value) -> str:
    value = re.sub(r"[（(][^）)]*[）)]", "", clean(value))
    return re.sub(r"[\s_%％/：:]", "", value).casefold()


def normalize_code(value) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(value) or value != int(value):
            return ""
        value = str(int(value)).zfill(6)
    text = clean(value).upper().replace(" ", "")
    # Excel sometimes exports codes as decimal strings.
    if re.fullmatch(r"\d{1,6}\.0+", text):
        text = text.split(".")[0].zfill(6)
    match = re.fullmatch(r"(?:(SH|SZ|BJ)[.:]?)?(\d{6})(?:[.](SH|SZ|BJ))?", text)
    if not match:
        return ""
    prefix, number, suffix = match.groups()
    if prefix and suffix and prefix != suffix:
        return ""
    exchange = suffix or prefix
    if not exchange:
        if number.startswith(("4", "8", "92")):
            exchange = "BJ"
        elif number.startswith(("5", "6", "9")):
            exchange = "SH"
        else:
            exchange = "SZ"
    return f"{number}.{exchange}"


def number(value, multiplier=1.0):
    text = clean(value).replace(",", "").replace("，", "").replace(" ", "")
    if not text or text in {"-", "--", "—", "N/A", "None"}:
        return None
    # A unit in the cell overrides the column/default unit.
    for suffix, scale in (("亿元", 1e8), ("万元", 1e4), ("万股", 1e4), ("亿", 1e8), ("万", 1e4), ("股", 1), ("元", 1)):
        if text.endswith(suffix):
            text, multiplier = text[:-len(suffix)], scale
            break
    try:
        result = float(text) * multiplier
        return result if math.isfinite(result) and result >= 0 else None
    except (TypeError, ValueError):
        return None


@dataclass
class Cell:
    value: object = None
    format: str = ""
    formula: bool = False


def read_rows(path: Path, sheet_name: str):
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("读取 .xls 需要 xlrd，请运行 install.cmd") from exc
        book = xlrd.open_workbook(str(path), formatting_info=True)
        try:
            name = next((n for n in book.sheet_names() if n.strip() == sheet_name), None)
            if name is None:
                raise ValueError(f"缺少「{sheet_name}」Sheet；已有：{'、'.join(book.sheet_names())}")
            sheet = book.sheet_by_name(name)
            for row in range(sheet.nrows):
                cells = []
                for col in range(sheet.ncols):
                    cell = sheet.cell(row, col)
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    cells.append(Cell(cell.value, fmt))
                yield cells
        finally:
            book.release_resources()
    else:
        from openpyxl import load_workbook
        values = load_workbook(path, read_only=True, data_only=True)
        formulas = None
        try:
            name = next((n for n in values.sheetnames if n.strip() == sheet_name), None)
            if name is None:
                raise ValueError(f"缺少「{sheet_name}」Sheet；已有：{'、'.join(values.sheetnames)}")
            formulas = load_workbook(path, read_only=True, data_only=False)
            for value_row, formula_row in zip(values[name].iter_rows(), formulas[name].iter_rows()):
                yield [Cell(v.value, v.number_format or "", f.data_type == "f") for v, f in zip(value_row, formula_row)]
        finally:
            values.close()
            if formulas:
                formulas.close()


def column_map(cells, overrides):
    found, ambiguous = {}, []
    for field, aliases in ALIASES.items():
        names = [overrides[field]] if field in overrides else aliases
        candidates = [i for i, cell in enumerate(cells) if heading(cell.value) in {heading(n) for n in names} and heading(cell.value)]
        if len(candidates) > 1:
            # Exact configured names can disambiguate columns with different units.
            exact = [i for i in candidates if clean(cells[i].value) == overrides.get(field)]
            if len(exact) == 1:
                candidates = exact
            else:
                ambiguous.append(field)
        if candidates:
            found[field] = candidates[0]
    return found, ambiguous


def parse_workbook(path: Path, config: dict):
    records, warnings = [], []
    mapping = None
    headers = []
    header_row = None
    for rownum, cells in enumerate(read_rows(path, config["sheet_name"]), 1):
        if mapping is None:
            candidate, ambiguous = column_map(cells, config["columns"])
            specified = config["header_row"]
            eligible = rownum == specified if specified else "code" in candidate and "quantity" in candidate
            if eligible:
                if ambiguous:
                    raise ValueError(f"第 {rownum} 行列名有歧义：{', '.join(ambiguous)}；请在 columns 中指定完整原始列名")
                if "code" not in candidate or "quantity" not in candidate:
                    raise ValueError("表头缺少证券代码或可融数量列")
                mapping, headers, header_row = candidate, [clean(c.value) for c in cells], rownum
                for field in ("nominal", "rate"):
                    if field not in mapping:
                        warnings.append(f"{path.name}：未识别{'名本' if field == 'nominal' else '出借费率'}列")
            elif rownum >= (specified or 50):
                raise ValueError("前 50 行未识别表头（需要证券代码、可融数量）；可配置 header_row 和 columns")
            continue
        def get(field):
            index = mapping.get(field)
            return cells[index] if index is not None and index < len(cells) else Cell()

        code = normalize_code(get("code").value)
        if not code:
            raw_code = clean(get("code").value)
            if raw_code and heading(raw_code) != heading(headers[mapping["code"]]) and not any(x in raw_code for x in ("合计", "说明", "备注", "总计")):
                warnings.append(f"{path.name}:{rownum} 无法识别证券代码「{raw_code}」，已跳过")
            continue
        q_header = headers[mapping["quantity"]]
        n_header = headers[mapping["nominal"]] if "nominal" in mapping else ""
        q_scale = 1e4 if "万股" in q_header or ("股" not in q_header and config["quantity_unit"] == "万股") else 1
        n_scale = 1e8 if "亿元" in n_header else 1e4 if "万元" in n_header else 1 if "元" in n_header else {"元": 1, "万元": 1e4, "亿元": 1e8}[config["nominal_unit"]]
        quantity = number(get("quantity").value, q_scale)
        nominal = number(get("nominal").value, n_scale)
        price = number(get("price").value)
        nominal_source = "券单" if nominal is not None else "缺失"
        if nominal is None and quantity is not None and price is not None:
            nominal, nominal_source = quantity * price, "估算：可融数量 × 券单价格"
        rate_cell = get("rate")
        rate_raw = clean(rate_cell.value)
        rate = number(rate_raw.replace("%", "").replace("％", ""))
        if rate is not None and "%" not in rate_raw and "％" not in rate_raw:
            if "%" in rate_cell.format or config["numeric_rate_unit"] == "decimal":
                rate *= 100
        for field in ("quantity", "nominal", "rate"):
            cell = get(field)
            if cell.formula and cell.value is None:
                warnings.append(f"{path.name}:{rownum} {field} 公式没有缓存值，请用 Excel 重新计算并保存")
        if quantity is None:
            warnings.append(f"{path.name}:{rownum} 可融数量缺失或无效，保留该记录但不计入合计")
        records.append({
            "code": code, "name": clean(get("name").value), "quantity": quantity,
            "nominal": nominal, "nominal_source": nominal_source, "rate": rate,
            "rate_raw": rate_raw, "industry": clean(get("industry").value),
            "industry_source": "券单" if clean(get("industry").value) else "待补充",
            "term": clean(get("term").value), "file": path.name, "row": rownum,
        })
    if mapping is None:
        raise ValueError("Sheet 为空或未识别表头")
    return records, warnings, {"header_row": header_row, "columns": {k: headers[v] for k, v in mapping.items()}}
